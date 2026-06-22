"""Native scatter (XY) chart builder (design §9a, Task 5.8).

Each category in the SeriesResult becomes one (x, y) data point.
x = cell(category, x_segment).<statistic>
y = cell(category, y_segment).<statistic>
"""
from __future__ import annotations
from pptx.chart.data import XyChartData
from pptx.enum.chart import XL_CHART_TYPE
from reportbuilder.render.base import RenderContext
from reportbuilder.stats.series import SeriesResult


def xy_chart_data(series: SeriesResult, scatter_xy: tuple[str, str]) -> XyChartData:
    """Build XyChartData from a SeriesResult using two named segments as axes.

    Each category becomes one data point: x from x_segment, y from y_segment.
    """
    x_seg, y_seg = scatter_xy
    xd = XyChartData()
    s = xd.add_series("points")
    for cat in series.categories:
        x = getattr(series.cell(cat, x_seg), series.statistic)
        y = getattr(series.cell(cat, y_seg), series.statistic)
        s.add_data_point(float(x), float(y))
    return xd


def build_scatter(ctx: RenderContext):
    """RenderContext-based builder for XY scatter charts.

    Returns a graphic frame containing an XY_SCATTER chart. Does NOT call
    apply_elements — deck assembly (Task 5.14) handles annotations.
    """
    if ctx.spec.scatter_xy is None:
        raise ValueError("scatter requires scatter_xy (two numeric axis segments)")
    xd = xy_chart_data(ctx.series, ctx.spec.scatter_xy)
    gf = ctx.slide.shapes.add_chart(
        XL_CHART_TYPE.XY_SCATTER,
        ctx.slot.left,
        ctx.slot.top,
        ctx.slot.width,
        ctx.slot.height,
        xd,
    )
    return gf
