"""Native funnel chart builder: stacked-bar approximation with transparent centering spacer (Task 5.9).

The funnel silhouette is achieved by:
  1. Adding a leading transparent (noFill/background) spacer series whose value
     equals (max - value) / 2, so each bar is visually centred.
  2. Adding the real value series on top of the spacer.
  3. Setting plot overlap = 100 so both series stack correctly on the same bar.

The result is a native c:barChart (BAR_STACKED), satisfying the editability
gate REQ-C-23a — zero picture shapes are produced.

REQ-C-13, C-23a.
"""
from __future__ import annotations
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.dml.color import RGBColor
from reportbuilder.render.base import RenderContext


def funnel_spacer_values(values: tuple[float, ...]) -> tuple[float, ...]:
    """Per-category leading spacer = (max - value) / 2, centering each bar (funnel silhouette).

    Args:
        values: Numeric values for each category, in category order.

    Returns:
        Tuple of spacer values; the largest bar gets spacer 0.0.
    """
    mx = max(values) if values else 0.0
    return tuple((mx - v) / 2.0 for v in values)


def build_funnel(ctx: RenderContext):
    """Build a native funnel chart as a stacked horizontal bar with a transparent spacer.

    The first series is an invisible centering spacer (noFill); the second is
    the visible value series. plot.overlap is set to 100 so bars are stacked.

    Args:
        ctx: RenderContext with slide, slot geometry, series data, and style.

    Returns:
        The graphic frame (python-pptx GraphicFrame) containing the chart.
    """
    series = ctx.series
    values = tuple(
        float(series.cell(c, "Total").value(series.statistic) or 0.0)
        for c in series.categories
    )
    spacers = funnel_spacer_values(values)

    cd = CategoryChartData()
    cd.categories = series.categories
    cd.add_series("_spacer", spacers)
    cd.add_series("value", values)

    gf = ctx.slide.shapes.add_chart(
        XL_CHART_TYPE.BAR_STACKED,
        ctx.slot.left,
        ctx.slot.top,
        ctx.slot.width,
        ctx.slot.height,
        cd,
    )

    plot = gf.chart.plots[0]
    plot.overlap = 100

    # Spacer: transparent (background / noFill) so only the value bar is visible
    plot.series[0].format.fill.background()

    # Value series: solid fill using palette color 0
    plot.series[1].format.fill.solid()
    plot.series[1].format.fill.fore_color.rgb = RGBColor.from_string(ctx.style.color_for(0))

    return gf
