"""De-novo native column chart with injected manual layout + data-label pos (R3)."""
from __future__ import annotations
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from reportbuilder.render.layout import solve_column_layout
from reportbuilder.stats.series import SeriesResult

_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"

def _value_for(series: SeriesResult, cat: str, seg: str) -> float:
    cell = series.cell(cat, seg)
    return float(getattr(cell, series.statistic))

def _manual_layout_xml(x: float, y: float, w: float, h: float) -> str:
    return (
        f'<c:layout xmlns:c="{_C}"><c:manualLayout>'
        '<c:layoutTarget val="inner"/>'
        '<c:xMode val="edge"/><c:yMode val="edge"/>'
        f'<c:x val="{x}"/><c:y val="{y}"/><c:w val="{w}"/><c:h val="{h}"/>'
        '</c:manualLayout></c:layout>'
    )

def _dlbls_xml() -> str:
    return (
        f'<c:dLbls xmlns:c="{_C}">'
        '<c:dLblPos val="outEnd"/>'
        '<c:showLegendKey val="0"/><c:showVal val="1"/>'
        '<c:showCatName val="0"/><c:showSerName val="0"/>'
        '<c:showPercent val="0"/><c:showBubbleSize val="0"/>'
        '</c:dLbls>'
    )

def build_column_chart(slide, slot, series: SeriesResult, *, point_size: int = 10):
    cd = CategoryChartData()
    cd.categories = series.categories
    for seg in series.segments:
        cd.add_series(seg, tuple(_value_for(series, c, seg) for c in series.categories))
    gf = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, slot.left, slot.top, slot.width, slot.height, cd,
    )
    chart = gf.chart
    cs = chart._chartSpace
    plot_area = cs.find(".//" + qn("c:plotArea"))
    # 1. manual plot-area layout as first child of plotArea
    lay = solve_column_layout(series.categories, series.segments, point_size=point_size).plot
    layout_el = parse_xml(_manual_layout_xml(lay.x, lay.y, lay.w, lay.h))
    plot_area.insert(0, layout_el)
    # 2. data-label position into the barChart element (before the c:axId refs)
    bar_chart = plot_area.find(qn("c:barChart"))
    dlbls = parse_xml(_dlbls_xml())
    first_axid = bar_chart.find(qn("c:axId"))
    if first_axid is not None:
        first_axid.addprevious(dlbls)
    else:
        bar_chart.append(dlbls)
    return gf
