"""A combo formats each measure among its own numbers.

The automatic decimal count is chosen from the values a number stands among, so
that a column of shares all carry the same number of decimals. The combo pooled
BOTH halves: a secondary mean of 2.8-3.4 needs a decimal, so every whole
percentage of the bars came out as "19.0 %", "21.0 %". A mean says nothing about
how precisely a share should be printed. (visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("25-34", "35-44", "45-54", "55-64")
_SHARES = (19.0, 21.0, 20.0, 40.0)
_MEANS = (3.37, 3.15, 3.0, 2.83)


def _labels(secondary_kind: str) -> list[str]:
    cells = {}
    for c, p, m in zip(_CATS, _SHARES, _MEANS):
        cells[(c, "Ikä")] = Cell(pct=p)
        cells[(c, "Luotettava")] = Cell(pct=m)
    series = SeriesResult(
        categories=_CATS, segments=("Ikä", "Luotettava"), cells=cells,
        base_n={"Ikä": 1000, "Luotettava": 800}, statistic="pct",
        segment_statistics={"Ikä": "pct", "Luotettava": "mean"})
    spec = ChartSpec(question_ref="q", chart_type="combo", statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(),
                     options={"combo_secondary": "Luotettava",
                              "combo_secondary_type": secondary_kind})
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    seen: list[str] = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        for axes in self.axes:
            seen.extend(t.get_text() for t in axes.texts
                        if t.get_gid() == VALUE_GID and t.get_visible())
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        _f.Figure.savefig = original
    return seen


@pytest.mark.parametrize("secondary_kind", ["line", "bar"])
def test_whole_shares_stay_whole_beside_a_mean(secondary_kind):
    labels = _labels(secondary_kind)
    shares = [t for t in labels if t.endswith("%")]
    assert shares, f"no share was labelled at all: {labels}"
    assert all("." not in t for t in shares), (
        f"whole percentages were printed with a decimal: {shares}")


@pytest.mark.parametrize("secondary_kind", ["line", "bar"])
def test_the_mean_keeps_its_decimal(secondary_kind):
    labels = _labels(secondary_kind)
    means = [t for t in labels if not t.endswith("%")]
    assert means and all("." in t for t in means), labels
