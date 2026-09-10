"""A cross-tab's group label stands clear of the category labels.

"Kuvassa olevat selitystekstit menevät välillä päällekkäin" — on a stacked bar
crossed by two classifiers, the rotated primary label ("Suomi") was printed
straight through the secondary tick labels ("hyvinvointialueen palveluksessa").

The rotated label sat at a fixed x of -0.13 in AXES coordinates, which means
"13% of the plot width to the left of the axis" — room enough for "Total", and
205px short for a long Finnish compound. Measured against the widest label
instead, so the two read as two columns.

Reproduced from the customer's own numbers before the fix: 12 collisions, the
worst 205.7px. (Johan, 2026-09-10)
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B

#: Long enough to reach past the old fixed offset — the customer's own values.
_SECTORS = ["yksityisellä sektorilla", "valtion palveluksessa",
            "kunnan palveluksessa", "hyvinvointialueen palveluksessa",
            "järjestön palveluksessa", "muualla"]
_COUNTRIES = ["Suomi", "Ruotsi", "Saksa"]


def _overlaps(slot_in: float = 11.6):
    """Rotated group labels that intersect a y tick label, as drawn."""
    scale = Variable(name="q", label="Merkitys", measurement="nominal",
                     missing_values=[],
                     value_labels=[ValueLabel(value=float(i + 1), label=str(i + 1))
                                   for i in range(7)])
    country = Variable(name="maa", label="country", measurement="nominal",
                       missing_values=[],
                       value_labels=[ValueLabel(value=float(i + 1), label=c)
                                     for i, c in enumerate(_COUNTRIES)])
    sector = Variable(name="sek", label="Sektori", measurement="nominal",
                      missing_values=[],
                      value_labels=[ValueLabel(value=float(i + 1), label=x)
                                    for i, x in enumerate(_SECTORS)])
    rows = [{"q": float(i % 7 + 1), "maa": float(i % 3 + 1), "sek": float(i % 6 + 1)}
            for i in range(2914)]
    model = QuestionModel(
        variables={"q": scale, "maa": country, "sek": sector},
        questions=[Question(qid="q", text="Merkitys", kind="single", variables=("q",)),
                   Question(qid="maa", text="country", kind="single", variables=("maa",)),
                   Question(qid="sek", text="Sektori", kind="single", variables=("sek",))])
    spec = ChartSpec(question_ref="q", chart_type="stacked_horizontal_bar",
                     statistic="pct", classifying_var="maa", classifying_var_2="sek",
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(slot_in), height=Inches(4.0), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    found: dict = {}
    original = B.render_png

    def _spy(fig):
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        ax = fig.axes[0]
        rotated = [(t, t.get_window_extent(r)) for t in ax.texts
                   if round(t.get_rotation()) == 90]
        ticks = [(t, t.get_window_extent(r)) for t in ax.get_yticklabels()
                 if t.get_text()]
        found["rotated"] = [t.get_text() for t, _b in rotated]
        found["hits"] = [
            (g.get_text(), lbl.get_text(), round(gb.x1 - lb.x0, 1))
            for g, gb in rotated for lbl, lb in ticks
            if gb.x1 > lb.x0 + 0.5 and gb.y1 > lb.y0 and lb.y1 > gb.y0]
        return original(fig)

    B.render_png = _spy
    try:
        B.build_image_bar_stacked(ctx)
    finally:
        B.render_png = original
    return found


def test_the_group_labels_are_drawn_at_all():
    got = _overlaps()
    assert set(_COUNTRIES) <= set(got["rotated"]), got["rotated"]


def test_no_group_label_touches_a_category_label():
    got = _overlaps()
    assert got["hits"] == [], got["hits"][:4]


@pytest.mark.parametrize("slot_in", [8.0, 9.5, 11.6, 13.3])
def test_it_holds_at_every_slide_width(slot_in):
    """The old offset was a fraction of the plot, so it failed differently at
    each width — the fix is measured, and has to hold across them."""
    assert _overlaps(slot_in)["hits"] == []
