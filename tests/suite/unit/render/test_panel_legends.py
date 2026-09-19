"""Small multiples: every panel's legend fits its panel, and a group is one colour.

Two classifiers drawn as panels gave each panel a legend of its own in up to
three columns whatever the panel's width. With names like
"Pääkaupunkiseudulla (n=118)" each legend was wider than its panel, and side by
side they printed through each other. Colours were keyed by position within a
panel, so a panel missing a group shifted every colour after it: one colour
meant two groups on the same slide. (visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("Erittäin huono", "Huono", "Hyvä", "Erittäin hyvä", "En osaa sanoa")
_REGIONS = ("Pääkaupunkiseudulla", "Muualla Etelä-Suomessa", "Länsi-Suomessa",
            "Pohjois-Suomessa", "Itä-Suomessa")
#: The fourth panel has only two of the regions — the one that shifted colours.
_PANELS = {"Mieheksi": _REGIONS, "Naiseksi": _REGIONS, "Muuksi": _REGIONS,
           "En halua sanoa": (_REGIONS[0], _REGIONS[3])}


def _drawn(chart_type: str):
    segs, primary, cells, base = [], {}, {}, {}
    for p, regions in _PANELS.items():
        for j, reg in enumerate(regions):
            seg = f"{p} · {reg}"
            segs.append(seg)
            primary[seg] = p
            base[seg] = 100 + j
            for i, c in enumerate(_CATS):
                cells[(c, seg)] = Cell(pct=float(5 + 7 * i + j))
    series = SeriesResult(categories=_CATS, segments=tuple(segs), cells=cells,
                          base_n=base, statistic="pct", segment_primary=primary)
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="p", classifying_var_2="r",
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(),
                     options={"xtab_layout": "small_multiples"})
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    got: dict = {"legends": [], "colours": {}}

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        for ax in self.axes:
            leg = ax.get_legend()
            if leg is None:
                continue
            got["legends"].append(leg.get_window_extent(r))
            for text, patch in zip(leg.get_texts(), leg.get_patches()):
                name = text.get_text().split(" (n=")[0]
                got["colours"].setdefault(name, set()).add(
                    tuple(round(c, 4) for c in patch.get_facecolor()))
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        _f.Figure.savefig = original
    return got


@pytest.mark.parametrize("chart_type", ["vertical_bar", "horizontal_bar"])
def test_no_panel_legend_runs_into_another(chart_type):
    boxes = _drawn(chart_type)["legends"]
    assert len(boxes) == len(_PANELS)
    for i, a in enumerate(boxes):
        for b in boxes[i + 1:]:
            overlap = (min(a.x1, b.x1) - max(a.x0, b.x0) > 0.5
                       and min(a.y1, b.y1) - max(a.y0, b.y0) > 0.5)
            assert not overlap, "two panels' legends are printed over each other"


@pytest.mark.parametrize("chart_type", ["vertical_bar", "horizontal_bar"])
def test_a_group_has_one_colour_in_every_panel(chart_type):
    colours = _drawn(chart_type)["colours"]
    assert set(colours) == set(_REGIONS)
    for name, seen in colours.items():
        assert len(seen) == 1, f"{name} is drawn in {len(seen)} different colours"
    assert len({next(iter(c)) for c in colours.values()}) == len(_REGIONS), \
        "two different groups share one colour"
