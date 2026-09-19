"""Every line a combo draws stays inside the axis it is drawn on.

Split by a classifier and with no secondary variable, the first group is the
bars and every other group a line — the same measure, so the lines share the
bars' ruler rather than getting one of their own. That ruler was built from the
BARS alone. On Attendo's "Missä päin Suomea asut?" by gender the bars top out at
30 % and the other groups reach 50 % and 100 %, so those lines ran off the top
of the plot and the chart showed numbers it had cut away. (visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("Etelä", "Länsi", "Pohjoinen", "Itä")
#: The bars (first group) stay low; the lines climb far above them.
_VALUES = {"Mies": (25.0, 30.0, 12.0, 11.0),
           "Muu": (25.0, 50.0, 0.0, 25.0),
           "EOS": (0.0, 0.0, 100.0, 0.0)}


def _series() -> SeriesResult:
    cells = {(c, g): Cell(pct=v[i]) for g, v in _VALUES.items()
             for i, c in enumerate(_CATS)}
    return SeriesResult(categories=_CATS, segments=tuple(_VALUES), cells=cells,
                        base_n={g: 100 for g in _VALUES}, statistic="pct")


def _plotted_lines(kind: str) -> list[tuple[float, float, list[float]]]:
    """(ylim bottom, ylim top, y data) of every data line on the saved figure."""
    spec = ChartSpec(question_ref="q", chart_type="combo", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(),
                     options={"combo_secondary_type": kind})
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(), fmt=NumberFormat())
    seen: list = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        for axes in self.axes:
            lo, hi = axes.get_ylim()
            for ln in axes.lines:
                if len(ln.get_xdata()) == len(_CATS):
                    seen.append((lo, hi, [float(y) for y in ln.get_ydata()]))
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        _f.Figure.savefig = original
    return seen


@pytest.mark.parametrize("kind", ["line", "area"])
def test_no_line_runs_off_the_top_of_its_axis(kind):
    lines = _plotted_lines(kind)
    assert lines, "the other groups were not drawn as lines at all"
    for lo, hi, ys in lines:
        for y in ys:
            assert lo <= y <= hi, (
                f"a {kind} value of {y} is drawn on an axis that ends at {hi}")
