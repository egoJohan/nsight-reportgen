"""A demographics grid's cells: the question above its chart, the chart legible.

Two defects the visual QA matrix found on the same slide (2026-09-19):

* each cell's question had a fixed 0.32in band, so a question that wraps to
  three lines printed straight over the top of its own chart;
* each cell's chart was drawn on a figure of at least 9 x 4.5 inches — the size
  a slide-wide chart is drawn at — and scaled down into a cell a fraction of
  that, so a six-chart grid printed its 9pt labels at about 3pt.
"""
from __future__ import annotations

import io

import pandas as pd
import pytest
from PIL import Image
from pptx.enum.shapes import MSO_SHAPE_TYPE

from reportbuilder.export.pptx_build import build_presentation
from reportbuilder.model.question import Question, QuestionModel, ValueLabel, Variable
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, Report, SortSpec
from reportbuilder.render.image.slide_chrome import wrapped_line_count

pytestmark = pytest.mark.unit

_LONG = ("Onko sinulla kokemusta erilaisista hoivapalveluista, kuten vanhusten tai "
         "vammaisten ihmisten hoivakodeista tai erilaisista kotihoidon palveluista? "
         "Valitse kaikki sopivat")
#: What a chart drawn at 200 dpi may be shrunk to on the slide, at least. At
#: 0.6 a 9pt label lands at 5.4pt; the old cells went to about 0.35.
_MIN_SCALE = 0.6


def _model(n: int):
    variables, questions, data = {}, [], {}
    for i in range(n):
        name = f"q{i}"
        labels = tuple(ValueLabel(float(k), f"Vaihtoehto {k}") for k in range(1, 6))
        variables[name] = Variable(name, _LONG if i == n - 1 else f"Kysymys {i}",
                                   "categorical", labels, frozenset())
        questions.append(Question(qid=name, kind="single", variables=(name,),
                                  text=variables[name].label))
        data[name] = [float(1 + (r * (i + 1)) % 5) for r in range(200)]
    return QuestionModel(variables=variables, questions=questions), pd.DataFrame(data)


def _slide(n: int):
    model, df = _model(n)
    grid = ChartSpec(question_ref="", chart_type="demographics_grid", statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="auto",
                     elements=ElementToggles(), slide_title="Vastaajat",
                     options={"charts": [{"question_ref": f"q{i}",
                                          "chart_type": "vertical_bar" if i % 2 else "horizontal_bar"}
                                         for i in range(n)]})
    prs = build_presentation(Report(name="r", render_mode="image", template_ref="",
                                    charts=(grid,)), model, df)
    return prs.slides[0]


@pytest.mark.parametrize("n", [4, 6])
def test_every_question_ends_above_its_chart(n):
    slide = _slide(n)
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == n
    titles = [s for s in slide.shapes if s.has_text_frame
              and s.text_frame.text.startswith(("Kysymys", "Onko"))]
    assert len(titles) == n
    for pic in pics:
        # The question box directly above this picture.
        centre = pic.left + pic.width // 2
        above = [t for t in titles
                 if t.left <= centre <= t.left + t.width and t.top < pic.top]
        box = max(above, key=lambda t: t.top)
        # Where the TEXT ends, not the box: a box keeps its height while the
        # words overflow it, which is exactly how the question reached the bars.
        text = box.text_frame.text
        size = box.text_frame.paragraphs[0].runs[0].font.size.pt
        needed = wrapped_line_count(text, box.width, size) * size * 1.2 / 72 * 914400
        assert box.top + needed <= pic.top, f"{text[:30]!r} runs into its chart"


@pytest.mark.parametrize("n", [4, 6])
def test_no_cell_chart_is_shrunk_out_of_legibility(n):
    for pic in [s for s in _slide(n).shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]:
        px_w, _px_h = Image.open(io.BytesIO(pic.image.blob)).size
        drawn_in = px_w / 200.0
        placed_in = pic.width / 914400.0
        assert placed_in / drawn_in >= _MIN_SCALE, (
            f"a cell chart drawn {drawn_in:.1f}in wide is placed at {placed_in:.1f}in")
