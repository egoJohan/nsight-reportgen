"""A lone summary category is not printed under the line chart's axis.

`_summary` builds ONE category — the question itself — times the classifier's
groups. `_category_ticks` already blanks that tick for bars, and says why: it
"stands under every one of them, tells none apart, and repeats the question the
subtitle already carries — printed rotated under the axis where it reads as an
axis title nobody asked for" (reported 2026-09-08).

The line builder set its tick labels itself and never called that rule, so the
question came back, rotated, under a line chart of a sum. Same rule, same
reason, one implementation. (Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

QUESTION = "Kuinka ylpeä olet omasta työstäsi?"
GROUPS = ("Alle 30", "30-44", "60 tai yli")


def _series(cats, statistic, groups=GROUPS):
    cells = {}
    for i, c in enumerate(cats):
        for j, g in enumerate(groups):
            cells[(c, g)] = (Cell(**{statistic: float(100 + i * 10 + j)})
                             if statistic in ("pct", "count", "mean")
                             else Cell(extra=((statistic, float(100 + i * 10 + j)),)))
    return SeriesResult(categories=tuple(cats), segments=tuple(groups), cells=cells,
                        base_n={"Total": 1018, **{g: 300 for g in groups}},
                        statistic=statistic)


def _xticks(cats, statistic, groups=GROUPS) -> list[str]:
    spec = ChartSpec(
        question_ref="q", chart_type="line", statistic=statistic,
        classifying_var="ikaryhma", number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(cats, statistic, groups),
        fmt=spec.number_format)
    seen: list[str] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        for axes in self.axes:
            seen.extend(t.get_text() for t in axes.get_xticklabels())
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["line"](ctx)
    finally:
        Figure.savefig = original
    # Unwrapped: a long tick is wrapped for drawing, which is not what
    # these tests are about.
    return [t.replace("\n", " ") for t in seen if t.strip()]


@pytest.mark.parametrize("statistic", ["sum", "mean", "median"])
def test_a_lone_summary_category_prints_no_tick(statistic):
    """The defect: the question reappeared rotated under the axis.

    With ONE segment, so there is nothing to transpose. A summary split by a
    classifier is turned into one point per group by `_as_one_series_per_group`
    and its ticks are then the groups' names, which must stay — see
    `test_line_transposes_a_lone_summary`. This is the un-split slide, where the
    lone category really is the question and really has nothing to say.
    """
    assert _xticks([QUESTION], statistic, groups=("Total",)) == [], statistic


def test_a_distribution_keeps_its_category_ticks():
    """Several categories are real answers and must stay named."""
    cats = ["1- Ei lainkaan", "2", "3"]
    assert _xticks(cats, "pct") == cats


def test_a_lone_ANSWER_category_keeps_its_tick():
    """One category left after the empty ones were hidden is still an answer,
    not the question — blanking it would leave an unnamed point."""
    assert _xticks(["Attendo"], "pct") == ["Attendo"]
