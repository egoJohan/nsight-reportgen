"""Growing a chart taller must not end up shrinking it.

A horizontal bar chart grows its figure taller as categories increase, so every
row keeps space for a wrapped label at a legible size instead of shrinking the
font (`bars.py`, `new_tall_figure`). But `place_picture_square` letterboxes:
it scales the PNG by `min(slot_w/px_w, slot_h/px_h)`. Once the figure is taller
than the slot's own aspect, height is the limiting dimension and the whole
picture is scaled DOWN — width goes unused and every font on it shrinks with it.

Measured on a 12.3x4.4in slot before the cap: 25 categories produced a
12.3x14.2in figure, placed at 31 % of the slot width, which turns a 9 pt tick
label into 2.8 pt on the slide. The growth rule was defeating its own purpose,
and not only in the pathological case — the loss starts at 7 categories.

So the figure may grow as tall as it likes up to the slot's aspect, and no
further. Past that point the rows compress and the builders' existing font
floors (`_MIN_LABEL_BAR_PT`, the value-label floors) do the work they were
written for — which leaves the labels BIGGER on the slide than growing did.
(Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import new_figure_grid, new_tall_figure
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

SLOT_W_IN, SLOT_H_IN = 12.3, 4.4
SLOT_ASPECT = SLOT_W_IN / SLOT_H_IN


def _ctx(series: SeriesResult | None = None, chart_type: str = "horizontal_bar",
         *, slot_in: tuple[float, float] = (SLOT_W_IN, SLOT_H_IN)):
    spec = ChartSpec(
        question_ref="q1", chart_type=chart_type, statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles(title=True, subtitle=True, legend=True, n=True,
                                axis_names=True, filter_var=True,
                                data_labels=True),
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    return RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(slot_in[0]), height=Inches(slot_in[1]), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)


def _aspect(fig: Figure) -> float:
    w, h = fig.get_size_inches()
    return w / h


# ---- the cap ---------------------------------------------------------------

def test_a_tall_figure_never_goes_past_the_slots_aspect():
    """25 categories asked for 14.2in against a 4.4in slot."""
    fig, _ax = new_tall_figure(_ctx(), 14.2)
    assert _aspect(fig) >= SLOT_ASPECT - 1e-6, (
        f"figure is {fig.get_size_inches()}, more portrait than the slot "
        f"({SLOT_ASPECT:.2f}) — everything past that is letterboxed away")


def test_a_grid_figure_is_capped_the_same_way():
    """Small multiples and separate panels grow by the same rule and are placed
    by the same one."""
    fig, _axes = new_figure_grid(_ctx(), 2, tall_in=14.2)
    assert _aspect(fig) >= SLOT_ASPECT - 1e-6, fig.get_size_inches()


def test_a_stacked_grid_gets_one_slot_aspect_per_ROW():
    """A second row of panels is a second chart's worth of content, not roomier
    rows. Capping the pair at one slot-aspect puts the second row's rotated
    ticks and legend on top of the first row's — the defect
    `test_vertical_stacked_panels_grow_figure_height` catches. So the ceiling is
    per band."""
    one = new_figure_grid(_ctx(), 2, tall_in=99.0, rows=1)[0].get_size_inches()[1]
    two = new_figure_grid(_ctx(), 4, tall_in=99.0, rows=2)[0].get_size_inches()[1]
    assert two == pytest.approx(one * 2)


def test_a_stacked_grid_is_still_capped():
    """Per-row is a ceiling too, not a licence — `rows=2` does not mean any
    height at all."""
    fig, _axes = new_figure_grid(_ctx(), 4, tall_in=99.0, rows=2)
    assert _aspect(fig) >= SLOT_ASPECT / 2 - 1e-6, fig.get_size_inches()


def test_on_a_full_width_slot_the_cap_is_the_slot_itself():
    """`w_in` is the slot's own width here, so the aspect ceiling works out to
    exactly the slot height: the figure is the slot, and the growth that used to
    run past it bought nothing the letterbox did not take straight back."""
    fig, _ax = new_tall_figure(_ctx(), 14.2)
    assert fig.get_size_inches() == pytest.approx([SLOT_W_IN, SLOT_H_IN])


def test_a_narrow_slot_still_has_room_to_grow_into():
    """The cap is an aspect, not a fixed size. A slot narrower than the 9in
    minimum figure width gets a figure WIDER than the slot, so there is real
    headroom before the aspect is reached — 4.5in here — and a request inside
    that headroom is honoured."""
    narrow = dict(slot_in=(6.0, 3.0))
    assert new_tall_figure(_ctx(**narrow), 4.0)[0].get_size_inches()[1] \
        == pytest.approx(4.0)
    assert new_tall_figure(_ctx(**narrow), 9.0)[0].get_size_inches()[1] \
        == pytest.approx(4.5)


def test_the_slot_height_is_still_the_floor():
    """A request SHORTER than the slot still fills the slot."""
    fig, _ax = new_tall_figure(_ctx(), 2.0)
    assert fig.get_size_inches()[1] == pytest.approx(SLOT_H_IN)


# ---- what the reader sees --------------------------------------------------

def _many_categories(n: int) -> SeriesResult:
    cats = tuple(str(i + 1) for i in range(n))
    return SeriesResult(
        categories=cats, segments=("Total",),
        cells={(c, "Total"): Cell(pct=4.0 + (i % 3)) for i, c in enumerate(cats)},
        base_n={"Total": 1018}, statistic="pct")


@pytest.fixture
def placed_width_fraction(monkeypatch):
    """How much of the slot's width the picture actually covers, computed the
    way `place_picture_square` computes it."""
    got: list[float] = []
    original = Figure.savefig

    def spy(self, path, *a, **k):
        r = original(self, path, *a, **k)
        from PIL import Image
        try:
            with Image.open(path) as im:
                px_w, px_h = im.size
        except Exception:  # noqa: BLE001 — a non-path target is not the picture
            return r
        scale = min(SLOT_W_IN / px_w, SLOT_H_IN / px_h)
        got.append(px_w * scale / SLOT_W_IN)
        return r

    monkeypatch.setattr(Figure, "savefig", spy)
    return got


@pytest.mark.parametrize("n", [7, 10, 15, 25])
def test_a_long_category_list_still_uses_the_slide(placed_width_fraction, n):
    IMAGE_BUILDERS["horizontal_bar"](_ctx(_many_categories(n)))
    assert placed_width_fraction, "nothing was saved; the fixture is wrong"
    used = max(placed_width_fraction)
    assert used >= 0.90, (
        f"{n} categories cover {used:.0%} of the slot width — the rest of the "
        f"slide is empty and the chart is scaled down to match")
