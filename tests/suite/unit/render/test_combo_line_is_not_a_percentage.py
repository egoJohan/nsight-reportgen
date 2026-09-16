"""The combo line is a mean, and a mean is not a percentage.

"Combo chartissa toisen muuttujan arvoihin kuvaajassa tulee mukaan %-merkki
vaikka ne eivät ole prosentteja."

A two-variable combo puts the question's distribution on the bars (percentages,
left axis) and the MEAN of a numeric secondary variable on the line (right
axis). `tyoelamaindeksi` on a 1-8 scale was printing `3.6 %` … `7.5 %`.

The cause is a shortcut in `_combo_two_var`: the secondary mean is stored in
`Cell.pct` and the whole result declares `statistic="pct"`, because that is what
made the renderer plot it on the right axis without changes. Storing it there is
fine — it is the value slot the renderer reads. Declaring the SERIES to be
percentages is not: the renderer formats both segments from that one word, so
the mean is handed a percent sign it never had.

Note the other combo path — a classifying variable rather than a secondary
variable — where both segments genuinely ARE percentages and must keep the sign.
That is what the last test here holds down. (Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pandas as pd
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.question import (
    Question, QuestionModel, ValueLabel, Variable,
)
from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats import engine
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

#: The reported slide: "Kuinka ylpeä olet omasta työstäsi?" 1-7, with the
#: työelämäindeksi (a 1-8 mean) as the line.
_CATS = ("1- Ei lainkaan ylpeä", "2", "3", "4", "5", "6", "7- Erittäin ylpeä")
_BARS = (3.0, 3.0, 6.0, 12.0, 23.0, 31.0, 22.0)     # a distribution, sums to 100
_LINE = (3.6, 4.6, 5.5, 6.0, 6.5, 7.1, 7.5)         # an index on a 1-8 scale
_QUESTION, _INDEX = "Kuinka ylpeä olet omasta työstäsi?", "tyoelamaindeksi"


def _two_var_series(**kw) -> SeriesResult:
    cells = {}
    for c, b, l in zip(_CATS, _BARS, _LINE):
        cells[(c, _QUESTION)] = Cell(pct=b)
        cells[(c, _INDEX)] = Cell(pct=l)      # the mean, in the value slot
    return SeriesResult(
        categories=_CATS, segments=(_QUESTION, _INDEX), cells=cells,
        base_n={"Total": 1018, _QUESTION: 1018, _INDEX: 1018},
        statistic="pct", **kw)


@pytest.fixture
def drawn(monkeypatch):
    """Every value text on the figure, per axes: index 0 is the bars, index 1
    the twinx the line lives on."""
    seen: list[tuple[int, str]] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        for i, ax in enumerate(self.axes):
            for t in ax.texts:
                if t.get_visible() and t.get_text().strip():
                    seen.append((i, t.get_text()))
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


def _render(series: SeriesResult) -> None:
    spec = ChartSpec(
        question_ref="q1", chart_type="combo", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles(title=True, data_labels=True, legend=True, n=True),
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    IMAGE_BUILDERS["combo"](RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format))


def _on(drawn, axes: int) -> list[str]:
    return [t for i, t in drawn if i == axes]


# ---- the engine says what each segment is ----------------------------------

def _computed() -> SeriesResult:
    primary = Variable("q", "Primary", "categorical",
                       (ValueLabel(1.0, "Low"), ValueLabel(2.0, "High")),
                       frozenset())
    sec = Variable("s", "tyoelamaindeksi", "scale", (), frozenset())
    model = QuestionModel(variables={"q": primary, "s": sec}, questions=[])
    q = Question(qid="q", kind="single", variables=("q",), text="Primary")
    df = pd.DataFrame({"q": [1, 1, 2, 2], "s": [3.0, 4.0, 7.0, 8.0]})
    spec = ChartSpec(
        question_ref="q", chart_type="combo", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s",
        elements=ElementToggles(), options={"combo_secondary": "s"})
    return engine.compute(q, spec, df, model)


def test_the_engine_marks_the_line_as_a_mean():
    r = _computed()
    bars, line = r.segments
    assert r.segment_statistics is not None, (
        "the result says nothing about what its segments are, so the renderer "
        "can only fall back to the series statistic — which is how the mean "
        "got a percent sign")
    assert r.segment_statistics[line] == "mean"
    assert r.segment_statistics[bars] == "pct"


def test_the_shortcut_that_stores_the_mean_in_pct_is_intact():
    """The fix must not move the value — the renderer reads `pct`."""
    r = _computed()
    assert r.cell("Low", "tyoelamaindeksi").pct == 3.5
    assert r.cell("High", "tyoelamaindeksi").pct == 7.5


# ---- what the reader sees --------------------------------------------------

def test_the_line_values_are_drawn_without_a_percent_sign(drawn):
    _render(_two_var_series(segment_statistics={_QUESTION: "pct", _INDEX: "mean"}))
    line = _on(drawn, 1)
    assert line, "the line drew no labels; the fixture is wrong, not the product"
    with_sign = [t for t in line if "%" in t]
    assert not with_sign, (
        f"the index is on a 1-8 scale and {len(with_sign)} of its {len(line)} "
        f"labels claim to be percentages: {with_sign}")


def test_the_bars_keep_their_percent_sign(drawn):
    """Widening must not strip the sign from the series that IS percentages."""
    _render(_two_var_series(segment_statistics={_QUESTION: "pct", _INDEX: "mean"}))
    bars = _on(drawn, 0)
    assert bars, "the bars drew no labels; the fixture is wrong, not the product"
    assert all("%" in t for t in bars), bars


def test_the_line_still_reads_as_the_index_it_is(drawn):
    """Not just "no %" — the number itself must survive, to one decimal."""
    _render(_two_var_series(segment_statistics={_QUESTION: "pct", _INDEX: "mean"}))
    line = " ".join(_on(drawn, 1))
    for v in ("3.6", "7.5"):
        assert v in line, f"{v} is missing from the line's labels: {line!r}"


def test_a_classifier_combo_is_percentages_on_both_axes(drawn):
    """The OTHER combo path: no secondary variable, two groups, both genuinely
    percentages. Nothing marks the segments, and both keep the sign."""
    _render(_two_var_series())
    every = _on(drawn, 0) + _on(drawn, 1)
    assert every
    assert all("%" in t for t in every), [t for t in every if "%" not in t]
