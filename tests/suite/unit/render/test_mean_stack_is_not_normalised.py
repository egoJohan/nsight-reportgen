"""A stack of MEANS is drawn at its true widths, not filled to 100 %.

`_stack_scaling` normalises when every bar it can judge partitions its base.
`_bar_is_measurable` deliberately refuses to judge a mean-statistic bar — cells
hold only `mean`, which says nothing about overlap — so that such a bar cannot
be the reason a chart abandons the 100 % reading. Sound for its own question,
but when NOTHING is measurable the list is empty and `all([])` is True: the
100 % reading was asserted with no evidence for it at all.

Työelämäindeksi crossed by pride drew eight identical full-width bars carrying
3.0, 3.6, 4.3, 4.9, 5.5, 6.1, 6.7 and 5.7. Every bar the same length, every
number different — groups that differ by more than a scale point look equal.
The printed number is true and the picture is false, which is the exact failure
the multi-response branch of that function already exists to prevent.

A mean is not a share of a whole; there is nothing for it to be a part of.
"""
from __future__ import annotations

import pandas as pd
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B

#: Mean index per pride level — the real shape, low group to high.
_MEANS = {1: 3.0, 2: 3.6, 3: 4.3, 4: 4.9, 5: 5.5, 6: 6.1, 7: 6.7}
_N = {1: 31, 2: 33, 3: 67, 4: 115, 5: 233, 6: 320, 7: 219}


def _ctx(chart_type="stacked_horizontal_bar"):
    idx = Variable(name="idx", label="Työelämäindeksi", measurement="scale",
                   missing_values=[], value_labels=[])
    pride = Variable(name="p", label="Ylpeys", measurement="nominal",
                     missing_values=[],
                     value_labels=[ValueLabel(value=float(c), label=f"{c}")
                                   for c in _MEANS])
    rows = [{"idx": float(_MEANS[c]), "p": float(c)}
            for c, n in _N.items() for _ in range(n)]
    model = QuestionModel(
        variables={"idx": idx, "p": pride},
        questions=[Question(qid="idx", text=idx.label, kind="single", variables=("idx",)),
                   Question(qid="p", text=pride.label, kind="single", variables=("p",))])
    spec = ChartSpec(question_ref="idx", chart_type=chart_type, statistic="mean",
                     classifying_var="p", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("idx"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    return RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                         slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                   width=Inches(11.6), height=Inches(4.0), name="s1"),
                         style=StyleSpec(), spec=spec, series=series,
                         fmt=spec.number_format), series


def _widths():
    ctx, series = _ctx()
    seen: dict = {}
    original = B.render_png

    def _spy(fig):
        ax = fig.axes[0]
        seen["widths"] = sorted({round(float(p.get_width()), 4)
                                 for p in ax.patches if p.get_width() > 0})
        seen["xmax"] = ax.get_xlim()[1]
        return original(fig)

    B.render_png = _spy
    try:
        B.build_image_bar_stacked(ctx)
    finally:
        B.render_png = original
    return seen


def test_the_bars_are_not_all_the_same_length():
    got = _widths()
    assert len(got["widths"]) > 1, (
        f"every bar drawn at {got['widths']} while the means differ")


def test_the_longest_bar_is_about_twice_the_shortest():
    """3.0 against 6.7 — the picture has to carry that difference."""
    got = _widths()
    lo, hi = got["widths"][0], got["widths"][-1]
    assert hi / lo > 1.8, f"{lo} .. {hi} does not show a 3.0 vs 6.7 spread"


def test_the_axis_reads_the_real_scale_not_a_hundred():
    got = _widths()
    assert got["xmax"] < 50, f"axis still runs to {got['xmax']} for a 1..7 index"


def test_a_percentage_stack_still_fills_to_a_hundred():
    """The control: the ordinary 100 %-stacked reading must not change."""
    cats = ["A", "B", "C"]
    var = Variable(name="q", label="Q", measurement="nominal", missing_values=[],
                   value_labels=[ValueLabel(value=float(i + 1), label=c)
                                 for i, c in enumerate(cats)])
    grp = Variable(name="g", label="G", measurement="nominal", missing_values=[],
                   value_labels=[ValueLabel(value=float(i + 1), label=f"G{i+1}")
                                 for i in range(2)])
    rows = [{"q": float((i % 3) + 1), "g": float((i % 2) + 1)} for i in range(300)]
    model = QuestionModel(
        variables={"q": var, "g": grp},
        questions=[Question(qid="q", text="Q", kind="single", variables=("q",)),
                   Question(qid="g", text="G", kind="single", variables=("g",))])
    spec = ChartSpec(question_ref="q", chart_type="stacked_horizontal_bar",
                     statistic="pct", classifying_var="g",
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    assert B._stack_scaling(series, list(series.segments), list(series.categories),
                            {c: [series.cell(c, s).value("pct") for s in series.segments]
                             for c in series.categories}) == (True, 100.0)
