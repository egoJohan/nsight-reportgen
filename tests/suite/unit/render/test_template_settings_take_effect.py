"""Every setting the layout editor offers changes the slide.

The editor is a promise: drag a box, pick a font, choose a colour, and the
sample redraws to show what you did. A control that stores a value and changes
nothing is worse than no control — it teaches an author that the feature is
broken and they stop trusting the ones that work.

An audit on 2026-09-08 found six settings that reached the style and never the
slide (title font, title colour, content colour, footer font, accent, and the
content font inside the CHART, as opposed to the text around it), plus four
that the request body dropped before they were ever stored. This file is the
whole table, so the next one that stops working says so.

Measured on the drawn PNG, not on the style object: "the value was stored" was
true of every one of those six.
"""
from __future__ import annotations

import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from reportbuilder.render.style_spec import apply_template_overrides, load_style_spec

_TEMPLATE = "src/reportbuilder/render/assets/nsight_default.pptx"
#: A colour no house palette contains, so finding it proves it was applied.
_PINK = "#B3005E"
_SERIF = "DejaVu Serif"


def _slide(overrides: dict):
    from reportbuilder.api.routes_templates import _sample_report
    from reportbuilder.export.pptx_build import build_presentation

    style = load_style_spec(_TEMPLATE)
    apply_template_overrides(style, overrides)
    report, model, df = _sample_report()
    return style, build_presentation(report, model, df, style=style).slides[0]


def _rendered(overrides: dict) -> bytes:
    """The whole slide as the editor's sample draws it."""
    from reportbuilder.render.image.fast_preview import compose_from_slide

    style, slide = _slide(overrides)
    return compose_from_slide(style, slide).tobytes()


def _chart_png(overrides: dict) -> bytes:
    """Just the chart image — the text INSIDE the chart, not the slide's."""
    _style, slide = _slide(overrides)
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert pics, "the sample slide is supposed to carry a chart"
    return pics[0].image.blob


@pytest.fixture(scope="module")
def plain() -> bytes:
    return _rendered({})


@pytest.fixture(scope="module")
def plain_chart() -> bytes:
    return _chart_png({})


# ── the areas the request body must accept ──────────────────────────────────
def test_the_request_body_keeps_every_area_the_editor_edits():
    """subtitle and footer were absent, so those panels saved nothing."""
    from reportbuilder.api.routes_templates import TemplateLayoutBody

    for area in ("title", "content", "subtitle", "footer"):
        assert area in TemplateLayoutBody.model_fields, f"{area} is dropped on save"


# ── title ───────────────────────────────────────────────────────────────────
def test_title_box_moves_the_headline(plain):
    assert _rendered({"title": {"x": 2.0, "y": 2.0}}) != plain


def test_title_size_is_applied(plain):
    assert _rendered({"title": {"size": 40}}) != plain


def test_title_font_is_applied(plain):
    assert _rendered({"title": {"font": _SERIF}}) != plain


def test_title_colour_is_applied(plain):
    assert _rendered({"title": {"colour": _PINK}}) != plain


# ── content: the chart's own typography ─────────────────────────────────────
def test_content_box_moves_the_chart(plain):
    assert _rendered({"content": {"x": 1.0, "y": 3.0, "w": 5.0, "h": 2.0}}) != plain


def test_content_size_reaches_the_chart(plain_chart):
    assert _chart_png({"content": {"size": 22}}) != plain_chart


def test_content_font_reaches_the_chart(plain_chart):
    """The reported fault: it styled the text AROUND the chart and not in it."""
    assert _chart_png({"content": {"font": _SERIF}}) != plain_chart


def test_content_colour_reaches_the_chart(plain_chart):
    assert _chart_png({"content": {"colour": _PINK}}) != plain_chart


# ── subtitle and footer ─────────────────────────────────────────────────────
@pytest.mark.parametrize("given", [{"font": _SERIF}, {"size": 22}, {"colour": _PINK}])
def test_subtitle_settings_are_applied(plain, given):
    assert _rendered({"subtitle": given}) != plain


@pytest.mark.parametrize("given", [{"font": _SERIF}, {"colour": _PINK}])
def test_footer_settings_are_applied(plain, given):
    assert _rendered({"footer": given}) != plain


# ── deck-wide ───────────────────────────────────────────────────────────────
def test_accent_is_applied(plain):
    assert _rendered({"accent": _PINK}) != plain


def test_background_is_applied(plain):
    assert _rendered({"background": "#102030"}) != plain


# ── the face a chart draws in ───────────────────────────────────────────────
def test_a_font_this_host_lacks_falls_back_to_the_house_face():
    """Not to matplotlib's DejaVu. An unavailable setting should land on a
    deliberate choice — the rule the removed admin-wide picker also applied."""
    from reportbuilder.render.house_style import _DEFAULT_CHART_FONT
    from reportbuilder.render.image._mpl import chart_text_font

    class _Style:
        body_font = "Definitely Not Installed"

    assert chart_text_font(_Style()) == _DEFAULT_CHART_FONT


def test_the_chart_face_is_per_figure_not_process_wide():
    """Two templates render concurrently on FastAPI's threadpool. The family is
    stashed on the FIGURE; putting it in matplotlib's global rcParams would let
    one report's font land on another's chart."""
    import matplotlib

    from reportbuilder.render.image._mpl import _remember_font

    class _Fig:
        pass

    class _Ctx:
        class style:
            body_font = "DejaVu Serif"

    before = matplotlib.rcParams["font.family"]
    fig = _Fig()
    _remember_font(fig, _Ctx())
    assert fig._nsight_chart_font == "DejaVu Serif"
    assert matplotlib.rcParams["font.family"] == before, "rcParams was mutated"


def test_the_admin_wide_chart_font_setting_is_gone():
    """It could not tell one customer's deck from another's. The face is a
    property of the template now, so the global setting had to go with it."""
    from reportbuilder.api import routes_settings

    assert not hasattr(routes_settings, "apply_chart_font")
    assert not hasattr(routes_settings, "chart_font_for")
    paths = {r.path for r in routes_settings.settings_router.routes}
    assert "/settings/chart-font" not in paths
