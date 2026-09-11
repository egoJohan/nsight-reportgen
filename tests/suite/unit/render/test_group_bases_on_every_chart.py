"""Every chart that compares groups says how many people each group is.

"Kohderyhmäkohtaiset n-luvut tulevat nyt esille vertical ja horizontal bar
kaaviotyypeissä, mutta ei stacked kaavioissa. Hienoa jos voisi varmistaa, että
kohderyhmien n-luvut tulevat esille kaikissa kaaviotyypeissä."

The grouped bars got it through their legend, because there the series ARE the
groups. A stacked bar's legend is the answer scale, and was rightly left alone —
but its BARS are the groups, and their names said nothing. The same was true of
every layout that names a group somewhere other than a series legend: the
panels of a cross-tab, a small-multiples row sharing one legend.

So this is checked on the figure as saved, for every chart type that splits by
a group and every cross-tab layout: each group's name is on it, and so is its
own base. The groups are deliberately of different sizes, so an n printed
against the wrong group does not pass by coincidence. (Johan, 2026-09-11)
"""
from __future__ import annotations

import re

import matplotlib
matplotlib.use("Agg")
import pandas as pd
import pytest
from matplotlib.figure import Figure
from matplotlib.text import Text
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.engine import compute

_SCALE = ["Erittäin huono", "Melko huono", "Ei hyvä eikä huono", "Melko hyvä", "Erittäin hyvä"]
#: (gender, age) -> people. Every group, and every crossing, a different size.
_CELLS = {("Naiset", "18-39"): 240, ("Naiset", "40-64"): 261,
          ("Miehet", "18-39"): 255, ("Miehet", "40-64"): 247}


def _var(name, label, values):
    return Variable(name=name, label=label, measurement="nominal", missing_values=[],
                    value_labels=[ValueLabel(value=float(i + 1), label=v)
                                  for i, v in enumerate(values)])


def _model():
    genders, ages = ["Naiset", "Miehet"], ["18-39", "40-64"]
    rows, i = [], 0
    for (g, a), n in _CELLS.items():
        for _ in range(n):
            rows.append({"q": float(i % 5 + 1), "sp": float(genders.index(g) + 1),
                         "ika": float(ages.index(a) + 1)})
            i += 1
    model = QuestionModel(
        variables={"q": _var("q", "Mielipide", _SCALE),
                   "sp": _var("sp", "Sukupuoli", genders),
                   "ika": _var("ika", "Ikä", ages)},
        questions=[Question(qid="q", text="Mielipide", kind="single", variables=("q",)),
                   Question(qid="sp", text="Sukupuoli", kind="single", variables=("sp",)),
                   Question(qid="ika", text="Ikä", kind="single", variables=("ika",))])
    return model, pd.DataFrame(rows)


def _render(chart_type, *, var="sp", var_2=None, layout=None, question="q",
            model_df=None):
    model, df = model_df or _model()
    spec = ChartSpec(question_ref=question, chart_type=chart_type, statistic="pct",
                     classifying_var=var, classifying_var_2=var_2,
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(),
                     options={"xtab_layout": layout} if layout else {})
    series = compute(model.question(question), spec, df, model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(12.3), height=Inches(4.6), name="s1"),
                        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS[chart_type](ctx)
    return series


@pytest.fixture
def printed(monkeypatch):
    """Every piece of text on each figure as it is saved."""
    seen: list[list[str]] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        seen.append([t.get_text() for t in self.findobj(Text)
                     if t.get_visible() and t.get_text().strip()])
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


def _groups(series):
    """(name as drawn, base) for every group the chart compares — the combos of a
    cross-tab are named by their second part beside the first."""
    out = []
    for seg in series.segments:
        if "Total" in seg:
            continue
        out.append((seg.split(" · ", 1)[-1], series.base_n[seg]))
    return out


def _missing(texts, groups):
    """Groups whose name, or whose own base, is nowhere on the figure."""
    n_pat = lambda n: re.compile(rf"\bn\s*=\s*{n}\b")
    return [(name, n) for name, n in groups
            if not any(name in t for t in texts)
            or not any(n_pat(n).search(t) for t in texts)]


# ── one classifier, every chart type that splits by it ──────────────────────

@pytest.mark.parametrize("chart_type", [
    "vertical_bar", "horizontal_bar", "stacked_vertical_bar", "stacked_horizontal_bar",
    "line", "radar", "pie", "doughnut", "funnel",
])
def test_each_group_s_base_is_on_the_chart(printed, chart_type):
    series = _render(chart_type)
    texts = [t for fig in printed for t in fig]
    assert _missing(texts, _groups(series)) == [], texts


@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar"])
def test_a_stacked_bar_says_it_on_the_bar_s_own_name(printed, chart_type):
    """Not in the legend — that is the answer scale, and "Melko hyvä" has no base."""
    series = _render(chart_type)
    texts = printed[-1]
    for name, n in _groups(series):
        assert any(name in t and f"n={n}" in t for t in texts), (name, n, texts)
    assert not any(lvl in t and "n=" in t for lvl in _SCALE for t in texts), texts


# ── two classifiers, in every layout a cross-tab can take ───────────────────

@pytest.mark.parametrize("layout", ["grouped", "separate", "small_multiples"])
@pytest.mark.parametrize("chart_type", [
    "vertical_bar", "horizontal_bar", "stacked_horizontal_bar", "stacked_vertical_bar",
])
def test_a_cross_tab_names_every_group_s_base_in_every_layout(printed, chart_type, layout):
    series = _render(chart_type, var="sp", var_2="ika", layout=layout)
    texts = [t for fig in printed for t in fig]
    assert _missing(texts, _groups(series)) == [], texts


def test_small_multiples_give_each_panel_its_own_bases(printed):
    """One legend under a row of panels cannot say "Naiset" is 240 people in one
    panel and 261 in the next. Each panel says its own."""
    series = _render("vertical_bar", var="sp", var_2="ika", layout="small_multiples")
    texts = printed[-1]
    for seg in series.segments:
        if "Total" in seg:
            continue
        assert any(f"n={series.base_n[seg]}" in t for t in texts), (seg, texts)


# ── and where there is no group, no n is made up ─────────────────────────────

def test_a_chart_split_by_nothing_invents_no_group(printed):
    """One bar of everyone: the slide's own N already says how many."""
    _render("stacked_horizontal_bar", var=None)
    assert not any("(n=" in t for t in printed[-1]), printed[-1]


def test_a_battery_s_statements_are_not_groups(printed):
    """A battery's bars are its statements. Each has its own count of answers,
    but it is not a group of people the slide compares — printing it thirteen
    times over is noise, and not what was asked for."""
    stmts = ["maku", "rapeus", "hinta"]
    scale = ["Ei lainkaan", "Vähän", "Paljon"]
    variables = {f"s{i}": _var(f"s{i}", f"{s}:Kuinka tärkeää?", scale)
                 for i, s in enumerate(stmts)}
    rows = [{f"s{i}": float((r + i) % 3 + 1) for i in range(3)} for r in range(300)]
    model = QuestionModel(
        variables=variables,
        questions=[Question(qid="bat", text="Kuinka tärkeää?", kind="battery",
                            variables=tuple(variables))])
    _render("stacked_horizontal_bar", var=None, question="bat",
            model_df=(model, pd.DataFrame(rows)))
    assert not any("(n=" in t for t in printed[-1]), printed[-1]
