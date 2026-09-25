"""A key-themes slide reports its N, like every chart slide.

"Key themes kysymystyypissä ei näy n-lukua." (2026-09-25) Themes are bullets
written from an open question's answers; they were drawn by the special-slide
renderer, which has no methodology line. The N is the respondents who wrote an
answer — the same count a word cloud of the same question reports — in the same
place and look, obeying the same "N" toggle and footer note.
"""
from __future__ import annotations

import os
import tempfile

import pandas as pd
import pytest
from pptx import Presentation

from reportbuilder.ingest.sav_reader import Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import report_from_json
from reportbuilder.stats.engine import text_respondents

pytestmark = pytest.mark.unit

_ANSWERS = ["Hyvä palvelu", "", "Hinta", None, "  ", "Nopea toimitus", "En osaa sanoa"]


def _model():
    model = QuestionModel(
        variables={"avoin": Variable(name="avoin", label="Miksi?", measurement="nominal",
                                     missing_values=[], value_labels=[])},
        questions=[Question(qid="avoin", text="Miksi valitsit meidät?", kind="text",
                            variables=("avoin",))])
    return model, pd.DataFrame({"avoin": _ANSWERS})


def _slide_texts(chart: dict) -> list[str]:
    from reportbuilder.export.pptx_build import build_pptx

    model, df = _model()
    base = {"question_ref": "avoin", "chart_type": "themes", "statistic": "pct",
            "number_format": {}, "template_slot": "s1", "elements": {},
            "sort": {"basis": "data_order"},
            "options": {"bullets": ["- Palvelu kiittää", "- Hinta mietityttää"]}}
    report = report_from_json({"name": "r", "render_mode": "image", "template_ref": "",
                               "charts": [{**base, **chart}]})
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "s.pptx")
        build_pptx(report, model, df, path)
        prs = Presentation(path)
    return [sh.text_frame.text.strip() for sh in prs.slides[0].shapes
            if sh.has_text_frame and sh.text_frame.text.strip()]


def test_the_count_is_the_respondents_who_wrote_something():
    model, df = _model()
    assert text_respondents(model.question("avoin"), df) == 4


def test_a_themes_slide_shows_its_n():
    texts = _slide_texts({})
    assert "N = 4" in texts, texts


def test_the_n_toggle_hides_it():
    assert not any(t.startswith("N =") for t in _slide_texts({"elements": {"n": False}}))


def test_an_authors_footer_note_replaces_it():
    texts = _slide_texts({"footer_note": "Avoimet vastaukset, n = {n}"})
    assert "Avoimet vastaukset, n = 4" in texts, texts


def test_an_overview_slide_has_no_n():
    texts = _slide_texts({"question_ref": "special_overview",
                          "chart_type": "special_overview"})
    assert not any(t.startswith("N =") for t in texts), texts
