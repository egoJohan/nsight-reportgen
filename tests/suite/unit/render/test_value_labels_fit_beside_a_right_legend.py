"""Column numbers are fitted to the bars as they are finally drawn.

Six groups move the legend to the right of the plot, which keeps the plot to
72 % of its width. The numbers were fitted before that happened, against bars a
quarter wider than the ones drawn, so neighbouring groups' "58 %" and "59 %"
were printed through each other. (visual QA, 2026-09-19 — Holiday Club, by age)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("En tunne omistajaetuja ollenkaan tai erittäin vähän",
         "Tunnen omistajaedut jokseenkin, mutta en muista kaikkia yksityiskohtia",
         "Tunnen omistajaedut hyvin")
_AGES = ("20-24 vuotta", "25-34 vuotta", "35-44 vuotta", "45-54 vuotta",
         "55-64 vuotta", "65-70 vuotta")
_VALUES = ((17, 18, 18, 25, 21, 25), (58, 59, 57, 54, 55, 51), (25, 23, 25, 21, 24, 24))


def test_neighbouring_groups_numbers_do_not_touch():
    cells = {(c, g): Cell(pct=float(_VALUES[i][j]))
             for i, c in enumerate(_CATS) for j, g in enumerate(_AGES)}
    series = SeriesResult(categories=_CATS, segments=_AGES, cells=cells,
                          base_n={g: 200 for g in _AGES}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type="vertical_bar", statistic="pct",
                     classifying_var="ika", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    boxes: list = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        boxes.extend((t.get_text(), t.get_window_extent(r)) for ax in self.axes
                     for t in ax.texts if t.get_gid() == VALUE_GID and t.get_visible())
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["vertical_bar"](ctx)
    finally:
        _f.Figure.savefig = original
    assert boxes, "no number was printed"
    for i, (ta, a) in enumerate(boxes):
        for tb, b in boxes[i + 1:]:
            assert not (min(a.x1, b.x1) - max(a.x0, b.x0) > 0.5
                        and min(a.y1, b.y1) - max(a.y0, b.y0) > 0.5), f"{ta!r} over {tb!r}"
