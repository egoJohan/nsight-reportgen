"""A column keeps its percentage as long as the number fits on it.

"Vertical bar chartissa ei näy prosenttilukuja kun tarkastelee lukuja eri
taustaryhmissä" — a chart split by four or more background groups printed no
numbers at all. The rule was a bare count: past `_MAX_LABELED_SEGMENTS_V = 4`
segments every label was dropped, on the assumption that the columns had become
too narrow to carry one.

A count cannot know that. The same five segments are roomy on a wide slide with
three categories and hopeless on a narrow one with twelve, and the reader loses
the numbers in both. So it is measured instead — the same move the pie and
stacked callouts made: draw the number where it fits, turn it on its side when
the column is narrow but tall enough, and drop it only when neither works.
(Johan, 2026-09-10)
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B


def _labels(n_groups: int, n_cats: int = 3, slot_in: float = 9.0) -> dict:
    """Render a grouped column chart and report the value labels it drew."""
    q = Variable(name="q", label="Q", measurement="nominal", missing_values=[],
                 value_labels=[ValueLabel(value=float(i + 1), label=f"Vaihtoehto {i+1}")
                               for i in range(n_cats)])
    g = Variable(name="g", label="Ryhmä", measurement="nominal", missing_values=[],
                 value_labels=[ValueLabel(value=float(i + 1), label=f"Ryhmä {i+1}")
                               for i in range(n_groups)])
    rows = [{"q": float(i % n_cats + 1), "g": float(i % n_groups + 1)}
            for i in range(600)]
    model = QuestionModel(
        variables={"q": q, "g": g},
        questions=[Question(qid="q", text="Q", kind="single", variables=("q",)),
                   Question(qid="g", text="Ryhmä", kind="single", variables=("g",))])
    spec = ChartSpec(question_ref="q", chart_type="vertical_bar", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1),
                                  width=Inches(slot_in), height=Inches(4.2), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    caught: dict = {}
    original = B.render_png

    def _spy(fig):
        ax = fig.axes[0]
        vals = [t for t in ax.texts if "%" in t.get_text()]
        caught["n"] = len(vals)
        caught["rotations"] = sorted({round(t.get_rotation()) for t in vals})
        caught["segments"] = len(series.segments)
        return original(fig)

    B.render_png = _spy
    try:
        B.build_image_column(ctx)
    finally:
        B.render_png = original
    return caught


def test_two_groups_still_read_flat():
    """The common case must look exactly as it always has."""
    got = _labels(2)
    assert got["n"] > 0
    assert got["rotations"] == [0]


def test_five_groups_keep_their_numbers():
    """The reported case: four background groups plus Total drew nothing."""
    got = _labels(5)
    assert got["segments"] >= 5
    assert got["n"] > 0, "no percentages were drawn at all"


def test_a_crowded_chart_turns_the_numbers_on_their_side():
    got = _labels(8, n_cats=4)
    assert got["n"] > 0
    assert 90 in got["rotations"], got["rotations"]


def test_every_column_that_gets_a_number_gets_it():
    """Never a partial row — some columns labelled and others not reads as a
    rendering fault rather than a choice."""
    got = _labels(5, n_cats=3)
    assert got["n"] % 3 == 0, got["n"]


def test_hopeless_crowding_still_drops_them():
    """Twenty segments on a narrow slide cannot carry a number either way, and
    a wall of overlapping text is worse than the grid and the legend."""
    got = _labels(20, n_cats=6, slot_in=4.0)
    assert got["n"] == 0


def _label_boxes(n_groups: int, n_cats: int = 4, slot_in: float = 9.0):
    """Where the value labels actually landed, in device pixels."""
    import pandas as pd
    from pptx import Presentation
    from pptx.util import Inches

    from reportbuilder.ingest.sav_reader import ValueLabel, Variable
    from reportbuilder.model.question import Question, QuestionModel
    from reportbuilder.render.base import RenderContext, Slot, StyleSpec

    q = Variable(name="q", label="Q", measurement="nominal", missing_values=[],
                 value_labels=[ValueLabel(value=float(i + 1), label=f"V{i+1}")
                               for i in range(n_cats)])
    g = Variable(name="g", label="R", measurement="nominal", missing_values=[],
                 value_labels=[ValueLabel(value=float(i + 1), label=f"R{i+1}")
                               for i in range(n_groups)])
    rows = [{"q": float(i % n_cats + 1), "g": float(i % n_groups + 1)}
            for i in range(900)]
    model = QuestionModel(
        variables={"q": q, "g": g},
        questions=[Question(qid="q", text="Q", kind="single", variables=("q",)),
                   Question(qid="g", text="R", kind="single", variables=("g",))])
    spec = ChartSpec(question_ref="q", chart_type="vertical_bar", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1),
                                  width=Inches(slot_in), height=Inches(4.2), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    boxes: dict = {}
    original = B.render_png

    def _spy(fig):
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        ax = fig.axes[0]
        boxes["b"] = sorted(
            (t.get_window_extent(r).x0, t.get_window_extent(r).x1,
             round(t.get_window_extent(r).y0))
            for t in ax.texts if "%" in t.get_text())
        return original(fig)

    B.render_png = _spy
    try:
        B.build_image_column(ctx)
    finally:
        B.render_png = original
    return boxes["b"]


@pytest.mark.parametrize("groups", [2, 3, 5, 6, 8])
def test_value_labels_never_overlap_each_other(groups):
    """The measurement exists to prevent exactly this, and asserting that
    labels were DRAWN does not check it — the first version of this fix printed
    "30 %30 %" and passed its own tests."""
    boxes = _label_boxes(groups)
    on_a_row: dict = {}
    for x0, x1, y in boxes:
        on_a_row.setdefault(y, []).append((x0, x1))
    for y, spans in on_a_row.items():
        spans.sort()
        for (a0, a1), (b0, b1) in zip(spans, spans[1:]):
            assert b0 >= a1 - 0.5, (
                f"{groups} groups: labels overlap at y={y} "
                f"({a0:.0f}-{a1:.0f} then {b0:.0f})")


def test_a_number_on_its_side_stands_on_its_own_bar():
    """Turned for a crowded chart, a number still sits ABOVE its bar and over
    its middle. Aligned before turning (`rotation_mode="anchor"`), the middle
    of the upright number sat on the bar top and it was pushed to the left:
    half of "43 %" inside a dark bar, reported from staging as "pylväiden
    numeroiden formaatissa on jotain outoa". (2026-09-19)"""
    import matplotlib.patches as mpatches

    found: dict = {}
    original = B.render_png

    def _spy(fig):
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        ax = fig.axes[0]
        bars = [p.get_window_extent(r) for p in ax.patches
                if isinstance(p, mpatches.Rectangle)]
        texts = [t for t in ax.texts if "%" in t.get_text()]
        found["pairs"] = []
        for t in texts:
            box = t.get_window_extent(r)
            cx = (box.x0 + box.x1) / 2
            bar = min(bars, key=lambda b: abs((b.x0 + b.x1) / 2 - cx))
            found["pairs"].append((t.get_rotation(), box, bar, t.get_text()))
        return original(fig)

    B.render_png = _spy
    try:
        _labels(8, n_cats=4)
    finally:
        B.render_png = original

    turned = [p for p in found["pairs"] if p[0] == 90]
    assert turned, "the crowded chart should turn its numbers"
    for _rot, box, bar, text in turned:
        assert box.y0 >= bar.y1 - 1, f"{text!r} starts inside its bar"
        assert bar.x0 - 1 <= (box.x0 + box.x1) / 2 <= bar.x1 + 1, (
            f"{text!r} is not over its bar")
