"""A template is parsed once per file version, and every caller gets its own copy.

Two parses were repeated on every preview (perf, 2026-09-19):

* the deck builder and the preview ground opened the customer's whole file —
  Attendo's is 56 slides — and deleted its slides again: 183-660 ms;
* a template whose author had chosen a layout was harvested again on every
  request, cached previews included: 0.8-1.9 s.

Both are now cached on the file's (path, size, mtime), like
`template_cache.resolve`. What must not change: a replaced file is read again,
and no two callers share an object one of them could alter.
"""
from __future__ import annotations

import os
import pathlib
import shutil

import pytest

_DECK = pathlib.Path("src/reportbuilder/render/assets/nsight_default.pptx")


@pytest.fixture
def template(tmp_path) -> pathlib.Path:
    path = tmp_path / "pohja.pptx"
    shutil.copy(_DECK, path)
    return path


def test_the_stripped_template_has_no_slides_and_keeps_its_layouts(template):
    from pptx import Presentation

    from reportbuilder.render.deck import open_stripped_template

    prs = open_stripped_template(str(template))
    assert len(prs.slides) == 0
    assert len(prs.slide_layouts) == len(Presentation(str(template)).slide_layouts)


def test_each_caller_gets_its_own_presentation(template):
    from reportbuilder.render.deck import open_stripped_template

    a = open_stripped_template(str(template))
    b = open_stripped_template(str(template))
    assert a is not b
    a.slides.add_slide(a.slide_layouts[0])
    assert len(b.slides) == 0, "one caller's slide appeared in another's deck"


def test_the_file_is_stripped_once_per_version(template):
    from reportbuilder.render import deck

    deck._stripped_bytes.cache_clear()
    deck.open_stripped_template(str(template))
    deck.open_stripped_template(str(template))
    assert deck._stripped_bytes.cache_info().misses == 1
    # A replaced file — new bytes, new modification time — is read again.
    with open(template, "ab") as fh:
        fh.write(b"\0")
    os.utime(template, ns=(template.stat().st_atime_ns, template.stat().st_mtime_ns + 10**9))
    try:
        deck.open_stripped_template(str(template))
    except Exception:  # noqa: BLE001 — the padded file need not parse; it must be READ
        pass
    assert deck._stripped_bytes.cache_info().misses == 2


def test_a_chosen_layout_is_harvested_once_and_copied(template, monkeypatch):
    from reportbuilder.render import style_spec, template_cache

    template_cache._with_layout.cache_clear()
    calls = []
    real = style_spec.load_style_spec

    def counting(path, force_layout=None):
        calls.append(force_layout)
        return real(path, force_layout=force_layout)

    monkeypatch.setattr(style_spec, "load_style_spec", counting)
    a = template_cache.style_with_overrides(str(template), {"layout_index": 0})
    b = template_cache.style_with_overrides(str(template), {"layout_index": 0})
    assert calls == [0], calls
    assert a is not b
    a.brand_palette.append("123456")
    assert "123456" not in b.brand_palette, "a caller's change leaked"
