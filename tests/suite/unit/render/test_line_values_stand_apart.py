"""A line chart split into groups prints each group's numbers where they can be read.

Every point's number sat a fixed 7pt above it, so wherever two lines ran close
their numbers were printed through each other — "50 %" over "49 %" over "46 %".
Every classified line chart in the visual QA matrix had it, up to 105
collisions on one slide. Nothing may be hidden to make room (Johan,
2026-09-13), so the numbers are moved apart instead. (visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("Erittäin huono", "Huono", "Hyvä", "Erittäin hyvä", "En osaa sanoa")
#: Four groups whose lines run close together, and share points at 0 %.
_GROUPS = {"Mieheksi": (5.0, 26.0, 46.0, 12.0, 12.0),
           "Naiseksi": (4.0, 25.0, 49.0, 9.0, 12.0),
           "Muuksi": (0.0, 25.0, 50.0, 0.0, 25.0),
           "EOS": (0.0, 100.0, 0.0, 0.0, 0.0)}


def _values():
    cells = {(c, g): Cell(pct=v[i]) for g, v in _GROUPS.items() for i, c in enumerate(_CATS)}
    series = SeriesResult(categories=_CATS, segments=tuple(_GROUPS), cells=cells,
                          base_n={g: 100 for g in _GROUPS}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type="line", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    got: dict = {}

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        texts = [t for ax in self.axes for t in ax.texts
                 if t.get_gid() == VALUE_GID and t.get_visible()]
        got["boxes"] = [(t.get_text(), t.get_window_extent(r)) for t in texts]
        got["points"] = {(round(float(t.xy[0])), t.get_text()) for t in texts}
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["line"](ctx)
    finally:
        _f.Figure.savefig = original
    return got


def test_no_number_is_printed_over_another():
    boxes = _values()["boxes"]
    for i, (ta, a) in enumerate(boxes):
        for tb, b in boxes[i + 1:]:
            overlap = (min(a.x1, b.x1) - max(a.x0, b.x0) > 0.5
                       and min(a.y1, b.y1) - max(a.y0, b.y0) > 0.5)
            assert not overlap, f"{ta!r} is printed over {tb!r}"


def test_every_point_keeps_its_number():
    """One number per point: the same value at the same marker is printed once,
    every other value is printed."""
    points = _values()["points"]
    expected = {(i, f"{v:.0f} %") for vals in _GROUPS.values() for i, v in enumerate(vals)}
    assert points == expected
