"""Category labels are never printed on top of each other — on any chart.

"Etenkin stacked horizontal kaaviossa olisi tärkeä varmistaa ettei tekstit mene
päällekkäin." The Taffel battery: thirteen statements on a 100%-stacked bar,
each row name wrapped at a fixed 30 characters and set at a fixed 11.5pt,
whatever height a row actually got. Thirteen rows in a slide-sized chart leave
about 18pt a row; two lines of 11.5pt need 28. Six pairs overlapped, at every
slide size tried.

Nothing about that is particular to a stacked bar. Every bar, column and line
builder labels its categories at a fixed size and a fixed wrap, so the guard is
checked through each of them rather than through the one that was reported.

Reproduced from the customer's own counts (staging, read-only). (Johan, 2026-09-11)
"""
from __future__ import annotations

import math

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

#: var81-var93, "Kuinka todennäköisesti ostaisit Taffelin tuotteita useammin,
#: jos Taffel…", N=1046: counts per level, least to most likely.
_TAFFEL = [
    ("parantaisi nykyisten tuotteiden makua", (129, 377, 401, 139)),
    ("parantaisi tuotteiden rapeutta", (117, 373, 419, 137)),
    ("toisi oman suosikkimakuni kauppaan, jossa asioin", (115, 302, 425, 204)),
    ("parantaisi tuotteiden laatua", (87, 335, 451, 173)),
    ("uudistaisi pakkausten ilmettä", (207, 465, 280, 94)),
    ("tarjoaisi enemmän kampanjoita ja tarjouksia", (80, 213, 439, 314)),
    ("lisäisi tuotteiden näkyvyyttä kaupoissa", (161, 444, 327, 114)),
    ("korostaisi enemmän tuotteiden kotimaisuutta", (153, 384, 372, 137)),
    ("tarjoaisi nykyistä pienempiä pakkauskokoja", (241, 368, 292, 145)),
    ("tarjoaisi nykyistä suurempia pakkauskokoja", (192, 390, 271, 193)),
    ("tarjoaisi tuotteet edullisempaan hintaan kuin Estrella", (79, 206, 409, 352)),
    ("laskisi kilohintaa", (59, 197, 408, 382)),
    ("toisi uusia makuvaihtoehtoja valikoimaan", (109, 294, 412, 231)),
]
_LEVELS = ("Erittäin epätodennäköisesti", "Melko epätodennäköisesti",
           "Melko todennäköisesti", "Erittäin todennäköisesti")


def _taffel_battery() -> SeriesResult:
    """The battery as the stacked engine hands it over: statements are the bars
    (segments), the scale levels the stack (categories)."""
    cells, base, top2 = {}, {}, []
    for stmt, counts in _TAFFEL:
        n = sum(counts)
        base[stmt] = n
        for lvl, c in zip(_LEVELS, counts):
            cells[(lvl, stmt)] = Cell(pct=100.0 * c / n, count=float(c), mean=None)
        top2.append(100.0 * (counts[2] + counts[3]) / n)
    stmts = tuple(s for s, _c in _TAFFEL)
    return SeriesResult(categories=_LEVELS, segments=stmts, cells=cells, base_n=base,
                        statistic="pct", row_summaries=tuple(top2), row_summary_keys=stmts,
                        segments_are_groups=False)   # as the engine says for a battery


def _one_series(labels) -> SeriesResult:
    """One value per category — the plain bar/column/line shape."""
    cells = {(c, "Total"): Cell(pct=20.0 + (i * 7) % 60, count=10.0, mean=None)
             for i, c in enumerate(labels)}
    return SeriesResult(categories=tuple(labels), segments=("Total",), cells=cells,
                        base_n={"Total": 1000}, statistic="pct")


def _render(chart_type, series, w_in, h_in, **spec_kw) -> bytes:
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(), **spec_kw)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(slide=slide,
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(w_in), height=Inches(h_in), name="s1"),
                        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS[chart_type](ctx)
    return [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE][0].image.blob


# ── the oracle: what is on the figure when it is saved ───────────────────────

def _oriented(t, r):
    """(centre, width, height, angle) of a label as the rectangle it really is —
    a rotated label's screen box is much bigger than its text, and two of them
    side by side would read as colliding when they are not."""
    deg = t.get_rotation()
    box = t.get_window_extent(r)
    t.set_rotation(0)
    flat = t.get_window_extent(r)
    t.set_rotation(deg)
    return ((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2), flat.width, flat.height, math.radians(deg)


def _collisions(fig):
    r = fig.canvas.get_renderer()
    hits = []
    for ax in fig.axes:
        for labels in (ax.get_xticklabels(), ax.get_yticklabels()):
            shown = [t for t in labels if t.get_visible() and t.get_text().strip()]
            rects = [(t.get_text(), *_oriented(t, r)) for t in shown]
            for i, (ta, ca, wa, ha, ang) in enumerate(rects):
                for tb, cb, wb, hb, _ang in rects[i + 1:]:
                    dx, dy = cb[0] - ca[0], cb[1] - ca[1]
                    # into the labels' own frame, where they are axis-aligned
                    u = dx * math.cos(ang) + dy * math.sin(ang)
                    v = -dx * math.sin(ang) + dy * math.cos(ang)
                    if abs(u) < (wa + wb) / 2 - 0.5 and abs(v) < (ha + hb) / 2 - 0.5:
                        hits.append((ta.replace("\n", " / "), tb.replace("\n", " / ")))
    return hits


@pytest.fixture
def saved(monkeypatch):
    """Every figure as it goes to disk: its colliding label pairs and its labels."""
    seen: list[dict] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        labels = [t for ax in self.axes
                  for t in (*ax.get_xticklabels(), *ax.get_yticklabels())
                  if t.get_visible() and t.get_text().strip()]
        # Space between neighbouring ROW names, in units of their own type size.
        r = self.canvas.get_renderer()
        gaps = []
        for ax in self.axes:
            rows = sorted((t for t in ax.get_yticklabels()
                           if t.get_visible() and t.get_text().strip()),
                          key=lambda t: t.get_window_extent(r).y0)
            for lo, hi in zip(rows, rows[1:]):
                gap_pt = (hi.get_window_extent(r).y0 - lo.get_window_extent(r).y1) * 72.0 / self.dpi
                gaps.append(gap_pt / lo.get_fontsize())
        # Names a legend is printed over.
        legends = [ax.get_legend() for ax in self.axes if ax.get_legend()] + list(self.legends)
        covered = []
        for leg in legends:
            lb = leg.get_window_extent(r)
            for t in labels:
                (cx, cy), w, h, ang = _oriented(t, r)
                # the name's corners, tested against the legend's box
                pts = [(cx + dx * math.cos(ang) - dy * math.sin(ang),
                        cy + dx * math.sin(ang) + dy * math.cos(ang))
                       for dx in (-w / 2, 0.0, w / 2) for dy in (-h / 2, 0.0, h / 2)]
                if any(lb.x0 + 0.5 < x < lb.x1 - 0.5 and lb.y0 + 0.5 < y < lb.y1 - 0.5
                       for x, y in pts):
                    covered.append(t.get_text().replace("\n", " / "))
        seen.append({"hits": _collisions(self),
                     "texts": [t.get_text() for t in labels],
                     "sizes": [t.get_fontsize() for t in labels],
                     "row_gaps": gaps,
                     "under_legend": covered})
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


# ── the reported chart ───────────────────────────────────────────────────────

@pytest.mark.parametrize("w_in,h_in", [(12.3, 4.4), (12.3, 5.0), (9.0, 4.5), (12.3, 3.6)])
def test_the_customer_s_battery_prints_every_statement_clear(saved, w_in, h_in):
    _render("stacked_horizontal_bar", _taffel_battery(), w_in, h_in,
            row_summary_fn="top2_sum")
    assert saved[-1]["hits"] == [], saved[-1]["hits"][:4]


@pytest.mark.parametrize("w_in,h_in", [(12.3, 4.4), (9.0, 4.5)])
def test_a_two_line_name_cannot_be_read_as_its_neighbour_s(saved, w_in, h_in):
    """Not touching is not enough. With the boxes flush, "…hintaan kuin
    Estrella" sat on "laskisi kilohintaa" and read as one name: the space
    BETWEEN two names has to be plainly more than the leading WITHIN one."""
    _render("stacked_horizontal_bar", _taffel_battery(), w_in, h_in,
            row_summary_fn="top2_sum")
    assert min(saved[-1]["row_gaps"]) >= 0.35, sorted(saved[-1]["row_gaps"])[:3]


def test_it_says_every_statement_in_full(saved):
    """Clear by being cut short would be no fix: these fit whole."""
    _render("stacked_horizontal_bar", _taffel_battery(), 12.3, 4.4,
            row_summary_fn="top2_sum")
    assert not any("…" in t for t in saved[-1]["texts"])


def test_it_stays_readable(saved):
    """Clear by being microscopic would be no fix either."""
    _render("stacked_horizontal_bar", _taffel_battery(), 12.3, 4.4,
            row_summary_fn="top2_sum")
    assert min(saved[-1]["sizes"]) >= 7.5


# ── the same guarantee, through every builder that names categories ─────────

_LONG = [s for s, _c in _TAFFEL] + [
    "valikoimassa olisi enemmän gluteenittomia vaihtoehtoja",
    "tuotteiden alkuperä olisi paremmin merkitty pakkaukseen",
    "hinta-laatusuhde olisi nykyistä parempi",
    "pakkaus olisi helpompi avata ja sulkea uudelleen",
]
_MEDIUM = ["Helsinki-Uusimaa", "Varsinais-Suomi", "Pohjois-Pohjanmaa", "Kanta-Häme",
           "Etelä-Pohjanmaa", "Pohjois-Savo", "Keski-Suomi", "Satakunta",
           "Pirkanmaa", "Päijät-Häme", "Kymenlaakso", "Etelä-Karjala",
           "Pohjois-Karjala", "Keski-Pohjanmaa"]


@pytest.mark.parametrize("chart_type,series,w_in,h_in", [
    ("stacked_vertical_bar", _taffel_battery(), 12.3, 4.4),
    ("horizontal_bar", _one_series(_LONG), 12.3, 4.4),
    ("vertical_bar", _one_series(_MEDIUM), 9.0, 4.5),
    ("line", _one_series(_MEDIUM), 9.0, 4.5),
], ids=["stacked-column", "bar", "column", "line"])
def test_no_chart_prints_its_categories_over_each_other(saved, chart_type, series, w_in, h_in):
    _render(chart_type, series, w_in, h_in)
    assert saved[-1]["hits"] == [], saved[-1]["hits"][:4]


@pytest.mark.parametrize("chart_type,series,w_in,h_in", [
    ("stacked_vertical_bar", _taffel_battery(), 12.3, 4.4),
    ("stacked_horizontal_bar", _taffel_battery(), 12.3, 4.4),
    ("horizontal_bar", _one_series(_LONG), 12.3, 4.4),
    ("vertical_bar", _one_series(_MEDIUM), 9.0, 4.5),
], ids=["stacked-column", "stacked-bar", "bar", "column"])
def test_the_legend_moves_out_of_the_way_rather_than_the_names_being_cut(
        saved, chart_type, series, w_in, h_in):
    """The stacked column's legend sat a fixed 8% below the axis, on top of the
    rotated names. Kept as an obstacle, the only way left to clear it was the
    ellipsis — "tarjoaisi…" four times over, which names nothing. The names are
    what a reader needs beside the bars; the legend can go lower."""
    _render(chart_type, series, w_in, h_in)
    assert saved[-1]["under_legend"] == [], saved[-1]["under_legend"][:4]
    assert not any("…" in t for t in saved[-1]["texts"]), saved[-1]["texts"][:4]


# ── and where there was nothing to fix, nothing changes ─────────────────────

@pytest.mark.parametrize("chart_type,labels", [
    ("stacked_horizontal_bar", None),
    ("horizontal_bar", ["Kyllä", "Ei", "En osaa sanoa"]),
    ("vertical_bar", ["Kyllä", "Ei", "En osaa sanoa"]),
])
def test_a_chart_whose_labels_already_sit_clear_is_drawn_exactly_as_before(
        monkeypatch, chart_type, labels):
    """This runs on every chart, and nearly all of them have nothing to fix."""
    from reportbuilder.render.image import label_fit

    if labels is None:
        few = _TAFFEL[:3]
        stmts = tuple(s for s, _c in few)
        cells = {(lvl, s): Cell(pct=100.0 * c / sum(cs), count=float(c), mean=None)
                 for s, cs in few for lvl, c in zip(_LEVELS, cs)}
        series = SeriesResult(categories=_LEVELS, segments=stmts, cells=cells,
                              base_n={s: sum(cs) for s, cs in few}, statistic="pct",
                              segments_are_groups=False)
    else:
        series = _one_series(labels)
    fitted = _render(chart_type, series, 12.3, 4.4)
    monkeypatch.setattr(label_fit, "fit_category_labels", lambda fig: None)
    untouched = _render(chart_type, series, 12.3, 4.4)
    assert fitted == untouched
