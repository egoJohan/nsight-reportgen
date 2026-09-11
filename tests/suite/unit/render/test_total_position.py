"""Where the Total sits is the author's choice: at the top, or at the bottom.

"Olisiko Total vastausten näkymiseen mahdollista saada mahdollisuus valita
näkyykö Total alimpana vai ylimpänä kaaviossa?"

`total_position` is "auto" (every chart exactly as it was drawn before),
"top" (Total first as the reader meets it: the top row of a horizontal chart,
the left column of a vertical one, the first entry of a line's legend) or
"bottom" (last). It is checked on the drawn figure, not in the data, because
the two orders differ: a grouped horizontal bar stacks its series from the
bottom up, so there the FIRST series is the LOWEST bar — and Total, last in
the data, has always been drawn on top.

The editor offers the control only while a Total is actually drawn. It works
that out from the chart's own settings, with a copy of the engine's rule; both
copies are pinned to one table, here and in web/src/lib/totalPosition.test.ts.
(Johan, 2026-09-11)
"""
from __future__ import annotations

from types import SimpleNamespace

import matplotlib
matplotlib.use("Agg")
import pandas as pd
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.config_schema import (
    clustered_bar_schema, combo_schema, single_series_schema, stacked_schema,
    standard_schema,
)
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.engine import compute
from reportbuilder.stats.percent_base import resolve_show_total


def _var(name, label, values):
    return Variable(name=name, label=label, measurement="nominal", missing_values=[],
                    value_labels=[ValueLabel(value=float(i + 1), label=v)
                                  for i, v in enumerate(values)])


def _model():
    rows = [{"q": float(i % 3 + 1), "sp": float(i % 2 + 1), "ika": float(i % 3 % 2 + 1)}
            for i in range(602)]
    model = QuestionModel(
        variables={"q": _var("q", "Mielipide", ["Hyvä", "Huono", "Ei kumpikaan"]),
                   "sp": _var("sp", "Sukupuoli", ["Naiset", "Miehet"]),
                   "ika": _var("ika", "Ikä", ["18-39", "40-64"])},
        questions=[Question(qid=q, text=q, kind="single", variables=(q,))
                   for q in ("q", "sp", "ika")])
    return model, pd.DataFrame(rows)


def _render(chart_type, position=None, *, var_2=None, layout=None):
    model, df = _model()
    extra = {} if position is None else {"total_position": position}
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="sp", classifying_var_2=var_2,
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(), show_total="on",
                     options={"xtab_layout": layout} if layout else {}, **extra)
    series = compute(model.question("q"), spec, df, model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(12.3), height=Inches(4.6), name="s1"),
                        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS[chart_type](ctx)


@pytest.fixture
def drawn(monkeypatch):
    """Each saved figure, per axes: its tick labels and bar centres where they are
    drawn, and its legend entries in order."""
    seen: list[list[dict]] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()

        def ticks(labels):
            out = []
            for t in labels:
                if t.get_visible() and t.get_text().strip():
                    b = t.get_window_extent(r)
                    out.append((t.get_text(), (b.x0 + b.x1) / 2, (b.y0 + b.y1) / 2))
            return out

        per_ax = []
        for ax in self.axes:
            bars = [(c.get_label(), c.patches[0].get_x() + c.patches[0].get_width() / 2,
                     c.patches[0].get_y() + c.patches[0].get_height() / 2)
                    for c in ax.containers if getattr(c, "patches", None)]
            leg = ax.get_legend()
            per_ax.append({"x": ticks(ax.get_xticklabels()), "y": ticks(ax.get_yticklabels()),
                           "bars": bars,
                           "legend": [t.get_text() for t in leg.get_texts()] if leg else []})
        seen.append(per_ax)
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


# ── where the reader meets the Total: "first" or "last" ────────────────────

def _rank(items, key, reverse=False) -> str:
    ranked = sorted(items, key=key, reverse=reverse)
    at = [i for i, it in enumerate(ranked) if it[0].startswith("Total")]
    assert len(at) == 1, ranked
    return "first" if at[0] == 0 else "last" if at[0] == len(ranked) - 1 else f"middle {at[0]}"


def rows(ax):        # a horizontal chart's bars, read top to bottom
    return _rank(ax["y"], key=lambda t: t[2], reverse=True)


def columns(ax):     # a vertical chart's bars, read left to right
    return _rank(ax["x"], key=lambda t: t[1])


def bars_up(ax):     # the bars within one group of a horizontal chart, top to bottom
    return _rank(ax["bars"], key=lambda t: t[2], reverse=True)


def bars_across(ax):  # the bars within one group of a vertical chart, left to right
    return _rank(ax["bars"], key=lambda t: t[1])


def legend(ax):
    return _rank([(t, i, 0) for i, t in enumerate(ax["legend"])], key=lambda t: t[1])


#: (chart type, how to read it, where Total has always been)
_CHARTS = [
    ("stacked_horizontal_bar", rows, "last"),
    ("stacked_vertical_bar", columns, "last"),
    ("horizontal_bar", bars_up, "first"),
    ("vertical_bar", bars_across, "last"),
    ("line", legend, "last"),
    ("radar", legend, "last"),
]
_IDS = [c[0] for c in _CHARTS]


@pytest.mark.parametrize("chart_type,where,today", _CHARTS, ids=_IDS)
@pytest.mark.parametrize("position", [None, "auto"])
def test_the_default_leaves_total_where_it_always_was(drawn, chart_type, where, today, position):
    """No existing slide moves: a saved report without the field, or with "auto"."""
    _render(chart_type, position)
    assert where(drawn[-1][0]) == today


@pytest.mark.parametrize("chart_type,where,_today", _CHARTS, ids=_IDS)
@pytest.mark.parametrize("position,expected", [("top", "first"), ("bottom", "last")])
def test_the_author_puts_total_at_the_top_or_the_bottom(drawn, chart_type, where, _today,
                                                         position, expected):
    _render(chart_type, position)
    assert where(drawn[-1][0]) == expected


@pytest.mark.parametrize("chart_type,layout", [
    ("stacked_horizontal_bar", None), ("vertical_bar", "grouped"),
    ("horizontal_bar", "small_multiples"),
])
def test_a_crossed_chart_has_no_total_to_place(chart_type, layout):
    """Two variables CROSSED into combos ("Naiset · 18-39") carry no Total, even
    with "Total column" set to Show — which is why the editor does not offer a
    position for one (web/src/lib/totalPosition.ts)."""
    model, df = _model()
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="sp", classifying_var_2="ika",
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(), show_total="on",
                     options={"xtab_layout": layout} if layout else {})
    series = compute(model.question("q"), spec, df, model)
    assert not any(s == "Total" or s.endswith(" · Total") for s in series.segments), series.segments


@pytest.mark.parametrize("position,expected", [("top", "first"), ("bottom", "last")])
def test_each_separate_panel_moves_its_own_total(drawn, position, expected):
    _render("stacked_horizontal_bar", position, var_2="ika", layout="separate")
    panels = [ax for ax in drawn[-1] if any(t[0].startswith("Total") for t in ax["y"])]
    assert panels, drawn[-1]
    for ax in panels:
        assert rows(ax) == expected


@pytest.fixture
def colours(monkeypatch):
    """Each saved figure's series colours by name: bars, lines and legend handles."""
    seen: list[dict[str, tuple]] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        from matplotlib.colors import to_rgba
        got: dict[str, tuple] = {}
        for ax in self.axes:
            for c in ax.containers:
                if getattr(c, "patches", None) and not c.get_label().startswith("_"):
                    got[c.get_label()] = to_rgba(c.patches[0].get_facecolor())
            for line in ax.get_lines():
                if not line.get_label().startswith("_"):
                    got[line.get_label()] = to_rgba(line.get_color())
        seen.append(got)
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


@pytest.mark.parametrize("chart_type", ["vertical_bar", "horizontal_bar", "line", "radar"])
def test_moving_the_total_does_not_recolour_the_groups(colours, chart_type):
    """Colours were handed out by position, so putting Total on top gave
    "Naiset" the colour Total had and moved every group along one. The author
    moved one bar; nothing else on the slide should change. Colours follow the
    series, not the slot it is drawn in."""
    by_position = {}
    for position in (None, "top", "bottom"):
        _render(chart_type, position)
        by_position[position] = colours[-1]
    assert by_position[None], by_position
    assert by_position["top"] == by_position[None]
    assert by_position["bottom"] == by_position[None]


# ── the control ──────────────────────────────────────────────────────────────

@pytest.mark.parametrize("schema", [standard_schema(), clustered_bar_schema(),
                                    stacked_schema(), stacked_schema(with_row_summary=True)])
def test_the_control_sits_right_after_the_total_column_control(schema):
    keys = [f.key for f in schema]
    assert keys[keys.index("show_total") + 1] == "total_position"
    field = schema[keys.index("total_position")]
    assert [v for v, _l in field.options] == ["auto", "top", "bottom"]
    assert field.default == "auto"


@pytest.mark.parametrize("schema", [single_series_schema(), combo_schema()])
def test_charts_without_a_movable_total_do_not_offer_it(schema):
    """A row of pies has no Total. A combo's first series is its bars and its
    second its line, so moving Total would change what is drawn as which."""
    assert "total_position" not in [f.key for f in schema]


#: (show_total, chart_type, statistic, percent_base, drawn) — the SAME table as
#: web/src/lib/totalPosition.test.ts. The editor shows the control only while a
#: Total is drawn, and decides that with a copy of this rule.
_DRAWS_TOTAL = [
    ("on", "vertical_bar", "pct", "classifier", True),
    ("off", "stacked_horizontal_bar", "pct", "total", False),
    ("auto", "stacked_horizontal_bar", "pct", "classifier", True),
    ("auto", "stacked_vertical_bar", "pct", "question", True),
    ("auto", "vertical_bar", "pct", "classifier", False),
    ("auto", "horizontal_bar", "pct", "auto", False),
    ("auto", "vertical_bar", "pct", "total", True),
    ("auto", "line", "mean", "classifier", True),
    ("auto", "radar", "count", "question", True),
]


@pytest.mark.parametrize("show_total,chart_type,statistic,percent_base,expected", _DRAWS_TOTAL)
def test_the_editor_s_rule_is_the_engine_s(show_total, chart_type, statistic, percent_base,
                                           expected):
    spec = SimpleNamespace(show_total=show_total, chart_type=chart_type,
                           statistic=statistic, percent_base=percent_base)
    assert resolve_show_total(spec, True) is expected
