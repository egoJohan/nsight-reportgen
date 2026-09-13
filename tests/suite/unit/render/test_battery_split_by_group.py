"""A battery split by a group names every row, and says each statement once.

"Kuvassa olevat selitystekstit menevät välillä päällekkäin." Six statements
split by country make eighteen bars, and each was named "<statement> ·
<country>" — up to 241 characters. Eighteen of those do not fit a slide, so the
names were printed over each other; and once the names were fitted, every one
was cut to its first words: "Olen luottavainen, että saan…" three times, and
two different "Teen nykyisin sivutöitä…" statements that could no longer be
told apart. Cutting from the end threw away the country, which is exactly
what tells the three rows apart.

So a split battery is drawn as the grouped chart it is: each statement ONCE,
beside its rows — written out and wrapped where it is too long to stand
rotated the way a short group name like "Suomi" does — and each row named by
its group with its own base: "Suomi (n=1016)". A short group name keeps its
rotated label exactly as before. (Johan, 2026-09-11)
"""
from __future__ import annotations

import math

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS

from suite.unit.render._customer_series import STATEMENTS, battery_by_country, sector_by_country

_SLOTS = [(12.3, 4.2), (12.3, 4.6), (9.0, 4.5), (12.3, 3.6)]


def _render(series, chart_type="stacked_horizontal_bar", w_in=12.3, h_in=4.2, *,
            classifying_var_2=None, options=None, row_summary_fn="none"):
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="country", classifying_var_2=classifying_var_2,
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(), options=options or {},
                     row_summary_fn=row_summary_fn)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(w_in), height=Inches(h_in), name="s1"),
                        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS[chart_type](ctx)


def _rect(t, r):
    """A text's own rectangle as (centre, width, height, angle) — rotated ones too."""
    deg = t.get_rotation()
    box = t.get_window_extent(r)
    t.set_rotation(0)
    flat = t.get_window_extent(r)
    t.set_rotation(deg)
    return ((box.x0 + box.x1) / 2, (box.y0 + box.y1) / 2), flat.width, flat.height, math.radians(deg)


def _corners(rect):
    (cx, cy), w, h, ang = rect
    ca, sa = math.cos(ang), math.sin(ang)
    return [(cx + dx * ca - dy * sa, cy + dx * sa + dy * ca)
            for dx, dy in ((-w / 2, -h / 2), (w / 2, -h / 2), (w / 2, h / 2), (-w / 2, h / 2))]


def _meet(a, b) -> bool:
    """Whether two rectangles intersect by more than half a pixel — at ANY two
    angles. A name set at 30° under a column beside a statement set flat has to
    be tested on the edges of both; projecting onto one rectangle's frame alone
    reported overlaps that were not there."""
    pa, pb = _corners(a), _corners(b)
    for poly in (pa, pb):
        for i in range(4):
            (x0, y0), (x1, y1) = poly[i], poly[(i + 1) % 4]
            nx, ny = y0 - y1, x1 - x0
            norm = math.hypot(nx, ny) or 1.0
            nx, ny = nx / norm, ny / norm
            proj_a = [x * nx + y * ny for x, y in pa]
            proj_b = [x * nx + y * ny for x, y in pb]
            if min(max(proj_a), max(proj_b)) - max(min(proj_a), min(proj_b)) <= 0.5:
                return False                      # a separating axis: apart
    return True


@pytest.fixture
def drawn(monkeypatch):
    """Per saved figure: row names, group labels and every other text, as drawn."""
    seen: list[dict] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        names, groups = [], []
        for ax in self.axes:
            for labels in (ax.get_yticklabels(), ax.get_xticklabels()):
                names += [t for t in labels if t.get_visible() and t.get_text().strip()
                          and not t.get_text().replace(".", "").isdigit()]
            groups += [t for t in ax.texts if t.get_gid() == "nsight-group"]
        name_rects = [(t.get_text(), _rect(t, r)) for t in names]
        group_rects = [(t.get_text(), _rect(t, r)) for t in groups]
        legend_boxes = [leg.get_window_extent(r) for leg in
                        [ax.get_legend() for ax in self.axes if ax.get_legend()]]
        hits = []
        everything = name_rects + group_rects
        for i, (ta, ra) in enumerate(everything):
            for tb, rb in everything[i + 1:]:
                if _meet(ra, rb):
                    hits.append((ta.replace("\n", " "), tb.replace("\n", " ")))
        for tg, g in groups and [(t.get_text(), t.get_window_extent(r)) for t in groups] or []:
            for lb in legend_boxes:
                if g.overlaps(lb):
                    hits.append((tg.replace("\n", " "), "legend"))
        seen.append({"names": [t.get_text() for t in names],
                     "groups": [t.get_text() for t in groups],
                     "sizes": [t.get_fontsize() for t in names + groups],
                     "hits": hits})
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


def _flat(text: str) -> str:
    return " ".join(text.split())


# ── the reported slide ───────────────────────────────────────────────────────

@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar"])
@pytest.mark.parametrize("w_in,h_in", _SLOTS)
def test_every_row_is_named_by_its_own_group_and_base(drawn, chart_type, w_in, h_in):
    series = battery_by_country()
    _render(series, chart_type, w_in, h_in)
    want = [f"{bar.split(' · ')[1]} (n={series.base_n[bar]})" for bar in series.segments]
    got = [_flat(t) for t in drawn[-1]["names"]]
    assert sorted(got) == sorted(want), got
    assert not any("…" in t for t in got), got


@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar"])
@pytest.mark.parametrize("w_in,h_in", _SLOTS)
def test_each_statement_is_said_once_and_can_be_told_apart(drawn, chart_type, w_in, h_in):
    _render(battery_by_country(), chart_type, w_in, h_in)
    groups = [_flat(t) for t in drawn[-1]["groups"]]
    assert len(groups) == len(STATEMENTS), groups
    for stmt, shown in zip(STATEMENTS, groups):
        full = _flat(stmt)
        # In full — or, for a paragraph too long for any slide, a clear start of
        # it: at least five words of it, word for word.
        assert shown == full or (shown.endswith("…") and full.startswith(shown[:-1].rstrip())
                                 and len(shown[:-1].split()) >= 5), (stmt, shown)
    # And no two are shown alike — which is the whole point.
    assert len(set(groups)) == len(groups), groups
    # The two statements that start alike are both still what they are.
    assert any("ansaitakseni" in g for g in groups), groups
    assert any("oppiakseni" in g for g in groups), groups


@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar"])
@pytest.mark.parametrize("w_in,h_in", _SLOTS)
def test_nothing_is_printed_over_anything(drawn, chart_type, w_in, h_in):
    _render(battery_by_country(), chart_type, w_in, h_in)
    assert drawn[-1]["hits"] == [], drawn[-1]["hits"][:4]


@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar"])
def test_it_stays_readable(drawn, chart_type):
    _render(battery_by_country(), chart_type, 12.3, 4.2)
    assert min(drawn[-1]["sizes"]) >= 7.5


def test_a_top_box_column_still_sits_beside_its_rows(drawn):
    """The row-summary column is keyed by bar, so grouping must not shift it."""
    _render(battery_by_country(), row_summary_fn="top2_sum")
    assert drawn[-1]["hits"] == [], drawn[-1]["hits"][:4]


# ── and what must not change ─────────────────────────────────────────────────

def test_a_short_group_name_keeps_its_rotated_label(drawn):
    """Country crossed with sector: "Suomi" stands beside its rows as it did."""
    _render(sector_by_country(), classifying_var_2="Sektori")
    assert set(drawn[-1]["groups"]) == {"Suomi", "Ruotsi", "Saksa"}


def test_a_short_group_name_is_still_drawn_rotated(monkeypatch):
    original = Figure.savefig
    rotations = []

    def spy(self, *a, **k):
        rotations.extend(round(t.get_rotation()) for ax in self.axes for t in ax.texts
                         if t.get_gid() == "nsight-group")
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    _render(sector_by_country(), classifying_var_2="Sektori")
    assert rotations and all(r == 90 for r in rotations), rotations


def test_separate_panels_need_a_second_variable(monkeypatch):
    """A battery keeps its grouping in `segment_primary` now; "Separate panels"
    left in a chart's options from an earlier second variable must not start
    splitting it into one panel per statement."""
    original = Figure.savefig
    axes = []

    def spy(self, *a, **k):
        axes.append(len(self.axes))
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    _render(battery_by_country(), options={"xtab_layout": "separate"})
    assert axes == [1]


# ── the fit's last resort keeps what tells the rows apart ───────────────────

def test_a_name_that_must_be_cut_keeps_its_group_and_base():
    from reportbuilder.render.image.label_fit import cap_keeping_tail

    name = "Olen luottavainen, että saan työstä riittävän suuren toimeentulon · Suomi (n=1016)"
    cut = cap_keeping_tail(name, 30, 1)
    assert cut.endswith(" · Suomi (n=1016)") and "…" in cut, cut
    assert cut.startswith("Olen"), cut


def test_a_plain_name_is_cut_from_the_end_as_before():
    from reportbuilder.render.image._mpl import wrap_label_capped
    from reportbuilder.render.image.label_fit import cap_keeping_tail

    text = "parantaisi nykyisten tuotteiden makua ja rapeutta huomattavasti"
    assert cap_keeping_tail(text, 30, 1) == wrap_label_capped(text, 30, 1)


def test_the_base_alone_is_kept_too():
    from reportbuilder.render.image.label_fit import cap_keeping_tail

    cut = cap_keeping_tail("hyvinvointialueen palveluksessa työskentelevät (n=97)", 20, 1)
    assert cut.endswith(" (n=97)") and "…" in cut, cut


def test_a_name_that_fits_is_left_whole():
    from reportbuilder.render.image.label_fit import cap_keeping_tail

    assert cap_keeping_tail("Suomi (n=1016)", 30, 1) == "Suomi (n=1016)"
