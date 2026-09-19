"""Column names are printed flat, and rotated only when flat would lose one.

Every column builder rotated its category names by 30 degrees whatever they
were: "Mieheksi" and "Naiseksi", two short names under columns half the slide
wide, came out tilted, and "Muualla Etelä-Suomessa" wrapped AND rotated. The
line chart rotated whenever there were more than four names. The combo already
tried flat first; now every builder does, through one helper.
(visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_FEW = ("Mieheksi", "Naiseksi")
_REGIONS = ("Pääkaupunkiseudulla", "Muualla Etelä-Suomessa", "Länsi-Suomessa",
            "Pohjois-Suomessa", "Itä-Suomessa")
#: Attendo's brand battery: fourteen statements, too many to stand flat in full.
_MANY = ("Mahdollistaa hyvän arjen", "Tarjoaa laadukkaita hoivapalveluita",
         "Tarjoaa monipuolista ruokaa", "Inhimillinen", "Luotettava", "Välittävä",
         "Ammattitaitoinen", "Kohtelee ihmisiä arvostavasti", "Edelläkävijä",
         "Rehellinen", "Kehittää toimintatapojaan", "Välinpitämätön", "Ahne",
         "Kodikas")


def _names(chart_type: str, cats) -> list:
    cells = {(c, "Total"): Cell(pct=10.0 + i) for i, c in enumerate(cats)}
    series = SeriesResult(categories=cats, segments=("Total",), cells=cells,
                          base_n={"Total": 1000}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    seen: list = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        for ax in self.axes:
            seen.extend((t.get_text(), t.get_rotation()) for t in ax.get_xticklabels()
                        if t.get_visible() and t.get_text())
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        _f.Figure.savefig = original
    return seen


@pytest.mark.parametrize("chart_type", ["vertical_bar", "line"])
@pytest.mark.parametrize("cats", [_FEW, _REGIONS], ids=["two", "regions"])
def test_names_that_fit_are_flat(chart_type, cats):
    names = _names(chart_type, cats)
    assert len(names) == len(cats)
    assert all(rot == 0 for _t, rot in names), f"rotated although they fit: {names}"


@pytest.mark.parametrize("chart_type", ["vertical_bar", "line"])
def test_names_that_do_not_fit_flat_are_rotated_whole(chart_type):
    """Flat at any size above the floor would cut the statements short; rotated
    they are printed in full, which is the better of the two."""
    names = _names(chart_type, _MANY)
    assert len(names) == len(_MANY)
    assert all(rot != 0 for _t, rot in names), names
    assert not any("…" in t for t, _rot in names), names


@pytest.mark.parametrize("chart_type", ["vertical_bar", "line"])
def test_a_flat_name_keeps_its_words_whole(chart_type):
    """ "Pääkaupunkiseudulla" was printed as "Pääkaupunkiseu / dulla"."""
    names = _names(chart_type, _REGIONS)
    words = {w for t, _rot in names for w in t.split()}
    assert "Pääkaupunkiseudulla" in words, names


def test_a_stacked_columns_group_names_are_flat_when_they_fit():
    """The stacked column had its own copy of the unconditional rotation:
    "Mieheksi (n=478)" under a column a sixth of the slide wide was printed at
    30 degrees. (visual QA, 2026-09-19)"""
    groups = ("Mieheksi", "Naiseksi", "Muuksi")
    cats = ("Erittäin huono", "Huono", "Hyvä", "Erittäin hyvä")
    cells = {(c, g): Cell(pct=25.0) for c in cats for g in groups}
    series = SeriesResult(categories=cats, segments=groups, cells=cells,
                          base_n={g: 400 for g in groups}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type="stacked_vertical_bar", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    seen: list = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        seen.extend((t.get_text(), t.get_rotation()) for ax in self.axes
                    for t in ax.get_xticklabels() if t.get_visible() and t.get_text())
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["stacked_vertical_bar"](ctx)
    finally:
        _f.Figure.savefig = original
    assert seen and all(rot == 0 for _t, rot in seen), seen
