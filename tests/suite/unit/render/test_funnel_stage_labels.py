"""A funnel's stage names do not print through each other.

The stage names sit in the gutter to the right of the bars, one per row,
wrapped at a fixed 28 characters with no limit on how many lines that took and
a fixed 11pt whatever the row was worth. On Attendo's brand battery — fourteen
stages, several of them two lines — consecutive names were drawn on top of one
another: "Tarjoaa laadukkaita / hoivapalveluita" straight through "Kohtelee
ihmisiä / arvostavasti", measured at 18.3px by the overlap oracle.

The horizontal bar builder solved this long ago: size the label to the row band
and cap it to the lines that band can hold, ellipsising as a last resort,
because "an ellipsis is preferred over labels overlapping each other"
(`wrap_label_capped`). The funnel never had it.
"""
from __future__ import annotations

import pathlib
import sys

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.funnel as F

# The project's own overlap rule — see test_combo_category_names for why a
# plain box intersection is the wrong instrument here.
_ACCEPTANCE = pathlib.Path(__file__).resolve().parents[4] / "acceptance"
if str(_ACCEPTANCE) not in sys.path:
    sys.path.insert(0, str(_ACCEPTANCE))
overlap = pytest.importorskip("oracles.overlap")

#: The customer's own statements, which is what makes them this long.
MANY = ["Välittävä", "Tarjoaa monipuolista ruokaa", "Inhimillinen",
        "Kehittää toimintatapojaan", "Mahdollistaa hyvän arjen",
        "Tarjoaa laadukkaita hoivapalveluita", "Kohtelee ihmisiä arvostavasti",
        "Kodikas", "Luotettava", "Ahne", "Edelläkävijä", "Rehellinen",
        "Ammattitaitoinen", "Välinpitämätön"]
FEW = ["Spontaani tunnettuus", "Autettu tunnettuus", "Harkinta", "Kokeilu"]


def _render(cats: list[str]) -> dict:
    var = Variable(name="q", label="Mielikuva", measurement="nominal",
                   missing_values=[],
                   value_labels=[ValueLabel(value=float(i + 1), label=c)
                                 for i, c in enumerate(cats)])
    rows = [{"q": float(i + 1)}
            for i in range(len(cats)) for _ in range(len(cats) * 10 - i * 5)]
    model = QuestionModel(
        variables={"q": var},
        questions=[Question(qid="q", text=var.label, kind="single", variables=("q",))])
    spec = ChartSpec(question_ref="q", chart_type="funnel", statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(11.6), height=Inches(4.0), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    seen: dict = {}
    original = F.render_png

    def _spy(fig):
        fig.canvas.draw()
        seen["collisions"] = [str(c) for c in overlap.collisions(fig)]
        seen["texts"] = [t.get_text() for a in fig.axes for t in a.texts]
        return original(fig)

    F.render_png = _spy
    try:
        F.build_image_funnel(ctx)
    finally:
        F.render_png = original
    return seen


def test_the_stage_names_are_drawn_at_all():
    """Guards the test below."""
    got = _render(MANY)
    assert any("Välittävä" in t for t in got["texts"])


def test_many_stages_do_not_print_through_each_other():
    hits = _render(MANY)["collisions"]
    assert hits == [], f"{len(hits)} collisions, e.g. {hits[:3]}"


def test_a_short_funnel_keeps_its_names_in_full():
    """Capping is for rows too thin to hold the text. Four roomy stages must
    still read in full — no ellipsis, nothing dropped."""
    got = _render(FEW)
    assert got["collisions"] == []
    joined = " ".join(got["texts"])
    for name in FEW:
        assert name.split()[-1] in joined, f"{name!r} was cut on a roomy funnel"
