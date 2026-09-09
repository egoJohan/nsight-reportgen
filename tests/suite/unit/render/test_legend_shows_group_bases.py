"""A chart comparing groups says how many people each group is.

"n-luvut eivät tule näkyviin kaavioihin jos kaaviossa näytetään usean eri
kohderyhmän tuloksia" — a slide comparing women with men gave the reader two
percentages and no way to know one was 501 people and the other 502. The
footer's N is the whole slide's base, which is not either group's.

So a SERIES legend names the base: "Naiset (n=501)", "Miehet (n=502)", and
"Total (n=1003)" when the Total column is drawn.

Only where the series ARE the groups. A stacked bar's legend is the answer
scale — its bars are the groups — so annotating it would attach a respondent
count to "Melko hyvä", which is not a group and has no base of its own.
(Johan, 2026-09-09)
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B

_GROUPS = ["Naiset", "Miehet"]


def _ctx(chart_type="vertical_bar"):
    opinion = Variable(name="mielipide", label="Mielipide", measurement="nominal",
                       missing_values=[],
                       value_labels=[ValueLabel(value=1.0, label="Hyvä"),
                                     ValueLabel(value=2.0, label="Huono"),
                                     ValueLabel(value=3.0, label="Ei kumpikaan")])
    sukup = Variable(name="sukupuoli", label="Sukupuoli", measurement="nominal",
                     missing_values=[],
                     value_labels=[ValueLabel(value=1.0, label="Naiset"),
                                   ValueLabel(value=2.0, label="Miehet")])
    # 501 women, 502 men — the numbers from the report
    rows = ([{"mielipide": float(i % 3 + 1), "sukupuoli": 1.0} for i in range(501)]
            + [{"mielipide": float(i % 3 + 1), "sukupuoli": 2.0} for i in range(502)])
    model = QuestionModel(
        variables={"mielipide": opinion, "sukupuoli": sukup},
        questions=[Question(qid="mielipide", text="Mielipide", kind="single",
                            variables=("mielipide",)),
                   Question(qid="sukupuoli", text="Sukupuoli", kind="single",
                            variables=("sukupuoli",))])
    spec = ChartSpec(question_ref="mielipide", chart_type=chart_type, statistic="pct",
                     classifying_var="sukupuoli", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("mielipide"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    return RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                         slot=Slot(slide_index=0, left=Inches(1), top=Inches(1),
                                   width=Inches(9), height=Inches(4.5), name="s1"),
                         style=StyleSpec(), spec=spec, series=series,
                         fmt=spec.number_format), series


def _legend_labels(chart_type, builder):
    ctx, series = _ctx(chart_type)
    caught: dict = {}
    original = B.render_png

    def _spy(fig):
        ax = fig.axes[0]
        _h, labels = ax.get_legend_handles_labels()
        caught["labels"] = labels
        leg = ax.get_legend()
        caught["drawn"] = [t.get_text() for t in leg.get_texts()] if leg else []
        return original(fig)

    B.render_png = _spy
    try:
        getattr(B, builder)(ctx)
    finally:
        B.render_png = original
    return caught, series


def test_a_grouped_column_names_each_groups_base():
    got, series = _legend_labels("vertical_bar", "build_image_column")
    assert "Naiset (n=501)" in got["labels"], got["labels"]
    assert "Miehet (n=502)" in got["labels"], got["labels"]


def test_the_total_column_says_its_base_too():
    got, series = _legend_labels("vertical_bar", "build_image_column")
    if "Total" in series.segments and any("Total" in l for l in got["labels"]):
        assert f"Total (n={series.base_n['Total']})" in got["labels"], got["labels"]


def test_the_horizontal_form_does_the_same():
    got, _s = _legend_labels("horizontal_bar", "build_image_bar")
    assert any(l.startswith("Naiset (n=") for l in got["labels"]), got["labels"]


def test_a_stacked_legend_is_the_scale_and_gets_no_base():
    """Its bars are the groups; the legend names answer categories, which have
    no base of their own."""
    got, _s = _legend_labels("stacked_horizontal_bar", "build_image_bar_stacked")
    assert got["labels"], "the stack should still have a legend"
    assert not any("(n=" in l for l in got["labels"]), got["labels"]


def test_the_drawn_legend_keeps_the_names():
    """`_legend_below` shortens a numeric SCALE to bare numbers. Group names are
    not a scale — an age band starts with a digit and must not become "25"."""
    got, _s = _legend_labels("vertical_bar", "build_image_column")
    assert any("Naiset" in t for t in got["drawn"]), got["drawn"]


def test_a_line_chart_names_its_groups_bases_too():
    """Its lines are the groups, same as a grouped bar's bars."""
    import reportbuilder.render.image.line as L

    ctx, _series = _ctx("line")
    caught: dict = {}
    original = L.render_png

    def _spy(fig):
        caught["labels"] = fig.axes[0].get_legend_handles_labels()[1]
        return original(fig)

    L.render_png = _spy
    try:
        L.build_image_line(ctx)
    finally:
        L.render_png = original
    assert any(l.startswith("Naiset (n=") for l in caught["labels"]), caught["labels"]
