"""The n= after a group's name can be switched off.

"Onko meillä mitään tapaa haluttaessa poistaa legendiin tulevat N-luvut?"

A series that IS a group carries its base — "Suomi (n=516)" — because a slide
comparing groups otherwise gives two percentages and no way to know one is 516
people and the other 448. Right default, and there was no way to decline it.

`elements.group_base` turns it off. One switch for every place a group states
its own base, not just the legend: a stacked chart says it on the BARS, because
there the bars are the groups and the legend is the answer scale, and a switch
that silenced one and not the other would look broken from whichever chart the
author happened to be on.

Default True, so nothing already saved changes. (Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("Päivittäin", "Viikoittain", "Harvemmin")
_GROUPS = ("Suomi", "Ruotsi", "Saksa")
_BASES = {"Suomi": 516, "Ruotsi": 448, "Saksa": 532}


def _series() -> SeriesResult:
    cells = {(c, g): Cell(pct=float(20 + i * 5 + j))
             for i, c in enumerate(_CATS) for j, g in enumerate(_GROUPS)}
    return SeriesResult(categories=_CATS, segments=_GROUPS, cells=cells,
                        base_n={"Total": 1496, **_BASES}, statistic="pct")


def _texts(chart_type: str, *, group_base: bool | None = None) -> list[str]:
    elements = ElementToggles(
        title=True, subtitle=True, legend=True, n=True, axis_names=True,
        filter_var=True, data_labels=True,
        **({} if group_base is None else {"group_base": group_base}))
    spec = ChartSpec(
        question_ref="q1", chart_type=chart_type, statistic="pct",
        classifying_var="country", number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=elements)
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(), fmt=spec.number_format)

    seen: list[str] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        for axes in self.axes:
            leg = axes.get_legend()
            if leg is not None:
                seen.extend(t.get_text() for t in leg.get_texts())
            seen.extend(lbl.get_text() for lbl in axes.get_yticklabels())
            seen.extend(lbl.get_text() for lbl in axes.get_xticklabels())
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        Figure.savefig = original
    return [t for t in seen if t.strip()]


CHARTS = ["vertical_bar", "horizontal_bar", "line", "stacked_horizontal_bar"]


@pytest.mark.parametrize("chart_type", CHARTS)
def test_the_base_is_shown_by_default(chart_type):
    """Unchanged for every report already saved."""
    assert any("(n=516)" in t for t in _texts(chart_type)), chart_type


@pytest.mark.parametrize("chart_type", CHARTS)
def test_turning_it_off_removes_the_base(chart_type):
    texts = _texts(chart_type, group_base=False)
    assert not [t for t in texts if "(n=" in t], texts


@pytest.mark.parametrize("chart_type", CHARTS)
def test_the_group_is_still_named(chart_type):
    """Only the number goes; a nameless series would be worse than a verbose one."""
    texts = _texts(chart_type, group_base=False)
    assert any("Suomi" in t for t in texts), texts


def test_explicitly_on_is_the_same_as_the_default():
    assert _texts("vertical_bar", group_base=True) == _texts("vertical_bar")


def test_the_toggle_defaults_to_on():
    assert ElementToggles().group_base is True
