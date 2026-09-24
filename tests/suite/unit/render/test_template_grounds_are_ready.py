"""A template's grounds are drawn once, ahead of time, and a box move keeps them.

The layout editor composes its sample on a "ground" — the template's empty
slide — exactly as a slide preview does. Only the ground needs LibreOffice, and
starting LibreOffice is 2.5-3.5 s on staging's one core. Two things made the
editor pay that again and again:

  * the ground's cache key carried the chart area's POSITION, which the empty
    slide never shows, so every drag or resize of a box in the editor started
    LibreOffice for a picture identical to the one already on disk;
  * a layout's ground was drawn the first time somebody picked it, while they
    waited.

So the position is out of the key, and `warm_grounds` draws every offered
layout's ground in the background — one at a time, at the lowest CPU priority,
skipping what is already drawn. The resident LibreOffice that was tried first
for this grew by ~130 MB per large export on a box with no swap; this needs no
process that stays up. (Johan, 2026-09-24)
"""
from __future__ import annotations

import copy
import pathlib

import pytest

pytestmark = pytest.mark.unit

_TEMPLATE = "input/Attendo Bränditutkimus Marraskuu 2025.pptx"


@pytest.fixture
def template():
    if not pathlib.Path(_TEMPLATE).exists():
        pytest.skip(f"{_TEMPLATE} not available locally")
    return _TEMPLATE


def _moved(style, dx: int, dy: int):
    from reportbuilder.render.base import Slot
    moved = copy.deepcopy(style)
    s = moved.chart_slot
    moved.chart_slot = Slot(slide_index=s.slide_index, left=s.left + dx, top=s.top + dy,
                            width=s.width - dx, height=s.height - dy, name=s.name)
    return moved


def test_moving_the_chart_area_keeps_the_ground(template):
    from reportbuilder.render.image.fast_preview import _key
    from reportbuilder.render.template_cache import resolve
    style = copy.deepcopy(resolve(template).style)
    assert style.chart_slot is not None
    assert _key(style, 110) == _key(_moved(style, 91440, 182880), 110)


def test_another_layout_is_another_ground(template):
    from reportbuilder.render.image.fast_preview import _key
    from reportbuilder.render.template_cache import resolve_layout
    from reportbuilder.render.template_check import rank_layouts
    from pptx import Presentation
    a, b = [c.index for c in rank_layouts(Presentation(template))][:2]
    assert _key(resolve_layout(template, a), 110) != _key(resolve_layout(template, b), 110)


def test_warming_draws_every_offered_layout_and_then_nothing_starts_libreoffice(
        template, require_soffice, monkeypatch, tmp_path):
    from pptx import Presentation

    from reportbuilder.render.image import fast_preview as FP
    from reportbuilder.render.template_cache import resolve_layout
    from reportbuilder.render.template_check import rank_layouts

    monkeypatch.setattr(FP, "_CACHE", tmp_path / "grounds")
    offered = [c.index for c in rank_layouts(Presentation(template))][:3]
    drawn = FP.warm_grounds(template, {}, layouts=offered, pause_s=0)
    assert drawn == len(offered)

    calls = []
    monkeypatch.setattr(FP, "pptx_to_pdf", lambda *a, **k: calls.append(a) or None)
    for i in offered:
        style = resolve_layout(template, i)
        assert FP.ground_image(style) is not None
        assert FP.ground_image(_moved(style, 45720, 45720)) is not None
    assert calls == [], "a warmed ground was drawn again"


def test_warming_twice_draws_nothing_the_second_time(template, require_soffice,
                                                      monkeypatch, tmp_path):
    from pptx import Presentation

    from reportbuilder.render.image import fast_preview as FP
    from reportbuilder.render.template_check import rank_layouts

    monkeypatch.setattr(FP, "_CACHE", tmp_path / "grounds")
    offered = [c.index for c in rank_layouts(Presentation(template))][:2]
    assert FP.warm_grounds(template, {}, layouts=offered, pause_s=0) == len(offered)
    assert FP.warm_grounds(template, {}, layouts=offered, pause_s=0) == 0


def test_one_pre_drawing_process_at_a_time_at_the_lowest_priority(monkeypatch):
    """A process of its own, so an export never waits for its Python; one for
    the whole server, so a busy morning of template uploads is one queue."""
    import subprocess

    from reportbuilder.render.image import fast_preview as FP

    started = []

    class Held:
        def __init__(self, args, **kw):
            started.append((args, kw))
        def wait(self):
            import time
            time.sleep(0.3)

    monkeypatch.setenv("NSIGHT_WARM_GROUNDS", "1")
    monkeypatch.setattr(subprocess, "Popen", Held)
    monkeypatch.setattr(FP, "_warming", {})
    monkeypatch.setattr(FP, "_pending", {})
    monkeypatch.setattr(FP, "_warmed", set())
    assert FP.start_warming("a.pptx", {}) is True
    assert FP.start_warming("b.pptx", {}) is False, "a second one started alongside"
    args, kw = started[0]
    assert args[1:4] == ["-m", "reportbuilder.render.image.fast_preview", "warm"]
    assert kw.get("preexec_fn") is not None


def test_a_template_asked_for_meanwhile_is_drawn_next_not_forgotten(monkeypatch):
    """Found measuring: the second template an author opened was dropped while
    the first was still being drawn, and its layouts stayed cold."""
    import subprocess
    import time

    from reportbuilder.render.image import fast_preview as FP

    started = []

    class Done:
        def __init__(self, args, **kw):
            started.append(args[4])
        def wait(self):
            time.sleep(0.2)
            return 0

    monkeypatch.setenv("NSIGHT_WARM_GROUNDS", "1")
    monkeypatch.setattr(subprocess, "Popen", Done)
    monkeypatch.setattr(FP, "_warming", {})
    monkeypatch.setattr(FP, "_pending", {})
    monkeypatch.setattr(FP, "_warmed", set())
    FP.start_warming("a.pptx", {})
    FP.start_warming("b.pptx", {})
    deadline = time.monotonic() + 3
    while len(started) < 2 and time.monotonic() < deadline:
        time.sleep(0.05)
    assert started == ["a.pptx", "b.pptx"]
    while FP._warming and time.monotonic() < deadline:
        time.sleep(0.05)
    assert FP.start_warming("a.pptx", {}) is False, "drawn already, started again"
    assert FP.start_warming("a.pptx", {"background": "000000"}) is True, \
        "new overrides are new grounds"


def test_switched_off_it_starts_nothing(monkeypatch):
    from reportbuilder.render.image import fast_preview as FP
    monkeypatch.setenv("NSIGHT_WARM_GROUNDS", "0")
    assert FP.start_warming("a.pptx", {}) is False
