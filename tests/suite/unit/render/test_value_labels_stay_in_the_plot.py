"""A number a chart prints has to be ON the chart.

A stacked bar calls a number out ABOVE its own bar when the segment is too
narrow to hold it — the gap below a bar belongs to the next bar down, so up is
the only direction available. The TOP row has no row above it, and the callouts
for its slivers were drawn above the plot frame altogether: two "2 %" floating
over the chart, attached to nothing a reader can follow, on a study whose first
group answered mostly at the top of the scale.

`annotation_clip=False` is why they appear at all rather than vanishing, and it
is the right setting — a number the chart computed must not be silently
dropped. The fix is room: when the top row carries callouts, the axis has to
keep headroom for them, exactly as `_STACK_BAR_H_CALLOUT` already thins the
bars to keep room between rows.
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image.bars import VALUE_GID
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B

#: A seven-point scale whose first two levels are slivers in every group — the
#: shape that forces a callout. Counts per (group, level).
_SHAPE = {
    "En lainkaan": [1, 1, 2, 3, 9, 20, 10],
    "En kovin":    [1, 1, 3, 5, 8, 14, 12],
    "En osaa":     [2, 1, 2, 5, 12, 11, 7],
    "Melko":       [1, 1, 2, 5, 9, 14, 9],
    "Erittäin":    [1, 2, 4, 3, 10, 13, 10],
}
_LEVELS = ["1- Ei lainkaan ylpeä", "2", "3", "4", "5", "6", "7- Erittäin ylpeä"]


def _draw(chart_type: str, builder: str, slot_in: float = 11.6) -> dict:
    """Render one chart and report every value label that left the plot."""
    scale = Variable(name="q", label="Kuinka ylpeä olet omasta työstäsi?",
                     measurement="nominal", missing_values=[],
                     value_labels=[ValueLabel(value=float(i + 1), label=lbl)
                                   for i, lbl in enumerate(_LEVELS)])
    group = Variable(name="g", label="Suosittelisitko", measurement="nominal",
                     missing_values=[],
                     value_labels=[ValueLabel(value=float(i + 1), label=g)
                                   for i, g in enumerate(_SHAPE)])
    rows = []
    for gi, counts in enumerate(_SHAPE.values()):
        for level, n in enumerate(counts):
            rows += [{"q": float(level + 1), "g": float(gi + 1)}] * n
    model = QuestionModel(
        variables={"q": scale, "g": group},
        questions=[Question(qid="q", text=scale.label, kind="single", variables=("q",)),
                   Question(qid="g", text=group.label, kind="single", variables=("g",))])
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(slot_in), height=Inches(4.0), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    found: dict = {}
    original = B.render_png

    def _spy(fig):
        fig.canvas.draw()
        r = fig.canvas.get_renderer()
        ax = fig.axes[0]
        area = ax.get_window_extent(r)
        escaped = []
        for t in ax.texts:
            if t.get_gid() != VALUE_GID or not t.get_visible():
                continue
            from matplotlib.text import Text
            b = Text.get_window_extent(t, r)
            # How far the box sticks out on each side; 0.5px of rounding is not
            # an escape.
            out = max(area.y0 - b.y0, b.y1 - area.y1,
                      area.x0 - b.x0, b.x1 - area.x1)
            if out > 0.5:
                escaped.append((t.get_text(), round(out, 1)))
        found["escaped"] = escaped
        found["labels"] = [t.get_text() for t in ax.texts if t.get_gid() == VALUE_GID]
        return original(fig)

    B.render_png = _spy
    try:
        getattr(B, builder)(ctx)
    finally:
        B.render_png = original
    return found


CHARTS = [("stacked_horizontal_bar", "build_image_bar_stacked"),
          ("stacked_vertical_bar", "build_image_column_stacked")]


@pytest.mark.parametrize("chart_type,builder", CHARTS)
def test_the_slivers_are_called_out_at_all(chart_type, builder):
    """Guards the test itself: with no callout there is nothing to escape."""
    got = _draw(chart_type, builder)
    assert len(got["labels"]) > 10, got["labels"]


@pytest.mark.parametrize("chart_type,builder", CHARTS)
def test_no_value_label_is_drawn_outside_the_plot(chart_type, builder):
    got = _draw(chart_type, builder)
    assert got["escaped"] == [], got["escaped"]


@pytest.mark.parametrize("slot_in", [8.0, 9.5, 11.6, 13.3])
def test_it_holds_at_every_slide_width(slot_in):
    """Headroom measured in data units alone would scale with the slide; the
    callout is a fixed number of points tall, so the two must be reconciled at
    whatever size the slot happens to be."""
    got = _draw("stacked_horizontal_bar", "build_image_bar_stacked", slot_in)
    assert got["escaped"] == [], got["escaped"]
