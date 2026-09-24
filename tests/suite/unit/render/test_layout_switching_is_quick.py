"""Choosing a layout in the template settings must not cost a re-parse or a boot.

Reported from staging, 2026-09-21: "When selecting a new layout from the
template settings the refresh takes a long time. The areas should update
immediately and the text + chart image rendered fast", then "Now it takes over
30s", then "the rendering of the previews are just like 1s. That would be
enough also for template layout change."

Three costs, measured on the real customer templates (this machine; a one-core
host is several times slower):

  * the .pptx was parsed on EVERY request for a forced layout — 1.4 s a time,
    because the cache was keyed on the file alone and a chosen layout missed it;
  * the boxes endpoint built a sample chart slide just to ask where the footer
    goes — 0.5-1.1 s, and `footer_top` never looks at what was drawn;
  * every new layout's ground was a fresh LibreOffice, and starting it is the
    whole cost — 2.5-3.5 s whether the deck is 7 MB or 0.12 MB. A ground is
    now drawn once per layout, ahead of time, and moving a box does not
    invalidate it; see `test_template_grounds_are_ready.py`.
"""
from __future__ import annotations

import pathlib

import pytest

pytestmark = pytest.mark.unit

_TEMPLATE = "input/Attendo Bränditutkimus Marraskuu 2025.pptx"
_HEAVY = "input/Holiday Club_Loyalty tutkimus_raportti_19.2.2026.pptx"


@pytest.fixture
def template() -> str:
    if not pathlib.Path(_TEMPLATE).exists():
        pytest.skip(f"{_TEMPLATE} not available locally")
    return _TEMPLATE


def test_a_chosen_layout_is_parsed_once(template, monkeypatch):
    from reportbuilder.render import template_cache as TC

    TC._resolve_forced.cache_clear()
    parses = []
    real = TC.load_style_spec

    def counting(path, force_layout=None):
        parses.append(force_layout)
        return real(path, force_layout=force_layout)

    monkeypatch.setattr(TC, "load_style_spec", counting)
    first = TC.resolve_layout(template, 4)
    again = TC.resolve_layout(template, 4)
    other = TC.resolve_layout(template, 5)

    assert parses == [4, 5], f"parsed {len(parses)} times for two layouts"
    assert first is not again, "each caller gets its own copy to correct"
    assert getattr(first.profile, "layout_index", None) == 4
    assert getattr(other.profile, "layout_index", None) == 5


def test_a_correction_does_not_reach_the_cache(template):
    """The copy is what callers patch; the next reader must not see it."""
    from reportbuilder.render.style_spec import apply_template_overrides
    from reportbuilder.render.template_cache import resolve_layout

    mine = resolve_layout(template, 4)
    apply_template_overrides(mine, {"accent": "FF5000"})
    assert resolve_layout(template, 4).accent != "FF5000"


@pytest.mark.parametrize("path,layout", [(_TEMPLATE, 4), (_HEAVY, 3)])
def test_a_ground_carries_only_the_layout_it_draws(path, layout, tmp_path):
    """A deck saved from a customer template brings every master, every layout
    and all their images — 7.2 MB for one empty slide on Holiday Club, and
    LibreOffice loads all of it. Nine layouts across three templates rendered
    pixel-identical without them."""
    from pptx import Presentation

    from reportbuilder.render.deck import _strip_slides
    from reportbuilder.render.image.fast_preview import keep_only_used_layout

    if not pathlib.Path(path).exists():
        pytest.skip(f"{path} not available locally")
    prs = Presentation(path)
    _strip_slides(prs)
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    whole = tmp_path / "whole.pptx"
    prs.save(str(whole))

    keep_only_used_layout(prs, slide)
    lean = tmp_path / "lean.pptx"
    prs.save(str(lean))

    assert len(prs.slide_masters) == 1
    assert len(prs.slide_masters[0].slide_layouts) == 1
    assert slide.slide_layout is prs.slide_masters[0].slide_layouts[0], (
        "the layout kept is the one the slide is drawn on")
    assert lean.stat().st_size < whole.stat().st_size / 2, (
        f"{lean.stat().st_size / 1e6:.2f} MB against {whole.stat().st_size / 1e6:.2f} MB")


def test_a_broken_deck_is_left_whole(template):
    """Reducing is an optimisation: anything unexpected keeps the full deck,
    because a ground drawn slowly is right and one drawn from a deck we broke
    is not."""
    from pptx import Presentation

    from reportbuilder.render.deck import _strip_slides
    from reportbuilder.render.image.fast_preview import keep_only_used_layout

    prs = Presentation(template)
    _strip_slides(prs)
    prs.slides.add_slide(prs.slide_layouts[4])
    masters_before = len(prs.slide_masters)
    keep_only_used_layout(prs, object())          # not a slide at all
    assert len(prs.slide_masters) == masters_before
