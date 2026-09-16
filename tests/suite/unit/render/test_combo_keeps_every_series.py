"""A combo chart draws every series it was given.

"Combo chartissa category variable ei tule mukaan kuvaan vaikka sen valitsee."

`build_image_combo` read exactly `segs[0]` as the bars and `segs[1]` as the
line, and never looked at the rest. Two defects fell out of that one shape:

* With a classifying variable AND a secondary variable, `_combo_two_var` built
  its bars with `classifying_var=None` — the classifier was discarded before it
  reached the renderer, and the result was byte-identical with and without it.
* With a classifying variable alone, every group after the second was dropped.
  On the reported slide the classifier was Sukupuoli with four groups, so
  `Muu` and `En halua vastata` were simply absent, with nothing to say so.

The rule now: the SECONDARY variable is the line, and the question is the bars,
one bar series per classifier group. With no secondary chosen, every group after
the first is a line — which keeps the existing two-group chart (group 1 bars,
group 2 line) exactly as it was, and stops the third group from vanishing.
(Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pandas as pd
import pytest
from matplotlib.container import BarContainer
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.question import (
    Question, QuestionModel, ValueLabel, Variable,
)
from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats import engine
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("1- Ei lainkaan ylpeä", "2", "3")
#: The reported slide's classifier: four groups, of which two were disappearing.
_GROUPS = ("Mies", "Nainen", "Muu", "En halua vastata")
_INDEX = "ikaluokka"


# ---- the engine keeps the classifier ---------------------------------------

def _model_and_data():
    pride = Variable("q", "Kuinka ylpeä olet omasta työstäsi?", "categorical",
                     tuple(ValueLabel(float(i + 1), c) for i, c in enumerate(_CATS)),
                     frozenset())
    sex = Variable("sp", "Sukupuolesi", "categorical",
                   tuple(ValueLabel(float(i + 1), g) for i, g in enumerate(_GROUPS)),
                   frozenset())
    idx = Variable("ika", _INDEX, "scale", (), frozenset())
    model = QuestionModel(variables={"q": pride, "sp": sex, "ika": idx}, questions=[])
    q = Question(qid="q", kind="single", variables=("q",), text=pride.label)
    n = 240
    df = pd.DataFrame({
        "q": [float(i % 3 + 1) for i in range(n)],
        "sp": [float(i % 4 + 1) for i in range(n)],
        "ika": [float(30 + i % 20) for i in range(n)],
    })
    return model, q, df


def _spec(**kw) -> ChartSpec:
    base = dict(question_ref="q", chart_type="combo", statistic="pct",
                classifying_var=None, number_format=NumberFormat(),
                sort=SortSpec(basis="data_order"), template_slot="s1",
                elements=ElementToggles(), options={})
    base.update(kw)
    return ChartSpec(**base)


def _computed(**kw) -> SeriesResult:
    model, q, df = _model_and_data()
    return engine.compute(q, _spec(**kw), df, model)


def test_the_classifier_reaches_the_chart():
    """The report: selecting it changed nothing at all."""
    r = _computed(classifying_var="sp", options={"combo_secondary": "ika"})
    for g in _GROUPS:
        assert g in r.segments, f"{g} is missing from {r.segments}"


def test_the_secondary_is_still_the_line():
    r = _computed(classifying_var="sp", options={"combo_secondary": "ika"})
    assert _INDEX in r.segments
    assert r.statistic_of(_INDEX) == "mean"
    for g in _GROUPS:
        # Membership first: `statistic_of` answers for a segment that is not
        # there by falling back to the series statistic, which is "pct" — so
        # without this the loop below passes while the group is missing.
        assert g in r.segments, f"{g} is missing from {r.segments}"
        assert r.statistic_of(g) == "pct"


def test_selecting_a_classifier_actually_changes_the_result():
    """The sharpest form of the bug: with and without were identical."""
    with_it = _computed(classifying_var="sp", options={"combo_secondary": "ika"})
    without = _computed(options={"combo_secondary": "ika"})
    assert with_it.segments != without.segments


def test_no_classifier_is_unchanged():
    """The plain two-variable combo keeps its exact shape."""
    r = _computed(options={"combo_secondary": "ika"})
    assert len(r.segments) == 2
    assert r.segments[1] == _INDEX
    assert r.statistic_of(r.segments[1]) == "mean"


# ---- the renderer draws all of them ----------------------------------------

def _series(groups: tuple[str, ...], *, secondary: bool) -> SeriesResult:
    segs = groups + ((_INDEX,) if secondary else ())
    cells = {}
    for i, c in enumerate(_CATS):
        for j, g in enumerate(groups):
            cells[(c, g)] = Cell(pct=float(20 + i * 5 + j * 3))
        if secondary:
            cells[(c, _INDEX)] = Cell(pct=float(34 + i))
    stats = {g: "pct" for g in groups}
    if secondary:
        stats[_INDEX] = "mean"
    return SeriesResult(
        categories=_CATS, segments=segs, cells=cells,
        base_n={"Total": 1018, **{s: 250 for s in segs}},
        statistic="pct", segment_statistics=stats)


def _render(series: SeriesResult):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(),
        spec=_spec(elements=ElementToggles(title=True, subtitle=True, legend=True,
                                           n=True, axis_names=True, filter_var=True,
                                           data_labels=True)),
        series=series, fmt=NumberFormat())
    drawn = {}

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        bar_series, line_series = 0, 0
        for axes in self.axes:
            bar_series += sum(1 for c in axes.containers
                              if isinstance(c, BarContainer))
            line_series += sum(1 for ln in axes.lines
                               if len(ln.get_xdata()) == len(_CATS))
        drawn["bars"], drawn["lines"] = bar_series, line_series
        drawn["legend"] = [t.get_text() for lg in self.legends for t in lg.get_texts()]
        for axes in self.axes:
            if axes.get_legend() is not None:
                drawn["legend"] += [t.get_text()
                                    for t in axes.get_legend().get_texts()]
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        _f.Figure.savefig = original
    return drawn


def test_every_classifier_group_is_drawn_as_its_own_bar_series():
    drawn = _render(_series(_GROUPS, secondary=True))
    assert drawn["bars"] == len(_GROUPS), (
        f"{len(_GROUPS)} groups were given and {drawn['bars']} bar series drawn")
    assert drawn["lines"] == 1, "the secondary variable is the one line"


def test_the_legend_names_the_groups():
    """With one bar series the legend names only the line, because the bars are
    the question and the subtitle already says so. With several, the reader has
    no other way to tell the groups apart."""
    drawn = _render(_series(_GROUPS, secondary=True))
    # Substring, not equality: a group states its own base ("Mies (n=250)") the
    # way every other chart's groups do — see `elements.group_base`. What this
    # test is about is that each group is NAMED, not how it is spelled.
    # (2026-09-16)
    for g in _GROUPS:
        assert any(g in entry for entry in drawn["legend"]), (
            f"{g} missing from legend {drawn['legend']}")


def test_no_group_is_dropped_when_there_is_no_secondary():
    """The second defect: groups 3 and 4 were drawn nowhere at all."""
    drawn = _render(_series(_GROUPS, secondary=False))
    assert drawn["bars"] + drawn["lines"] == len(_GROUPS), (
        f"{len(_GROUPS)} groups given, {drawn['bars']} bars + {drawn['lines']} "
        f"lines drawn — the rest vanished silently")


def test_the_two_group_chart_is_unchanged():
    """One bars, one line — exactly what it does today."""
    drawn = _render(_series(("Nainen", "Mies"), secondary=False))
    assert (drawn["bars"], drawn["lines"]) == (1, 1)


def test_the_plain_two_variable_combo_is_unchanged():
    drawn = _render(_series(("Kuinka ylpeä",), secondary=True))
    assert (drawn["bars"], drawn["lines"]) == (1, 1)
