"""Measuring a chart's text does not paint the chart.

Every number a stacked bar printed was measured with a full `canvas.draw()` —
every artist rasterised to read one text's width — and the label fitters drew
the whole figure again each time they wanted to look. A 14-statement battery
split four ways painted itself 76 times: 35 s of a preview on one core. The
measurements now settle the axes (`settle_axes`) or lay the figure out without
rasterising it (`Figure.draw_without_rendering`); only the saved picture is
painted. Pixel-identical on 259 charts of the visual QA matrix. (perf, 2026-09-19)
"""
from __future__ import annotations

import traceback

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.backends.backend_agg import FigureCanvasAgg
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_STATEMENTS = tuple(f"Väittämä numero {i}" for i in range(1, 15))
_SCALE = ("1", "2", "3", "4", "5")


@pytest.mark.parametrize("chart_type", ["stacked_horizontal_bar", "stacked_vertical_bar",
                                        "vertical_bar", "line", "pie"])
def test_a_chart_is_painted_once(chart_type, monkeypatch):
    if chart_type == "pie":
        cats, segs = _SCALE, ("Mieheksi", "Naiseksi", "Muuksi")
    elif chart_type.startswith("stacked"):
        cats, segs = _SCALE, tuple(f"{s} · {g}" for s in _STATEMENTS[:6]
                                   for g in ("Mieheksi", "Naiseksi"))
    else:
        cats, segs = _STATEMENTS, ("Mieheksi", "Naiseksi", "Muuksi")
    cells = {(c, g): Cell(pct=100.0 / len(cats)) for c in cats for g in segs}
    series = SeriesResult(categories=cats, segments=segs, cells=cells,
                          base_n={g: 300 for g in segs}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    painted = []
    real = FigureCanvasAgg.draw

    def counting(self, *a, **k):
        # Matplotlib obtains a renderer for `draw_without_rendering` (and for a
        # tight bounding box) by STARTING a canvas draw and aborting it before
        # anything is painted — `_get_renderer`. Those are not paints.
        if not any(f.name == "_get_renderer" for f in traceback.extract_stack()):
            painted.append(1)
        return real(self, *a, **k)

    monkeypatch.setattr(FigureCanvasAgg, "draw", counting)
    IMAGE_BUILDERS[chart_type](ctx)
    # The saved picture itself, and at most one more.
    assert len(painted) <= 2, f"{chart_type} painted itself {len(painted)} times"
