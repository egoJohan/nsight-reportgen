"""Every number drawn on a chart says that it is a number.

A `gid` is how anything downstream tells one kind of text from another, and two
things now need that:

* Telling "a number sits on its own bar" (expected) from "a number sits on
  ANOTHER bar" (a defect) — the first cannot be judged without knowing which
  texts are values.
* Choosing which font floor applies. The renderers enforce five different ones
  (7.5pt names, 6.5pt the stacked shrink floor, 5.5pt clustered-horizontal
  values), so "is this too small?" has no answer until you know what it is.

Today only `bars.py` tags anything, and only on its stacked paths: measured
through the real preview endpoint, a clustered horizontal bar reports
`gids={None: 6}`. The clustered paths (`bars.py:1250`, `:1376`), `line.py:61`
and `funnel.py:74,89` draw their numbers untagged.

Value labels are found by SHAPE here, not by matching formatted strings: the
formatter's space character is an implementation detail, and a near-miss would
make this pass while finding nothing. Categories are words and the value axis
ticks carry no percent sign, so the pattern separates them cleanly.
(Johan, 2026-09-12)
"""
from __future__ import annotations

import re

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image.bars import _VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

#: "60 %", "12,5 %", "7%" — a number, however it is spaced or decimalised.
_LOOKS_LIKE_A_VALUE = re.compile(r"^\s*-?\d+(?:[.,]\d+)?\s*%\s*$")


def _two_segment_series() -> SeriesResult:
    """Two series over four categories — what makes a bar chart CLUSTERED,
    which is the path that draws its numbers untagged."""
    cats = ("Toukokuu 2024", "Marraskuu 2024", "Toukokuu 2025", "Marraskuu 2025")
    segs = ("Attendo", "Esperi")
    cells = {(c, s): Cell(pct=float(60 + j * 10 + i * 2))
             for j, s in enumerate(segs) for i, c in enumerate(cats)}
    return SeriesResult(categories=cats, segments=segs, cells=cells,
                        base_n={"Attendo": 1001, "Esperi": 1001}, statistic="pct")


def _funnel_series() -> SeriesResult:
    cats = ("Spontaneous", "Aided", "Consideration", "Trial", "Loyalty")
    vals = (80.0, 60.0, 40.0, 25.0, 12.0)
    return SeriesResult(
        categories=cats, segments=("Total",),
        cells={(c, "Total"): Cell(pct=v) for c, v in zip(cats, vals)},
        base_n={"Total": 500}, statistic="pct")


@pytest.fixture
def drawn(monkeypatch):
    """Every text on the figure as it goes to disk, with its gid."""
    seen: list[tuple[str, str | None]] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        for ax in self.axes:
            for t in ax.texts:
                if t.get_visible() and t.get_text().strip():
                    seen.append((t.get_text(), t.get_gid()))
        return original(self, *a, **k)

    monkeypatch.setattr(Figure, "savefig", spy)
    return seen


def _render(chart_type: str, series: SeriesResult) -> None:
    spec = ChartSpec(
        question_ref="q1", chart_type=chart_type, statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles(title=True, data_labels=True, legend=True, n=True),
    )
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS[chart_type](ctx)


def _pie_series() -> SeriesResult:
    """One series carrying a SLIVER, so both of pie's paths run: the
    percentages matplotlib draws inside the wedges, and the callout a wedge
    too thin to hold its own number gets instead (2 % is under the 4 % floor)."""
    cats = ("Kyllä", "Ei", "En osaa sanoa", "Ei vastausta")
    vals = (55.0, 31.0, 12.0, 2.0)
    return SeriesResult(
        categories=cats, segments=("Total",),
        cells={(c, "Total"): Cell(pct=v) for c, v in zip(cats, vals)},
        base_n={"Total": 500}, statistic="pct")


# Not scatter: `scatter.py:54` annotates each point with its CATEGORY NAME,
# not a value, so it has no value labels to tag and tagging it would be wrong.
@pytest.mark.parametrize("chart_type,series", [
    ("vertical_bar", _two_segment_series()),
    ("horizontal_bar", _two_segment_series()),
    ("line", _two_segment_series()),
    ("funnel", _funnel_series()),
    ("combo", _two_segment_series()),
    ("pie", _pie_series()),
], ids=["clustered-column", "clustered-bar", "line", "funnel", "combo", "pie"])
def test_every_number_it_draws_is_tagged_as_a_value(drawn, chart_type, series):
    _render(chart_type, series)

    numbers = [(text, gid) for text, gid in drawn if _LOOKS_LIKE_A_VALUE.match(text)]

    # Without this the test passes by finding nothing at all.
    assert numbers, (
        f"{chart_type} drew no numbers to check; the fixture or the matcher is "
        f"wrong, not the product: {[t for t, _g in drawn][:8]}")
    untagged = [text for text, gid in numbers if gid != _VALUE_GID]
    assert not untagged, (
        f"{chart_type} draws {len(untagged)} of its {len(numbers)} numbers with "
        f"no gid, so nothing downstream can tell them from any other text: "
        f"{untagged[:6]}")
