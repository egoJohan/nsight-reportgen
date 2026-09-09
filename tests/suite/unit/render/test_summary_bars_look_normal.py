"""A mean chart draws like any other vertical bar.

One value per group is a bar chart whose CATEGORIES are the groups. The engine
reports it the other way round — one category (the measure) × N segments — and
the renderer took that literally: four segments split a single category slot, so
the bars came out flush against each other and, at one category, filling most of
the axis. Four fat blocks in four different colours, reported as "palkit ovat
tosi leveitä kuvaajan kokoon nähden … ne saavat olla irti toisistaan niin kuin
normaaleissa vertical bar kuvaajissa". (Johan, 2026-09-09)

So it is transposed before drawing: groups become the categories and there is
one series. Everything else then follows from the ordinary path — the same bar
width, the same gaps, the same single series colour, the group names on the
axis, and no legend, because a chart with one series has nothing to explain.
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.stats.engine import compute
import reportbuilder.render.image.bars as B

_GROUPS = ["Suomi", "Ruotsi", "Saksa"]


def _drawn(statistic="mean", chart_type="vertical_bar"):
    """Render and report the bars matplotlib actually made."""
    measure = Variable(name="TeknologianVaikutukset", label="TeknologianVaikutukset",
                       measurement="scale", value_labels=[], missing_values=[])
    opts = Variable(name="mielipide", label="Mielipide", measurement="nominal",
                    missing_values=[],
                    value_labels=[ValueLabel(value=float(i + 1), label=f"Vaihtoehto {i+1}")
                                  for i in range(4)])
    maa = Variable(name="maa", label="Maa", measurement="nominal", missing_values=[],
                   value_labels=[ValueLabel(value=float(i + 1), label=g)
                                 for i, g in enumerate(_GROUPS)])
    rows = [{"TeknologianVaikutukset": 6.0 + (i % 5) * 0.3,
             "mielipide": float(i % 4 + 1), "maa": float(i % 3 + 1)}
            for i in range(300)]
    qid = "tek" if statistic == "mean" else "mielipide"
    model = QuestionModel(
        variables={"TeknologianVaikutukset": measure, "mielipide": opts, "maa": maa},
        questions=[Question(qid="tek", text="TeknologianVaikutukset", kind="single",
                            variables=("TeknologianVaikutukset",)),
                   Question(qid="mielipide", text="Mielipide", kind="single",
                            variables=("mielipide",)),
                   Question(qid="maa", text="Maa", kind="single", variables=("maa",))])
    spec = ChartSpec(question_ref=qid, chart_type=chart_type, statistic=statistic,
                     classifying_var="maa", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question(qid), spec, pd.DataFrame(rows), model)
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(1), top=Inches(1),
                                  width=Inches(9), height=Inches(4.5), name="s1"),
                        style=StyleSpec(), spec=spec, series=series,
                        fmt=spec.number_format)
    caught: dict = {}
    original = B.render_png

    def _spy(fig):
        ax = fig.axes[0]
        rects = [p for p in ax.patches if getattr(p, "get_width", None)]
        caught["bars"] = [(p.get_x(), p.get_width(), p.get_facecolor()) for p in rects]
        caught["xticks"] = [t.get_text() for t in ax.get_xticklabels() if t.get_text()]
        caught["yticks"] = [t.get_text() for t in ax.get_yticklabels() if t.get_text()]
        caught["legend"] = ax.get_legend() is not None or bool(fig.legends)
        caught["n_categories"] = len(series.categories)
        caught["segments"] = list(series.segments)
        return original(fig)

    B.render_png = _spy
    try:
        (B.build_image_column if chart_type == "vertical_bar"
         else B.build_image_bar)(ctx)
    finally:
        B.render_png = original
    return caught


def test_the_engine_still_reports_one_category_times_groups():
    """The transposition is the RENDERER's; nothing upstream changes."""
    got = _drawn()
    assert got["n_categories"] == 1
    assert set(_GROUPS) <= set(got["segments"])


def test_there_is_one_bar_per_group():
    got = _drawn()
    assert len(got["bars"]) == len(got["segments"])


def test_the_bars_do_not_touch():
    got = _drawn()
    edges = sorted((x, x + w) for x, w, _c in got["bars"])
    gaps = [nxt[0] - cur[1] for cur, nxt in zip(edges, edges[1:])]
    assert gaps and min(gaps) > 0.05, f"bars are flush or overlapping: {gaps}"


def test_the_bars_are_the_ordinary_width():
    """0.5 of a category slot — what a single-series vertical bar always uses.
    They were 0.175 and flush, four of them sharing one slot."""
    got = _drawn()
    assert all(w == pytest.approx(0.5, abs=0.01) for _x, w, _c in got["bars"])


def test_every_bar_is_the_same_colour():
    """A vertical bar of one series draws one colour. Four groups in four
    different colours is what made this read as a different kind of chart."""
    got = _drawn()
    assert len({c for _x, _w, c in got["bars"]}) == 1


def test_the_groups_are_named_on_the_axis():
    got = _drawn()
    assert set(_GROUPS) <= set(got["xticks"]), got["xticks"]


def test_no_legend_is_drawn_for_one_series():
    assert _drawn()["legend"] is False


def test_a_distribution_is_untouched():
    """Only a summary chart is one-value-per-group. A pct chart keeps its
    categories, its grouped bars and its legend."""
    got = _drawn(statistic="pct")
    assert got["n_categories"] == 4
    # one bar per category per DRAWN series (a Total column is not always drawn)
    assert len(got["bars"]) > 4 and len(got["bars"]) % 4 == 0
    assert got["legend"] is True


def test_the_horizontal_form_is_transposed_the_same_way():
    """Whichever way the bars run, one value per group is a chart of groups."""
    got = _drawn(chart_type="horizontal_bar")
    assert set(_GROUPS) <= set(got["yticks"]), got["yticks"]
    assert len({c for _x, _w, c in got["bars"]}) == 1
