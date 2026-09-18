"""The accent bar beside a heading is the client's colour, not ours.

`theme_colours` exists to answer exactly this — "(background, ink, accent) for
a slide, preferring the template's own", with the reason in its own docstring:
"A deck built on the client's template should not carry nSight's cream ground
and teal accent bar". A chart slide honours it.

The demographics grid and the special slides (overview, conclusion, blank) call
it, take the background and the ink, and then discard the accent and paint the
bar house teal. So one deck had the client's colour beside its chart headings
and nSight's green beside its section headings.

Only where the template does NOT own the slide ground: a template that supplies
its own furniture draws no bar of ours at all, and that path is untouched.
"""
from __future__ import annotations

import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import Slot, StyleSpec
from reportbuilder.render.house_style import PX_TEAL

#: A brand nobody could mistake for house teal.
BRAND_HEX = "B3005E"


def _style(*, branded: bool) -> StyleSpec:
    style = StyleSpec()
    if branded:
        style.from_template = True
        style.brand_palette = [BRAND_HEX]
        style.accent = BRAND_HEX
    return style


def _spec(chart_type: str) -> ChartSpec:
    return ChartSpec(question_ref="", chart_type=chart_type, statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(), slide_title="Otsikko",
                     options={"bullets": ["Eka", "Toka"], "charts": []})


def _bar_fills(draw) -> list[str]:
    """Every solid-filled shape's colour on the drawn slide, as hex."""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slot = Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                width=Inches(11.6), height=Inches(4.0), name="s1")
    draw(slide, slot)
    out = []
    for sh in slide.shapes:
        try:
            if sh.fill.type is not None and sh.fill.fore_color.rgb is not None:
                out.append(str(sh.fill.fore_color.rgb).upper())
        except Exception:
            continue
    return out


def _special(style):
    from reportbuilder.render.image.special_slide import render_special_slide

    return lambda slide, slot: render_special_slide(slide, slot, style, _spec("special_conclusion"))


def _grid(style):
    from reportbuilder.render.image.demographics_grid import render_demographics_grid

    return lambda slide, slot: render_demographics_grid(
        slide, slot, style, _spec("demographics_grid"), {}, {})


@pytest.mark.parametrize("make", [_special, _grid], ids=["special-slide", "demographics-grid"])
def test_a_branded_deck_gets_the_brands_accent_bar(make):
    fills = _bar_fills(make(_style(branded=True)))
    assert BRAND_HEX in fills, f"no {BRAND_HEX} bar; drew {fills}"


@pytest.mark.parametrize("make", [_special, _grid], ids=["special-slide", "demographics-grid"])
def test_a_branded_deck_has_no_house_teal_bar(make):
    fills = _bar_fills(make(_style(branded=True)))
    assert str(PX_TEAL).upper() not in fills, f"still drew house teal; {fills}"


@pytest.mark.parametrize("make", [_special, _grid], ids=["special-slide", "demographics-grid"])
def test_an_untemplated_deck_is_still_house_teal(make):
    """The house deck must not change."""
    fills = _bar_fills(make(_style(branded=False)))
    assert str(PX_TEAL).upper() in fills, f"lost the house accent; {fills}"
