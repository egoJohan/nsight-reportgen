"""A chart is placed at the size it was drawn, or smaller — never bigger.

"Slide 6 — why is the chart font so big? There should be a limited maximum. It
is now bigger than the question subtitle."

Every font in a chart is chosen in POINTS: the row labels clamp at 11.5, the
value labels at 9.5, against a subtitle of 13.8. Those numbers were true of the
figure and not of the slide. `savefig(bbox_inches="tight")` trims the drawing to
its ink, so a chart with six short labels and bars reaching 25 % saves a much
narrower PNG than the figure it came from — and `place_picture_square` then
scaled that PNG to FILL the slot, magnifying every point size with it. Measured
on the reported slide: 11.5pt row labels arriving at roughly 17pt, larger than
the subtitle above them.

So the scale is capped at 1:1. A chart too big for its slot still shrinks — that
is what letterboxing is for — but one that came out small is left at the size it
was drawn, and a point size means on the slide what it says in the builder.
(Johan, 2026-09-17)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

#: What one rendered pixel is worth on a slide. `render_png` saves at 200 dpi,
#: so this is the size the figure was DRAWN at — the ceiling on placement.
EMU_PER_PX = 914400 / 200


def _placed(n_cats: int, label: str = "Amazon") -> tuple[int, int, int, int]:
    """(placed_w, placed_h, png_w_px, png_h_px) for a horizontal bar."""
    cats = tuple(f"{label} {i}" if n_cats > 1 else label for i in range(n_cats))
    series = SeriesResult(
        categories=cats, segments=("Total",),
        cells={(c, "Total"): Cell(pct=25.0 - i * 0.5) for i, c in enumerate(cats)},
        base_n={"Total": 1018}, statistic="pct")
    spec = ChartSpec(
        question_ref="q", chart_type="horizontal_bar", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    IMAGE_BUILDERS["horizontal_bar"](ctx)
    pic = next(sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    return pic.width, pic.height, pic.image.size[0], pic.image.size[1]


def test_a_short_chart_is_left_at_the_size_it_was_drawn():
    """The reported slide: six categories, bars to 25 %, a PNG much narrower
    than its figure — which used to be stretched across the whole slot."""
    w, _h, px_w, _px_h = _placed(6)
    assert w <= px_w * EMU_PER_PX + 2, (
        f"placed at {w / px_w:.0f} EMU/px against the {EMU_PER_PX:.0f} it was "
        f"drawn at — every point size on this chart is "
        f"{w / px_w / EMU_PER_PX:.2f}x what it says")


def test_a_tall_chart_still_shrinks_to_fit():
    """The cap is a ceiling, not a fixed size. A chart bigger than its slot must
    still letterbox down, or it would overflow the slide."""
    w, h, _px_w, _px_h = _placed(30, label="Pitkä nimi jonka pituus vaihtelee")
    assert w <= Inches(12.3) and h <= Inches(4.4)


def test_the_aspect_ratio_is_kept():
    """Whatever the scale, it is one scale — a chart is never stretched."""
    w, h, px_w, px_h = _placed(6)
    assert abs((w / h) - (px_w / px_h)) < 0.02
