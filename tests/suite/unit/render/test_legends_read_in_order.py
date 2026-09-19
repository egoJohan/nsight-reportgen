"""A legend of several columns reads, row by row, in its entries' own order.

Matplotlib fills a legend's columns first. Five categories in four columns put
the second alone on the second row, so a pie row's scale read "Erittäin huono,
Hyvä, Erittäin hyvä, En osaa sanoa … Huono", and the brand radar's eight
brands came out "Attendo, Mainio-kodit, Ykköskodit, Rinnekodit / Esperi, …".
(visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_SCALE = ("Erittäin huono", "Huono", "Hyvä", "Erittäin hyvä", "En osaa sanoa")
_BRANDS = ("Attendo", "Esperi", "Mainio-kodit", "Onnikodit", "Ykköskodit",
           "Humana", "Rinnekodit", "Validia")
_STATEMENTS = ("Luotettava", "Välittävä", "Ahne", "Kodikas", "Inhimillinen")


def _legend_reading(chart_type: str, cats, segs) -> list[str]:
    cells = {(c, g): Cell(pct=10.0 + i + 3 * j)
             for i, c in enumerate(cats) for j, g in enumerate(segs)}
    series = SeriesResult(categories=cats, segments=segs, cells=cells,
                          base_n={g: 100 for g in segs}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="g" if len(segs) > 1 else None,
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(group_base=False))
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    got: list = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        legends = list(self.legends) + [ax.get_legend() for ax in self.axes
                                        if ax.get_legend() is not None]
        for leg in legends:
            boxes = [(t.get_window_extent(r), t.get_text()) for t in leg.get_texts()]
            # Top row first, then left to right within a row.
            boxes.sort(key=lambda bt: (-round(bt[0].y1 / 4), bt[0].x0))
            got.append([text.replace("\n", " ") for _b, text in boxes])
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        _f.Figure.savefig = original
    assert got, "no legend was drawn"
    return got[0]


def test_a_row_of_pies_lists_its_scale_in_order():
    groups = ("Mieheksi", "Naiseksi", "Muuksi")
    assert _legend_reading("pie", _SCALE, groups) == list(_SCALE)


def test_the_brand_radar_lists_its_brands_in_order():
    assert _legend_reading("radar", _STATEMENTS, _BRANDS) == list(_BRANDS)


def test_a_grouped_horizontal_bar_lists_its_groups_top_down():
    """A horizontal chart draws each category's bars bottom-up, so a legend in
    plotting order listed "Mieheksi" first and drew it last. It now names the
    groups in the order the reader meets them. (visual QA, 2026-09-19)"""
    groups = ("Mieheksi", "Naiseksi", "Muuksi")
    cats = ("Attendo", "Esperi", "Humana")
    cells = {(c, g): Cell(pct=20.0 + 5 * i + j) for i, c in enumerate(cats)
             for j, g in enumerate(groups)}
    series = SeriesResult(categories=cats, segments=groups, cells=cells,
                          base_n={g: 100 for g in groups}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type="horizontal_bar", statistic="pct",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(group_base=False))
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    got: dict = {}

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        ax = self.axes[0]
        tops = sorted(((c[0].get_y(), c.get_label()) for c in ax.containers), reverse=True)
        got["bars"] = [label for _y, label in tops]
        got["legend"] = [t.get_text() for t in ax.get_legend().get_texts()]
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["horizontal_bar"](ctx)
    finally:
        _f.Figure.savefig = original
    assert got["legend"] == got["bars"], got
