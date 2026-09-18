"""A single-series chart is drawn in the CLIENT's colour, whatever its type.

On a branded template the bar, line, pie and radar builders take their colour
from the template's accents; the funnel drew every stage in a hard-coded house
TEAL, and the combo fell back to TEAL whenever it had only one series. So one
deck, one template, one question showed the client's blue as a bar and nSight's
teal as a funnel — the reader's conclusion being that the colour means
something.

The fix is not a new colour rule, it is the SAME rule: `series_colors(1, ...)`
already answers "the one colour for a single-series chart", and answers TEAL
when no template applies, so a house deck is unchanged.
"""
from __future__ import annotations

import pandas as pd
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.ingest.sav_reader import ValueLabel, Variable
from reportbuilder.model.question import Question, QuestionModel
from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.house_style import TEAL
from reportbuilder.stats.engine import compute

#: A brand nobody could mistake for the house teal or for Office's defaults.
BRAND = ["B3005E", "2E1A47", "C8A200", "5E7C16", "8C3B00", "00666E"]
_CATS = ["Spontaneous", "Aided", "Consideration", "Trial", "Loyalty"]


def _ctx(chart_type: str, *, branded: bool):
    var = Variable(name="q", label="Tunnettuus", measurement="nominal",
                   missing_values=[],
                   value_labels=[ValueLabel(value=float(i + 1), label=c)
                                 for i, c in enumerate(_CATS)])
    counts = [80, 60, 40, 25, 12]
    rows = [{"q": float(i + 1)} for i, n in enumerate(counts) for _ in range(n)]
    model = QuestionModel(
        variables={"q": var},
        questions=[Question(qid="q", text=var.label, kind="single", variables=("q",))])
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    series = compute(model.question("q"), spec, pd.DataFrame(rows), model)
    style = StyleSpec()
    if branded:
        style.from_template = True
        style.brand_palette = list(BRAND)
        style.accent = BRAND[0]
    prs = Presentation()
    return RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                         slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                   width=Inches(11.6), height=Inches(4.0), name="s1"),
                         style=style, spec=spec, series=series,
                         fmt=spec.number_format)


def _pixels(chart_type: str, builder_mod: str, builder: str, *, branded: bool):
    """How many PIXELS the finished chart image has of each colour.

    Counted, not collected: a set of distinct colours says a colour is present
    once whether it is one antialiased pixel or a whole bar, and "present" is
    not the question — "drawn in" is.
    """
    import importlib
    from collections import Counter

    from PIL import Image

    mod = importlib.import_module(f"reportbuilder.render.image.{builder_mod}")
    seen: dict = {}
    original = mod.render_png

    def _spy(fig):
        # render_png hands back a PATH, not the bytes.
        path = original(fig)
        with Image.open(path) as img:
            seen["colours"] = Counter(img.convert("RGB").getdata())
        return path

    mod.render_png = _spy
    try:
        getattr(mod, builder)(_ctx(chart_type, branded=branded))
    finally:
        mod.render_png = original
    return seen["colours"]


def _rgb(hex6: str) -> tuple[int, int, int]:
    h = hex6.lstrip("#")
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


#: A chart DRAWN in a colour covers real area. Antialiasing along an edge of
#: some other colour can put a handful of near-matching pixels anywhere, so
#: "was this chart drawn in X" has to mean area, not presence.
_AREA_PX = 400


def _near(colours, target: str, tol: int = 6, least: int = _AREA_PX) -> bool:
    t = _rgb(target)
    return sum(n for p, n in colours.items()
               if all(abs(p[i] - t[i]) <= tol for i in range(3))) >= least


SINGLE_SERIES = [
    ("funnel", "funnel", "build_image_funnel"),
    ("combo", "combo", "build_image_combo"),
    ("vertical_bar", "bars", "build_image_column"),
    ("horizontal_bar", "bars", "build_image_bar"),
    ("line", "line", "build_image_line"),
]


@pytest.mark.parametrize("chart_type,mod,builder", SINGLE_SERIES)
def test_a_branded_template_gets_the_brands_colour(chart_type, mod, builder):
    colours = _pixels(chart_type, mod, builder, branded=True)
    assert _near(colours, BRAND[0]), f"{chart_type} drew no {BRAND[0]}"


@pytest.mark.parametrize("chart_type,mod,builder", SINGLE_SERIES)
def test_a_branded_template_shows_no_house_teal(chart_type, mod, builder):
    colours = _pixels(chart_type, mod, builder, branded=True)
    assert not _near(colours, TEAL), f"{chart_type} still draws house teal"


@pytest.mark.parametrize("chart_type,mod,builder", SINGLE_SERIES)
def test_no_template_is_still_house_teal(chart_type, mod, builder):
    """The house deck must not change; this is what makes the fix safe."""
    colours = _pixels(chart_type, mod, builder, branded=False)
    assert _near(colours, TEAL), f"{chart_type} lost the house teal"


# ---------------------------------------------------------------------------
# Scatter, which needs two segments and so has its own fixture
# ---------------------------------------------------------------------------

def _scatter_pixels(*, branded: bool):
    """A scatter of attributes, two groups on the two axes."""
    from collections import Counter

    from PIL import Image

    import reportbuilder.render.image.scatter as mod
    from reportbuilder.stats.series import Cell, SeriesResult

    cats = ("Luotettava", "Inhimillinen", "Ahne", "Kodikas")
    segs = ("Mieheksi", "Naiseksi")
    cells = {(c, s): Cell(mean=3.0 + 0.1 * i + (0.2 if s == segs[1] else 0.0))
             for i, c in enumerate(cats) for s in segs}
    series = SeriesResult(categories=cats, segments=segs, cells=cells,
                          base_n={s: 400 for s in segs}, statistic="mean")
    spec = ChartSpec(question_ref="q", chart_type="scatter", statistic="mean",
                     classifying_var="g", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(), scatter_xy=segs)
    style = StyleSpec()
    if branded:
        style.from_template = True
        style.brand_palette = list(BRAND)
        style.accent = BRAND[0]
    prs = Presentation()
    ctx = RenderContext(slide=prs.slides.add_slide(prs.slide_layouts[6]),
                        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                                  width=Inches(11.6), height=Inches(4.0), name="s1"),
                        style=style, spec=spec, series=series, fmt=spec.number_format)
    seen: dict = {}
    original = mod.render_png

    def _spy(fig):
        path = original(fig)
        with Image.open(path) as img:
            seen["colours"] = Counter(img.convert("RGB").getdata())
        return path

    mod.render_png = _spy
    try:
        mod.build_image_scatter(ctx)
    finally:
        mod.render_png = original
    return seen["colours"]


def test_a_branded_scatter_draws_the_brands_colour():
    """The dots are the data. On a client template they were house teal while
    everything around them — typeface, title, footer — was the client's."""
    assert _near(_scatter_pixels(branded=True), BRAND[0], least=200)


def test_a_branded_scatter_shows_no_house_teal():
    assert not _near(_scatter_pixels(branded=True), TEAL, least=200)


def test_an_untemplated_scatter_is_still_house_teal():
    assert _near(_scatter_pixels(branded=False), TEAL, least=200)
