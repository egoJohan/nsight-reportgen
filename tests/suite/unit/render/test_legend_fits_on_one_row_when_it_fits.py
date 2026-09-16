"""A legend that fits on one row is drawn on one row.

"Slide 15 — why didn't legend fit to single line although there would have been
space?"

`_legend_below` decided by COUNT: one row only for five entries or fewer (seven
if every label is a bare scale number). A six-entry legend of short words —
Amazon, Salesforce, Aramco, nSight, Estrella, En osaa sanoa — was wrapped onto
two rows across the full width of a slide that had room for all six.

The codebase already learned this about value labels, in `_value_label_layout`:
"A count cannot tell a roomy chart from a cramped one — the same five groups are
comfortable on a wide slide with three categories and hopeless on a narrow one
with twelve." The legend was still counting.

So it measures. The numeric-scale shortening stays exactly as it was: that is
about what a legend SAYS, not how wide it is. (Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

SHORT = ("Amazon", "Salesforce", "Aramco", "nSight", "Estrella", "En osaa sanoa")
LONG = tuple(f"Erittäin pitkä vastausvaihtoehto numero {i}" for i in range(1, 8))


def _series(levels) -> SeriesResult:
    cells = {(l, "Total"): Cell(pct=100.0 / len(levels)) for l in levels}
    return SeriesResult(categories=tuple(levels), segments=("Total",), cells=cells,
                        base_n={"Total": 1018}, statistic="pct")


def _legend_rows(levels, width_in: float = 12.3) -> int:
    spec = ChartSpec(
        question_ref="q", chart_type="stacked_vertical_bar", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(width_in), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(levels),
        fmt=spec.number_format)
    rows: list[int] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        for axes in self.axes:
            leg = axes.get_legend()
            if leg is None:
                continue
            ys = {round(t.get_window_extent(self.canvas.get_renderer()).y0, 0)
                  for t in leg.get_texts()}
            rows.append(len(ys))
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["stacked_vertical_bar"](ctx)
    finally:
        Figure.savefig = original
    return max(rows, default=0)


def test_six_short_entries_fit_on_one_row():
    """The defect: five was the limit whatever the labels or the width."""
    assert _legend_rows(SHORT) == 1


def test_five_entries_still_fit():
    assert _legend_rows(SHORT[:5]) == 1


def test_labels_too_wide_still_wrap():
    """Measuring must not become 'always one row' — seven long labels do not
    fit across a slide, and squeezing them onto one line would overlap."""
    assert _legend_rows(LONG) > 1


def test_the_same_count_can_go_either_way():
    """The whole point of measuring rather than counting: six entries fit or do
    not depending on the WORDS, and no count can tell those two apart.

    Not by slot width — `new_figure` floors the figure at 9in, so a narrow slot
    still gets a wide figure and the legend still fits. The label text is what
    varies.
    """
    assert _legend_rows(SHORT) == 1
    assert _legend_rows(LONG[:6]) > 1
