"""A radar prints the value at each point — with up to three groups.

"Voisiko varmistaa, että kaikissa kaaviotyypeissä näkyy prosentit kun tuloksia
tarkastellaan kohderyhmittäin." (2026-09-25) A radar printed no values at all.
Now each point carries its value, in its group's colour, just outside it on its
spoke; values on one spoke that would overlap are pushed apart. With four or
more groups they would bury the shape, so none. The "Numbers" switch
(`elements.data_labels`) turns them off like on every other chart.
"""
from __future__ import annotations

import os

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_ATTRS = ("Maku", "Hinta", "Laatu", "Saatavuus", "Pakkaus", "Brändi", "Terveellisyys")


def _series(n_groups: int, close: bool = False) -> SeriesResult:
    groups = tuple(f"Ryhmä {g}" for g in range(1, n_groups + 1))
    cells = {}
    for i, a in enumerate(_ATTRS):
        for j, g in enumerate(groups):
            # `close`: every group within a point or two of the others — the
            # case where values on a spoke would print over each other.
            v = 40 + (i * 7) % 30 + (j * 1.5 if close else j * 12)
            cells[(a, g)] = Cell(pct=float(v))
    return SeriesResult(categories=_ATTRS, segments=groups, cells=cells,
                        base_n={g: 200 for g in groups}, statistic="pct")


def _draw(n_groups: int, *, close: bool = False, numbers: bool = True, save: str = ""):
    spec = ChartSpec(question_ref="q", chart_type="radar", statistic="pct",
                     classifying_var="ryhma" if n_groups > 1 else None,
                     number_format=NumberFormat(), sort=SortSpec(basis="data_order"),
                     template_slot="s1", elements=ElementToggles(data_labels=numbers))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(9.0), height=Inches(5.0), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(n_groups, close),
        fmt=spec.number_format, notes=[])
    seen, names = [], []
    orig = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        labels = [t for ax in self.axes for t in ax.texts
                  if t.get_gid() == VALUE_GID and t.get_visible()]
        seen.append([(t.get_text(), t.get_window_extent(r)) for t in labels])
        names.append([t.get_window_extent(r) for ax in self.axes
                      for t in ax.get_xticklabels() if t.get_visible() and t.get_text()])
        if save:
            orig(self, save, dpi=110)
        return orig(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["radar"](ctx)
    finally:
        Figure.savefig = orig
    _draw.names = names[-1]
    return seen[-1]


@pytest.mark.parametrize("n_groups", [1, 2, 3])
def test_each_point_carries_its_value(n_groups):
    assert len(_draw(n_groups)) == len(_ATTRS) * n_groups


def test_values_on_one_spoke_do_not_overlap():
    """Three groups within a point or two of each other on every spoke."""
    boxes = [b for _t, b in _draw(3, close=True)]
    for i, a in enumerate(boxes):
        for b in boxes[i + 1:]:
            overlap_x = min(a.x1, b.x1) - max(a.x0, b.x0)
            overlap_y = min(a.y1, b.y1) - max(a.y0, b.y0)
            assert not (overlap_x > 1 and overlap_y > 1), (a, b)


def test_four_groups_print_none():
    assert _draw(4) == []


def test_the_numbers_switch_turns_them_off():
    assert _draw(2, numbers=False) == []


def test_pictures_for_review(tmp_path_factory):
    """Not an assertion — writes the pictures a person looks at, when asked to."""
    out = os.environ.get("RADAR_REVIEW_DIR")
    if not out:
        pytest.skip("set RADAR_REVIEW_DIR to write the review pictures")
    for n, close in ((1, False), (2, False), (3, True), (4, False)):
        _draw(n, close=close, save=os.path.join(out, f"radar_{n}groups{'_close' if close else ''}.png"))


def test_values_stay_off_the_spoke_names():
    """Past the outer ring the spoke's name is printed; a value pushed out there
    printed over it ("64 %" over "Saatavuus"). Three close groups."""
    labels = _draw(3, close=True)
    assert labels
    for text, b in labels:
        for n in _draw.names:
            overlap_x = min(b.x1, n.x1) - max(b.x0, n.x0)
            overlap_y = min(b.y1, n.y1) - max(b.y0, n.y0)
            assert not (overlap_x > 1 and overlap_y > 1), (text, b, n)
