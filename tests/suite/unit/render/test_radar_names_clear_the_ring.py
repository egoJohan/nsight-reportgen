"""Every spoke's name is printed outside the radar's outer ring.

Matplotlib centres a polar tick label on its anchor, so a long name on the
right-hand spoke reached back across the ring by half its width and a
multi-line name above the circle by half its height: the ring was drawn
through "läheisenä" and "henkilökohtaista", and a series reaching the top of
the scale ran through the name at the end of its spoke. (visual QA, 2026-09-19)
"""
from __future__ import annotations

import math

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_LONG = ("Kyllä, minulla on henkilökohtaista kokemusta hoivapalveluista asiakkaana",
         "Kyllä, minulla on kokemusta hoivapalveluista asiakkaan läheisenä",
         "Kyllä, työskentelen tai olen työskennellyt hoivapalveluiden parissa",
         "Ei, minulla ei ole kokemusta hoivapalveluista")
_REGIONS = ("Pääkaupunkiseudulla", "Muualla Etelä-Suomessa", "Länsi-Suomessa",
            "Pohjois-Suomessa", "Itä-Suomessa")


def _geometry(cats):
    cells = {(c, "Total"): Cell(pct=20.0 + 15 * i) for i, c in enumerate(cats)}
    series = SeriesResult(categories=cats, segments=("Total",), cells=cells,
                          base_n={"Total": 1000}, statistic="pct")
    spec = ChartSpec(question_ref="q", chart_type="radar", statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    got: dict = {}

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        ax = self.axes[0]
        cx, cy = ax.transData.transform((0.0, 0.0))
        ex, ey = ax.transData.transform((0.0, ax.get_ylim()[1]))
        got["centre"], got["radius"] = (cx, cy), math.hypot(ex - cx, ey - cy)
        got["names"] = [(t.get_text(), t.get_window_extent(r))
                        for t in ax.get_xticklabels() if t.get_text()]
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS["radar"](ctx)
    finally:
        _f.Figure.savefig = original
    return got


@pytest.mark.parametrize("cats", [_LONG, _REGIONS], ids=["long", "regions"])
def test_no_name_reaches_inside_the_ring(cats):
    g = _geometry(cats)
    (cx, cy), radius = g["centre"], g["radius"]
    assert len(g["names"]) == len(cats)
    for text, b in g["names"]:
        # The point of the name's box nearest the centre.
        nx, ny = min(max(cx, b.x0), b.x1), min(max(cy, b.y0), b.y1)
        assert math.hypot(nx - cx, ny - cy) >= radius - 1.0, (
            f"{text!r} reaches inside the ring")


@pytest.mark.parametrize("cats", [_LONG, _REGIONS], ids=["long", "regions"])
def test_no_name_is_printed_over_another(cats):
    names = _geometry(cats)["names"]
    for i, (ta, a) in enumerate(names):
        for tb, b in names[i + 1:]:
            assert not (min(a.x1, b.x1) - max(a.x0, b.x0) > 0.5
                        and min(a.y1, b.y1) - max(a.y0, b.y0) > 0.5), (ta, tb)
