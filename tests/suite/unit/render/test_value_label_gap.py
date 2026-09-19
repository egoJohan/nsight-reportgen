"""A bar's number sits the same short distance past its end on every axis.

The gap was added to the VALUE — `max(0.5, max * 0.01)` data units, a size
chosen for a 0-100 axis. A mean lives on 0-5, where half a unit is a tenth of the
scale: on Attendo's brand battery every mean floated a finger's width above its
bar, "3.6" hovering over a bar that ended at 3.64 like a label for the gridline.
(visual QA, 2026-09-19)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("Luotettava", "Välittävä", "Ahne", "Kodikas")

#: The largest gap, in points, a number may keep from the end of its bar.
_MAX_GAP_PT = 8.0


def _gaps(chart_type: str, statistic: str, values) -> list[float]:
    cells = {(c, "Total"): Cell(pct=v, mean=v, count=v) for c, v in zip(_CATS, values)}
    series = SeriesResult(categories=_CATS, segments=("Total",), cells=cells,
                          base_n={"Total": 800}, statistic=statistic)
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic=statistic,
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=NumberFormat())
    gaps: list[float] = []

    import matplotlib.figure as _f
    original = _f.Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        px_per_pt = self.dpi / 72.0
        for ax in self.axes:
            bars = [p for c in ax.containers for p in c]
            labels = [t for t in ax.texts if t.get_gid() == VALUE_GID and t.get_visible()]
            for bar, t in zip(bars, labels):
                b, tb = bar.get_window_extent(r), t.get_window_extent(r)
                gap = (tb.y0 - b.y1) if chart_type == "vertical_bar" else (tb.x0 - b.x1)
                gaps.append(gap / px_per_pt)
        return original(self, *a, **k)

    _f.Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        _f.Figure.savefig = original
    return gaps


@pytest.mark.parametrize("chart_type", ["vertical_bar", "horizontal_bar"])
@pytest.mark.parametrize("statistic,values", [
    ("mean", (3.07, 3.03, 3.64, 3.16)),     # a 1-5 battery's means
    ("pct", (58.0, 70.0, 45.0, 36.0)),      # ordinary shares
    ("pct", (3.0, 4.0, 2.0, 1.0)),          # shares that never pass 5 %
])
def test_the_number_sits_just_past_the_bar(chart_type, statistic, values):
    gaps = _gaps(chart_type, statistic, values)
    assert len(gaps) == len(_CATS), f"expected a number per bar, got {gaps}"
    for g in gaps:
        assert 0.0 <= g <= _MAX_GAP_PT, (
            f"{statistic} {chart_type}: a number {g:.1f}pt from its bar")
