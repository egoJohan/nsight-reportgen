"""A bar chart viewed by target group shows its percentages.

"Horizontal ja vertical bareissa ei näy prosenttilukuja kun tuloksia
tarkastellaan kohderyhmittäin" (Taffel, 2026-09-25): six age groups over six
answers in a template's short chart area gave bars 3.6pt thick, and the rule
wanted 5pt, so every number was dropped. Now the bar is measured once the plot
is laid out, and when the digits would not fit, each answer's bars take more of
its row (85% instead of 70%) — only on a chart that would otherwise lose its
numbers, so every chart numbered before is drawn exactly as before.
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.render.image._mpl import VALUE_GID
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_ANSWERS = ("Perinteiset perunalastut", "Maalaisperunalastut tai Kettle-perunalastut",
            "Linssisipsit", "Juustonaksut", "Juures- tai kasvissipsit", "En mitään näistä")
_AGES = ("18–24", "25–34", "35–44", "45–54", "55–64", "65–74")


def _series(answers=_ANSWERS, groups=_AGES) -> SeriesResult:
    cells = {(a, g): Cell(pct=float(5 + (i * 7 + j * 11) % 90))
             for i, a in enumerate(answers) for j, g in enumerate(groups)}
    return SeriesResult(categories=tuple(answers), segments=tuple(groups), cells=cells,
                        base_n={g: 170 for g in groups}, statistic="pct")


def _draw(chart_type: str, height_in: float, series: SeriesResult):
    spec = ChartSpec(question_ref="q", chart_type=chart_type, statistic="pct",
                     classifying_var="ika", number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles())
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    notes: list = []
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(9.3), height=Inches(height_in), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format, notes=notes)
    seen = []
    orig = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        labels = [t for ax in self.axes for t in ax.texts if t.get_gid() == VALUE_GID]
        seen.append({"n": len(labels), "sizes": [t.get_fontsize() for t in labels],
                     "boxes": [(t.get_window_extent(r), t.get_text(), t.get_fontsize())
                               for t in labels], "dpi": self.dpi})
        return orig(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS[chart_type](ctx)
    finally:
        Figure.savefig = orig
    return seen[-1], notes


def test_the_reported_chart_shows_every_percentage():
    """The Taffel slide's shape: 36 bars in a 3.45in-tall chart."""
    f, notes = _draw("horizontal_bar", 3.45, _series())
    assert f["n"] == 36, f["n"]
    assert min(f["sizes"]) >= 5.0
    assert not any(getattr(n, "kind", "") == "unlabelled" for n in notes)


def test_stacked_numbers_do_not_touch():
    """One above the other, their digits must not meet: centres at least a
    digit's height apart."""
    from reportbuilder.render.image.bars import _ink_height_in

    f, _ = _draw("horizontal_bar", 3.45, _series())
    items = sorted(f["boxes"], key=lambda b: b[0].y0)
    for (a, ta, pa), (b, tb, pb) in zip(items, items[1:]):
        if a.x1 > b.x0 and b.x1 > a.x0:  # side by side horizontally
            ink = max(_ink_height_in(ta, pa), _ink_height_in(tb, pb)) * f["dpi"]
            assert (b.y0 + b.y1) / 2 - (a.y0 + a.y1) / 2 >= ink - 0.5, (ta, tb)


def test_a_chart_numbered_before_is_drawn_as_before():
    """Room enough by the old rule: the same size as before, nothing widened."""
    f, _ = _draw("horizontal_bar", 6.0, _series(groups=_AGES[:2]))
    assert f["n"] == 12
    assert min(f["sizes"]) > 5.5


def test_far_too_dense_still_says_so():
    """22 answers x 6 groups: bars under 2pt. No numbers — and the author is told."""
    answers = tuple(f"Vastaus {i}" for i in range(22))
    f, notes = _draw("horizontal_bar", 3.45, _series(answers=answers))
    assert f["n"] == 0
    assert any(getattr(n, "kind", "") == "unlabelled" for n in notes)
