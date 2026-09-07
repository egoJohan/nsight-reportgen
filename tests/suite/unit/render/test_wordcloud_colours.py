"""A word cloud is transparent, and its words are the TEMPLATE's colour.

It used to bake two house decisions into its own pixel raster: nSight cream
behind the words, and a hardcoded nSight-teal ramp for the words themselves.
Both survived every other mechanism for honouring a customer's template — the
figure is saved with `transparent=True`, but that blanks the matplotlib patches
and not an imshow'd array, and the ramp was a module constant no template could
reach. On a deck that is neither cream nor teal, the word cloud was the one
slide still wearing nSight's clothes. (Johan, 2026-09-07)
"""
from __future__ import annotations

import io

import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches

from reportbuilder.model.report import ChartSpec, ElementToggles, NumberFormat, SortSpec
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.house_style import ramp_from
from reportbuilder.render.image.wordcloud import build_image_wordcloud
from reportbuilder.stats.series import Cell, SeriesResult


class _Styled(StyleSpec):
    """A template that states its brand colour and its own ground."""
    def __init__(self, accent="", background=""):
        self.accent = accent
        self.background = background


def _ctx(style=None):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    spec = ChartSpec(question_ref="q", chart_type="wordcloud", statistic="count",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="slot1",
                     elements=ElementToggles())
    counts = {"kallis": 185.0, "huono": 75.0, "luotettava": 61.0,
              "hyvä": 60.0, "ahne": 57.0}
    series = SeriesResult(
        categories=tuple(counts), segments=("Total",),
        cells={(c, "Total"): Cell(pct=None, count=v, mean=None)
               for c, v in counts.items()},
        base_n={"Total": 817}, statistic="count")
    slot = Slot(slide_index=0, left=Inches(1), top=Inches(1),
                width=Inches(8), height=Inches(5), name="slot1")
    ctx = RenderContext(slide=slide, slot=slot, style=style or StyleSpec(),
                        spec=spec, series=series, fmt=spec.number_format)
    return slide, ctx


def _drawn(style=None) -> Image.Image:
    """The PNG the builder actually put on the slide, as RGBA."""
    slide, ctx = _ctx(style)
    build_image_wordcloud(ctx)
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    assert len(pics) == 1
    return Image.open(io.BytesIO(pics[0].image.blob)).convert("RGBA")


def _opaque_colours(im: Image.Image) -> set[tuple[int, int, int]]:
    return {px[:3] for px in im.getdata() if px[3] > 200}


def test_the_ground_shows_through_instead_of_a_painted_rectangle():
    im = _drawn()
    w, h = im.size
    corners = [im.getpixel(p) for p in ((0, 0), (w - 1, 0), (0, h - 1), (w - 1, h - 1))]
    assert all(px[3] == 0 for px in corners), f"corner is painted: {corners}"
    # not just the corners: most of a word cloud is gaps
    clear = sum(1 for px in im.getdata() if px[3] == 0)
    assert clear > im.size[0] * im.size[1] * 0.5, "the raster is mostly opaque"


def test_the_words_are_drawn_in_the_templates_colour():
    accent = "#B3005E"                      # a magenta no house ramp contains
    im = _drawn(_Styled(accent=accent))
    wanted = {tuple(int(c[i:i + 2], 16) for i in (1, 3, 5))
              for c in ramp_from(accent, steps=6)}
    drawn = _opaque_colours(im)
    hits = {c for c in drawn if any(
        abs(c[0] - w[0]) + abs(c[1] - w[1]) + abs(c[2] - w[2]) <= 12 for w in wanted)}
    assert hits, f"none of {sorted(wanted)} was drawn; got {sorted(drawn)[:12]}"
    # and NOT the house teal it used to hardcode
    assert not any(abs(c[0] - 0x13) + abs(c[1] - 0x61) + abs(c[2] - 0x5E) <= 12
                   for c in drawn), "house teal is still in there"


def test_a_template_that_states_no_colour_still_gets_the_house_ramp():
    drawn = _opaque_colours(_drawn())
    house = {tuple(int(c[i:i + 2], 16) for i in (1, 3, 5))
             for c in ramp_from("", steps=6)}
    assert any(any(abs(c[0] - h[0]) + abs(c[1] - h[1]) + abs(c[2] - h[2]) <= 12
                   for h in house) for c in drawn)
