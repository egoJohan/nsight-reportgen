"""An author's title box moves the drawn title.

The layout editor lets somebody drag the title area onto the place their
template actually puts its headline. That correction reached the style — the
`profile.title` box moved — and stopped there: every reader of the title
consults `slide.shapes.title` FIRST and only falls back to the profile, so on
any template whose layout carries a title placeholder (most of them) the text
stayed exactly where the placeholder was. The editor showed the box in its new
home and the slide ignored it.

Reported 2026-09-08: "the title does not reposition to title area when editing
template layout".

Harvested and authored are kept apart on purpose. A template nobody has touched
must keep inheriting the placeholder's own position, size, font and colour —
that is the whole point of using the customer's layout — so only an explicit
correction moves it.
"""
from __future__ import annotations

import pytest
from pptx.util import Emu, Inches

from reportbuilder.render.style_spec import apply_template_overrides, load_style_spec

_TEMPLATE = "src/reportbuilder/render/assets/nsight_default.pptx"
_EMU = 914400.0


def _title_shape(style):
    """Build the sample slide this style produces and return its title shape."""
    from reportbuilder.api.routes_templates import _sample_report
    from reportbuilder.export.pptx_build import build_presentation

    report, model, df = _sample_report()
    slide = build_presentation(report, model, df, style=style).slides[0]
    return slide.shapes.title


def _styled(overrides: dict):
    style = load_style_spec(_TEMPLATE)
    apply_template_overrides(style, overrides)
    return style


def test_the_untouched_template_keeps_its_own_title_position():
    """No correction, no movement: the placeholder is inherited as it is."""
    plain = _title_shape(_styled({}))
    assert plain is not None, "this fixture is supposed to have a placeholder"
    assert int(plain.left or 0) / _EMU == pytest.approx(0.34, abs=0.05)


def test_a_dragged_title_box_moves_the_headline():
    moved = _title_shape(_styled({"title": {"x": 2.5, "y": 3.0}}))
    assert moved is not None
    assert int(moved.left or 0) / _EMU == pytest.approx(2.5, abs=0.02), \
        f"left is {int(moved.left or 0) / _EMU:.2f}in — the placeholder won"
    assert int(moved.top or 0) / _EMU == pytest.approx(3.0, abs=0.02), \
        f"top is {int(moved.top or 0) / _EMU:.2f}in — the placeholder won"


def test_a_resized_title_box_takes_its_width_and_height_too():
    box = _title_shape(_styled({"title": {"x": 1.0, "y": 2.0, "w": 6.0, "h": 1.2}}))
    assert int(box.width or 0) / _EMU == pytest.approx(6.0, abs=0.02)
    # the height may GROW for a headline that wraps, but never shrink below it
    assert int(box.height or 0) / _EMU >= 1.2 - 0.02


def test_moving_it_defeats_the_pull_up_that_normally_reclaims_the_margin():
    """A template that parks its title a third of the way down is normally
    raised into the empty band above it. An author who dragged the box there ON
    PURPOSE must not have it silently moved again."""
    low = _title_shape(_styled({"title": {"x": 1.0, "y": 3.5}}))
    assert int(low.top or 0) / _EMU == pytest.approx(3.5, abs=0.02)
