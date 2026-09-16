"""A combo says which two kinds of chart it is a combo OF.

"Combo chartissa puuttuu määrittely minkä kuvaajien combo piirretään.
Ykköstyypin kuvaajana on aina vertical bar ja kakkosena Line."

`combo_primary_type` and `combo_secondary_type` pick the shape of each half —
bars, a line, or a filled area. Both default to what the chart drew before
(bars, then line), so no saved slide changes on upgrade.

Bars on BOTH axes is offered rather than refused. It needs the two groups to
share one cluster: each is on its own y-scale but they sit in the same category
slot, and two bar groups that each allocate the full slot draw straight through
each other. (Johan, 2026-09-16)
"""
from __future__ import annotations

import matplotlib
matplotlib.use("Agg")
import pytest
from matplotlib.collections import PolyCollection
from matplotlib.container import BarContainer
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.util import Inches

from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.base import RenderContext, Slot, StyleSpec
from reportbuilder.render.image import IMAGE_BUILDERS
from reportbuilder.stats.series import Cell, SeriesResult

pytestmark = pytest.mark.unit

_CATS = ("1", "2", "3", "4")
_GROUPS = ("Nainen", "Mies")
_INDEX = "tyoelamaindeksi"


def _series(groups=_GROUPS) -> SeriesResult:
    segs = tuple(groups) + (_INDEX,)
    cells = {}
    for i, c in enumerate(_CATS):
        for j, g in enumerate(groups):
            cells[(c, g)] = Cell(pct=float(20 + i * 5 + j * 3))
        cells[(c, _INDEX)] = Cell(pct=float(5 + i * 0.4))
    return SeriesResult(
        categories=_CATS, segments=segs, cells=cells,
        base_n={"Total": 1000, **{s: 500 for s in segs}}, statistic="pct",
        segment_statistics={**{g: "pct" for g in groups}, _INDEX: "mean"})


def _render(groups=_GROUPS, **options) -> dict:
    spec = ChartSpec(
        question_ref="q1", chart_type="combo", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles(title=True, subtitle=True, legend=True, n=True,
                                axis_names=True, filter_var=True, data_labels=True),
        options=dict(options))
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=_series(groups),
        fmt=spec.number_format)

    seen: dict = {}
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        per_axes = []
        for axes in self.axes:
            per_axes.append({
                "bars": [c for c in axes.containers if isinstance(c, BarContainer)],
                "lines": [ln for ln in axes.lines
                          if len(ln.get_xdata()) == len(_CATS)],
                "areas": [c for c in axes.collections
                          if isinstance(c, PolyCollection)],
                "texts": [t for t in axes.texts if t.get_text().strip()],
            })
        seen["axes"] = per_axes
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        Figure.savefig = original
    return seen


def _count(seen, what: str) -> int:
    return sum(len(a[what]) for a in seen["axes"])


def test_the_default_is_the_chart_it_always_drew():
    """No options set — bars on the left, a line on the right."""
    seen = _render()
    assert _count(seen, "bars") == len(_GROUPS)
    assert _count(seen, "lines") == 1
    assert _count(seen, "areas") == 0


def test_the_primary_can_be_a_line():
    seen = _render(combo_primary_type="line")
    assert _count(seen, "bars") == 0
    assert _count(seen, "lines") == len(_GROUPS) + 1


def test_the_secondary_can_be_bars():
    seen = _render(combo_secondary_type="bar")
    assert _count(seen, "lines") == 0
    assert _count(seen, "bars") == len(_GROUPS) + 1


def test_either_half_can_be_an_area():
    assert _count(_render(combo_primary_type="area"), "areas") >= len(_GROUPS)
    assert _count(_render(combo_secondary_type="area"), "areas") >= 1


def test_bars_on_both_axes_do_not_draw_through_each_other():
    """The case that needs the two groups to share one cluster."""
    seen = _render(combo_secondary_type="bar")
    centres = []
    for axes in seen["axes"]:
        for container in axes["bars"]:
            centres.append(tuple(round(p.get_x() + p.get_width() / 2, 6)
                                 for p in container.patches))
    assert len(centres) == len(_GROUPS) + 1
    assert len(set(centres)) == len(centres), (
        f"two bar series share the same x positions and overlap: {centres}")


def test_a_half_keeps_its_numbers_whatever_shape_it_is():
    """Changing the shape must not silently drop the values. Drawn as bars the
    primary half labels every bar; drawn as a line or an area it labelled
    nothing at all, so picking "Line" quietly cost the reader the numbers.

    One series, because that is the rule a line already followed for the
    secondary half: several lines put their labels in the same narrow band and
    land on each other, and unlike bars there is no column width to measure a
    fit against. `test_several_lines_are_left_to_the_axis` holds that down.
    """
    for kind in ("bar", "line", "area"):
        seen = _render(groups=("Total",), combo_primary_type=kind)
        drawn = [t.get_text() for a in seen["axes"] for t in a["texts"]]
        assert any("20" in t for t in drawn), (
            f"primary drawn as {kind!r} printed no values: {drawn}")


def test_several_lines_are_left_to_the_axis():
    """Two line series would stack their labels on each other."""
    seen = _render(combo_primary_type="line")      # two groups in the fixture
    drawn = [t.get_text() for a in seen["axes"] for t in a["texts"]]
    assert not any("20" in t for t in drawn), drawn


def test_a_value_is_printed_once():
    """The secondary line is placed by the anchored rule; letting the half
    label itself as well printed every number twice."""
    seen = _render()
    drawn = [t.get_text() for a in seen["axes"] for t in a["texts"]]
    line_labels = [t for t in drawn if t in {"5.0", "5.4", "5.8", "6.2"}]
    assert len(line_labels) == len(set(line_labels)), f"duplicated: {drawn}"


def test_an_unknown_type_falls_back_instead_of_drawing_nothing():
    """A saved slide from a future version, or a typo, still renders."""
    seen = _render(combo_primary_type="sparkline", combo_secondary_type="ufo")
    assert _count(seen, "bars") + _count(seen, "lines") + _count(seen, "areas") > 0
