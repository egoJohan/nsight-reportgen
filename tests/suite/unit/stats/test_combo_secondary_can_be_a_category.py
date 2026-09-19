"""A combo's secondary variable may be a categorical one — drawn as the share of
one of its groups.

"Combo chartissa secondary variableen tulee näkyviin vain rajallinen määrä
muuttujia. Muuttujalistasta puuttuu suuri joukko luokittelevia muuttujia jotka
olisivat käyttökelpoisia kun secondary series määrittely on Bars."

The secondary was only ever a MEAN per category, so only variables a mean could
be taken of were offered. A mean of a gender or a region code says nothing, so
the list was right to leave them out — what was missing was a measure that DOES
mean something for them: the share of respondents in each category who fall in
one chosen group ("Kyllä", "Nainen", "Uusimaa"). It is one series, so it draws
as bars, a line or an area alike, and it sits beside a classifier split rather
than duplicating it. (Johan, 2026-09-17)
"""
from __future__ import annotations

import dataclasses

import numpy as np
import pandas as pd
import pytest

from reportbuilder.model.question import (
    Question, QuestionModel, ValueLabel, Variable,
)
from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec,
)
from reportbuilder.render.image.combo import split_primary_and_secondary_segments
from reportbuilder.stats import engine

pytestmark = pytest.mark.unit


def _study():
    rng = np.random.default_rng(5)
    n = 900
    age = rng.choice([1, 2, 3], size=n)
    # "Kyllä" rises with the age group, so a wrong lookup cannot pass by luck.
    yes = rng.random(n) < np.array([0.2, 0.5, 0.8])[age - 1]
    had = np.where(yes, 1.0, 2.0)
    had[:30] = 9.0                                   # "En osaa sanoa" — missing
    q = Variable("age", "Ikäluokka", "categorical",
                 (ValueLabel(1.0, "Nuoret"), ValueLabel(2.0, "Keski-ikäiset"),
                  ValueLabel(3.0, "Vanhat")), frozenset())
    sec = Variable("had", "Onko ollut aiemmin MV", "categorical",
                   (ValueLabel(1.0, "Kyllä"), ValueLabel(2.0, "Ei"),
                    ValueLabel(9.0, "En osaa sanoa")), frozenset({9.0}))
    model = QuestionModel(variables={"age": q, "had": sec}, questions=[])
    question = Question(qid="age", kind="single", variables=("age",), text=q.label)
    df = pd.DataFrame({"age": age.astype(float), "had": had})
    return model, question, df


def _spec(**options) -> ChartSpec:
    return ChartSpec(
        question_ref="age", chart_type="combo", statistic="pct",
        classifying_var=None, number_format=NumberFormat(),
        sort=SortSpec(basis="data_order"), template_slot="s1",
        elements=ElementToggles(),
        options={"combo_secondary": "had", **options})


def _expected(df, cat_code):
    rows = df[(df["age"] == cat_code) & (df["had"] != 9.0)]
    return 100.0 * (rows["had"] == 1.0).mean()


def test_the_share_of_the_chosen_group_is_drawn_per_category():
    model, q, df = _study()
    r = engine.compute(q, _spec(combo_secondary_value="Kyllä"), df, model)
    sec = r.secondary_segments[0]
    for code, label in ((1.0, "Nuoret"), (2.0, "Keski-ikäiset"), (3.0, "Vanhat")):
        assert r.cell(label, sec).pct == pytest.approx(_expected(df, code))


def test_the_share_is_a_percentage_not_a_mean():
    model, q, df = _study()
    r = engine.compute(q, _spec(combo_secondary_value="Kyllä"), df, model)
    assert r.statistic_of(r.secondary_segments[0]) == "pct"


def test_the_series_names_the_group_it_counts():
    model, q, df = _study()
    r = engine.compute(q, _spec(combo_secondary_value="Kyllä"), df, model)
    assert "Kyllä" in r.secondary_segments[0]


def test_the_renderer_finds_the_share_although_both_halves_are_percentages():
    """Both halves are `pct`, so "measures something else" cannot tell them apart
    — the series has to say which half is the secondary variable."""
    model, q, df = _study()
    r = engine.compute(q, _spec(combo_secondary_value="Kyllä"), df, model)
    primary, secondary = split_primary_and_secondary_segments(r, list(r.segments))
    assert secondary == list(r.secondary_segments)
    assert primary and not set(primary) & set(secondary)


def test_with_a_classifier_every_group_stays_primary():
    model, q, df = _study()
    model.variables["sex"] = Variable(
        "sex", "Sukupuoli", "categorical",
        (ValueLabel(1.0, "Mies"), ValueLabel(2.0, "Nainen"), ValueLabel(3.0, "Muu")),
        frozenset())
    df["sex"] = np.resize([1.0, 2.0, 3.0], len(df))
    spec = dataclasses.replace(_spec(combo_secondary_value="Kyllä"),
                               classifying_var="sex")
    r = engine.compute(q, spec, df, model)
    primary, secondary = split_primary_and_secondary_segments(r, list(r.segments))
    assert len(secondary) == 1
    assert {"Mies", "Nainen", "Muu"} <= set(primary)


def test_without_a_group_it_is_still_the_mean():
    """Every slide saved before this keeps drawing what it drew."""
    model, q, df = _study()
    model.variables["idx"] = Variable("idx", "index", "scale", (), frozenset())
    df["idx"] = 4.0 + df["age"]
    spec = dataclasses.replace(_spec(), options={"combo_secondary": "idx"})
    r = engine.compute(q, spec, df, model)
    sec = r.secondary_segments[0]
    assert r.statistic_of(sec) == "mean"
    assert r.cell("Vanhat", sec).pct == pytest.approx(7.0)


def _drawn(spec, series) -> dict:
    """Render the combo and read back what a reader sees: the legend, and
    whether a right-hand axis line is drawn."""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from pptx import Presentation
    from pptx.util import Inches
    from reportbuilder.render.base import RenderContext, Slot, StyleSpec
    from reportbuilder.render.image import IMAGE_BUILDERS

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide,
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    seen: dict = {}
    original = Figure.savefig

    def spy(self, *a, **k):
        legends = [ax.get_legend() for ax in self.axes if ax.get_legend()]
        seen["legend"] = [t.get_text() for lg in legends for t in lg.get_texts()]
        seen["right_line"] = any(ax.spines["right"].get_visible() for ax in self.axes)
        seen["ylims"] = [ax.get_ylim() for ax in self.axes]
        self.canvas.draw()
        renderer = self.canvas.get_renderer()
        fig_box = self.bbox
        seen["legend_inside"] = all(
            lg.get_window_extent(renderer).x0 >= fig_box.x0 - 1
            and lg.get_window_extent(renderer).x1 <= fig_box.x1 + 1
            for lg in legends)
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        Figure.savefig = original
    return seen


@pytest.mark.parametrize("kind", ["bar", "line", "area"])
def test_the_legend_names_the_share(kind):
    """Seen on the first real slide: the share shares the bars' axis, and the
    legend only ever named a series on the SECOND axis — so the light bars
    were drawn with nothing on the slide to say what they were."""
    model, q, df = _study()
    spec = _spec(combo_secondary_value="Kyllä", combo_secondary_type=kind)
    seen = _drawn(spec, engine.compute(q, spec, df, model))
    assert any("Kyllä" in t for t in seen["legend"]), seen


def test_the_share_has_a_right_hand_axis_graduated_like_the_left():
    """Its own ruler, so it reads as apart from the bars — but the same ruler,
    so a 40 % is the same height on either side and nothing is misread."""
    model, q, df = _study()
    spec = _spec(combo_secondary_value="Kyllä", combo_secondary_type="bar")
    seen = _drawn(spec, engine.compute(q, spec, df, model))
    assert seen["right_line"]
    assert len(seen["ylims"]) == 2 and seen["ylims"][0] == pytest.approx(seen["ylims"][1])


def test_the_share_keeps_its_own_base():
    """Everyone who answered the secondary question — not the slide's N."""
    model, q, df = _study()
    spec = _spec(combo_secondary_value="Kyllä")
    r = engine.compute(q, spec, df, model)
    answered = int((df["had"] != 9.0).sum())
    assert r.base_n[r.secondary_segments[0]] == answered


def test_the_legend_states_how_many_are_in_the_group():
    """"Kyllä (n=…)" is read as how many said Kyllä. It stated the share's
    base, everyone who answered — "ei ole Kyllä-vastausten N, vaan kaikkien
    vastausten N-luku", reported from staging. (2026-09-19)"""
    model, q, df = _study()
    spec = _spec(combo_secondary_value="Kyllä")
    r = engine.compute(q, spec, df, model)
    said_yes = int((df["had"] == 1.0).sum())
    answered = int((df["had"] != 9.0).sum())
    assert r.secondary_group_n == said_yes
    legend = _drawn(spec, r)["legend"]
    assert any(f"n={said_yes}" in t for t in legend), legend
    assert not any(f"n={answered}" in t for t in legend), legend


def test_a_combo_without_a_secondary_draws_no_bare_right_hand_line():
    """The fallback — a classifier's groups split over the two halves, one
    measure — shares one axis, and must not show a right-hand line with no
    numbers on it."""
    from reportbuilder.stats.series import Cell, SeriesResult
    segs = ("A", "B", "C")
    series = SeriesResult(
        categories=("x", "y"), segments=segs,
        cells={(c, s): Cell(pct=30.0) for c in ("x", "y") for s in segs},
        base_n={s: 100 for s in segs} | {"Total": 300}, statistic="pct")
    spec = dataclasses.replace(_spec(), options={}, classifying_var="g")
    assert not _drawn(spec, series)["right_line"]


def test_an_unlabelled_segment_flag_resolves_the_groups_the_picker_offers():
    """Found on the Mobiilivarmenne data: `AikooJatkaa` is a 0/1 flag with no
    value labels. The picker names its groups "AikooJatkaa" and "Muut" — the
    classifier relabeller's names — while the share looked only at value labels,
    found no group, and the slide silently lost its secondary series."""
    model, q, df = _study()
    model.variables["Flag"] = Variable("Flag", "Flag", "categorical", (), frozenset())
    df["Flag"] = (df["had"] == 1.0).astype(float)
    spec = dataclasses.replace(
        _spec(), options={"combo_secondary": "Flag", "combo_secondary_value": "Flag"})
    r = engine.compute(q, spec, df, model)
    assert r.secondary_segments, "the secondary series was dropped"
    sec = r.secondary_segments[0]
    rows = df[df["age"] == 3.0]
    assert r.cell("Vanhat", sec).pct == pytest.approx(100.0 * (rows["Flag"] == 1.0).mean())


def test_unlabelled_codes_resolve_by_their_code():
    model, q, df = _study()
    model.variables["codes"] = Variable("codes", "Codes", "categorical", (), frozenset())
    df["codes"] = df["had"].replace({9.0: 3.0})
    spec = dataclasses.replace(
        _spec(), options={"combo_secondary": "codes", "combo_secondary_value": "1"})
    r = engine.compute(q, spec, df, model)
    assert r.secondary_segments


LONG = ("53 Aion jatkossakin pysyä DNA:n mobiilivarmenteen käyttäjänä, "
        "vaikka hinta nousisi")


@pytest.mark.parametrize("group", [None, "Kyllä"])
def test_the_secondary_name_is_never_cut(group):
    """"Legend is cut and cannot be edited" — the name was cut at 30
    characters, so the legend said "53 Aion jatkossakin pysyä…" and nobody
    could tell what the bars were. The whole name is kept."""
    model, q, df = _study()
    model.variables["had"] = dataclasses.replace(model.variables["had"], label=LONG)
    model.variables["idx"] = Variable("idx", LONG, "scale", (), frozenset())
    df["idx"] = 4.0 + df["age"]
    opts = ({"combo_secondary": "had", "combo_secondary_value": group} if group
            else {"combo_secondary": "idx"})
    r = engine.compute(q, dataclasses.replace(_spec(), options=opts), df, model)
    assert LONG in r.secondary_segments[0]
    assert "…" not in r.secondary_segments[0]


def test_a_long_legend_entry_wraps_inside_the_picture():
    model, q, df = _study()
    model.variables["had"] = dataclasses.replace(model.variables["had"], label=LONG)
    spec = _spec(combo_secondary_value="Kyllä", combo_secondary_type="bar")
    seen = _drawn(spec, engine.compute(q, spec, df, model))
    text = " ".join(" ".join(t.split()) for t in seen["legend"])
    assert LONG in text, seen["legend"]
    assert seen["legend_inside"], "the legend runs off the picture"


@pytest.mark.parametrize("kind", ["bar", "line"])
def test_the_legend_names_both_halves_when_there_is_a_secondary(kind):
    """"There is only one item in the legend." With a secondary variable the
    slide draws two different things, and the legend named only one of them —
    the question's own bars were left for the subtitle to explain, which it
    does not do beside a second set of bars. (Johan, 2026-09-17)"""
    model, q, df = _study()
    spec = _spec(combo_secondary_value="Kyllä", combo_secondary_type=kind)
    r = engine.compute(q, spec, df, model)
    legend = _drawn(spec, r)["legend"]
    assert len(legend) == 2, legend
    assert any("Ikäluokka" in t for t in legend), legend


def test_the_primary_name_is_never_cut():
    model, q, df = _study()
    long = "Ikäluokka, johon vastaaja kuului kyselyn tekohetkellä"
    model.variables["age"] = dataclasses.replace(model.variables["age"], label=long)
    r = engine.compute(q, _spec(combo_secondary_value="Kyllä"), df, model)
    assert long in r.segments


def test_a_line_half_labels_do_not_land_on_the_bar_labels():
    """"Combo chart piirtää välillä kuvia joita on vaikea ymmärtää." With the
    QUESTION drawn as a line and the secondary as bars, the line's numbers were
    placed a fixed distance above each point and landed on the bars' own
    numbers. Measured on the reported slide: "43 %" printed across "59.1". The
    secondary line has dodged the bars since 2026-09-16; the primary one does
    now too. (Johan, 2026-09-17)"""
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from matplotlib.text import Text
    from pptx import Presentation
    from pptx.util import Inches
    from reportbuilder.model.report import (ChartSpec, ElementToggles, NumberFormat,
                                            SortSpec)
    from reportbuilder.render.base import RenderContext, Slot, StyleSpec
    from reportbuilder.render.image import IMAGE_BUILDERS
    from reportbuilder.render.image._mpl import VALUE_GID
    from reportbuilder.stats.series import Cell, SeriesResult

    cats = ("Negatiivisia", "Neutraaleja", "Positiivisia")
    question, age = "50 Millaisia ajatuksia…", "Ikäluokka"
    pct = dict(zip(cats, (16.0, 43.0, 41.0)))
    mean = dict(zip(cats, (55.1, 59.1, 59.6)))
    series = SeriesResult(
        categories=cats, segments=(question, age),
        cells={**{(c, question): Cell(pct=pct[c]) for c in cats},
               **{(c, age): Cell(pct=mean[c]) for c in cats}},
        base_n={question: 1051, age: 1051, "Total": 1051}, statistic="pct",
        segment_statistics={question: "pct", age: "mean"},
        secondary_segments=(age,))
    spec = ChartSpec(question_ref="q", chart_type="combo", statistic="pct",
                     classifying_var=None, number_format=NumberFormat(),
                     sort=SortSpec(basis="data_order"), template_slot="s1",
                     elements=ElementToggles(),
                     options={"combo_secondary": "age", "combo_primary_type": "line",
                              "combo_secondary_type": "bar"})
    prs = Presentation()
    ctx = RenderContext(
        slide=prs.slides.add_slide(prs.slide_layouts[6]),
        slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                  width=Inches(12.3), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    seen: list = []
    original = Figure.savefig

    def spy(self, *a, **k):
        self.canvas.draw()
        r = self.canvas.get_renderer()
        for axes in self.axes:
            for t in axes.findobj(Text):
                if t.get_gid() == VALUE_GID and t.get_text().strip():
                    seen.append((t, t.get_window_extent(r), axes))
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        IMAGE_BUILDERS["combo"](ctx)
    finally:
        Figure.savefig = original

    assert len(seen) == 6, f"both halves label their values: {[t.get_text() for t, _b, _a in seen]}"
    # The line and its numbers must sit ABOVE the bars: a twinned axis paints
    # over the first one whatever the zorder, so the question's own line ran
    # behind the secondary variable's bars until its axis was raised.
    line_axes = {a for t, _b, a in seen if "%" in t.get_text()}
    bar_axes = {a for t, _b, a in seen if "%" not in t.get_text()}
    assert line_axes and bar_axes
    assert min(a.get_zorder() for a in line_axes) > \
        max(a.get_zorder() for a in bar_axes), "the line is painted over by the bars"
    overlaps = [(seen[i][0].get_text(), seen[j][0].get_text()) for i in range(len(seen))
                for j in range(i + 1, len(seen)) if seen[i][1].overlaps(seen[j][1])]
    assert not overlaps, f"value labels printed over each other: {overlaps}"
