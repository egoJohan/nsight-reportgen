"""The names a legend shows can be edited on the slide — not only the answers.

"Legend is cut and cannot be edited." Category labels edits a question's
ANSWERS, and on a pie or a stacked bar those are exactly what the legend shows.
But split by a classifying variable, the legend shows the GROUPS ("Mies",
"Nainen"), and a combo adds its secondary variable — and neither was in any
editor. `series_label_overrides` renames those.

Kept apart from the category overrides on purpose: a name can be an answer and
a group at once ("Kyllä", "Muu"), and renaming one must not rename the other.
(Johan, 2026-09-17)
"""
from __future__ import annotations

import dataclasses

import numpy as np
import pandas as pd
import pytest

from reportbuilder.model.question import Question, QuestionModel, ValueLabel, Variable
from reportbuilder.model.report import (
    ChartSpec, ElementToggles, NumberFormat, SortSpec, report_from_json, report_to_json,
)
from reportbuilder.stats import engine

pytestmark = pytest.mark.unit


def _study():
    rng = np.random.default_rng(11)
    n = 600
    q = Variable("q", "Mitä mieltä olet?", "categorical",
                 (ValueLabel(1.0, "Kyllä"), ValueLabel(2.0, "Ei")), frozenset())
    sex = Variable("sex", "Sukupuoli", "categorical",
                   (ValueLabel(1.0, "Mies"), ValueLabel(2.0, "Nainen")), frozenset())
    # A classifier whose group is ALSO called "Kyllä" — the collision case.
    yes = Variable("yes", "Onko", "categorical",
                   (ValueLabel(1.0, "Kyllä"), ValueLabel(2.0, "Muu")), frozenset())
    age = Variable("age", "Ikä", "categorical",
                   (ValueLabel(1.0, "Nuori"), ValueLabel(2.0, "Vanha")), frozenset())
    idx = Variable("idx", "Indeksi", "scale", (), frozenset())
    model = QuestionModel(variables={v.name: v for v in (q, sex, yes, age, idx)},
                          questions=[])
    question = Question(qid="q", kind="single", variables=("q",), text=q.label)
    df = pd.DataFrame({
        "q": rng.choice([1.0, 2.0], n), "sex": rng.choice([1.0, 2.0], n),
        "yes": rng.choice([1.0, 2.0], n), "age": rng.choice([1.0, 2.0], n),
        "idx": rng.normal(5, 1, n),
    })
    return model, question, df


def _spec(**kw) -> ChartSpec:
    base = dict(question_ref="q", chart_type="vertical_bar", statistic="pct",
                classifying_var="sex", number_format=NumberFormat(),
                sort=SortSpec(basis="data_order"), template_slot="s1",
                elements=ElementToggles())
    base.update(kw)
    return ChartSpec(**base)


def _compute(**kw):
    model, q, df = _study()
    return engine.compute(q, _spec(**kw), df, model)


# ---- nothing asked, nothing changes -----------------------------------------

@pytest.mark.parametrize("kw", [
    {}, {"classifying_var": None}, {"chart_type": "combo", "options": {"combo_secondary": "idx"}},
    {"classifying_var_2": "age"}, {"chart_type": "pie"},
])
def test_no_renames_is_exactly_the_series_it_was(kw):
    """Every saved slide has no such setting. Its series must be identical,
    field for field, to what it computed before the setting existed."""
    assert _compute(**kw) == _compute(series_label_overrides=(), **kw)


# ---- a group is renamed -----------------------------------------------------

def test_a_group_is_renamed_with_its_numbers():
    before = _compute()
    after = _compute(series_label_overrides=(("Mies", "Miehet"),))
    assert "Miehet" in after.segments and "Mies" not in after.segments
    for cat in before.categories:
        assert after.cell(cat, "Miehet") == before.cell(cat, "Mies")
    assert after.base_n["Miehet"] == before.base_n["Mies"]


def test_a_renamed_group_does_not_rename_an_answer_of_the_same_name():
    before = _compute(classifying_var="yes")
    after = _compute(classifying_var="yes", series_label_overrides=(("Kyllä", "Kyllä-ryhmä"),))
    assert "Kyllä" in after.categories, "the ANSWER was renamed along with the group"
    assert "Kyllä-ryhmä" in after.segments
    assert after.cell("Kyllä", "Kyllä-ryhmä") == before.cell("Kyllä", "Kyllä")


def test_a_renamed_answer_does_not_rename_a_group_renamed_separately():
    after = _compute(classifying_var="yes",
                     category_label_overrides=(("Kyllä", "K"),),
                     series_label_overrides=(("Kyllä", "Ryhmä K"),))
    assert "K" in after.categories
    assert "Ryhmä K" in after.segments


def test_total_is_not_renamed():
    after = _compute(series_label_overrides=(("Total", "Kaikki"),))
    assert "Kaikki" not in after.segments
    assert "Total" in after.base_n


def test_two_groups_renamed_to_one_name_keep_their_own_names():
    """A clash would merge two groups into one dict entry and drop one."""
    after = _compute(series_label_overrides=(("Mies", "X"), ("Nainen", "X")))
    assert {"Mies", "Nainen"} <= set(after.segments)


def test_the_combo_secondary_series_is_renamed_and_stays_the_secondary():
    before = _compute(chart_type="combo", options={"combo_secondary": "idx"})
    sec = before.secondary_segments[0]
    after = _compute(chart_type="combo", options={"combo_secondary": "idx"},
                     series_label_overrides=((sec, "Indeksi (ka.)"),))
    assert after.secondary_segments == ("Indeksi (ka.)",)
    assert after.statistic_of("Indeksi (ka.)") == "mean"
    for cat in before.categories:
        assert after.cell(cat, "Indeksi (ka.)") == before.cell(cat, sec)


def test_a_cross_tab_renames_the_group_in_every_combination():
    """A cross-tab's series are combinations ("Mies · Nuori"), and each part is
    drawn on its own — the bar's label and the group label above the bars. A
    rename applies to the PART wherever it occurs, and the grouping follows."""
    before = _compute(classifying_var_2="age")
    after = _compute(classifying_var_2="age",
                     series_label_overrides=(("Mies", "Miehet"), ("Nuori", "Nuoret")))
    assert "Miehet · Nuoret" in after.segments
    assert not any("Mies" in s.split(" · ") for s in after.segments)
    assert after.segment_primary["Miehet · Nuoret"] == "Miehet"
    old = next(s for s in before.segments if s == "Mies · Nuori")
    for cat in before.categories:
        assert after.cell(cat, "Miehet · Nuoret") == before.cell(cat, old)
    assert after.base_n["Miehet · Nuoret"] == before.base_n[old]


def test_a_cross_tab_part_rename_that_would_merge_two_series_is_not_applied():
    after = _compute(classifying_var_2="age",
                     series_label_overrides=(("Nuori", "Vanha"),))
    assert len(set(after.segments)) == len(after.segments)
    assert any("Nuori" in s for s in after.segments)


# ---- stored and read back ----------------------------------------------------

def _report(chart: ChartSpec):
    from reportbuilder.model.report import Report
    return Report(name="r", render_mode="image", template_ref="", charts=(chart,))


def test_the_setting_survives_saving_and_loading():
    chart = _spec(series_label_overrides=(("Mies", "Miehet"),))
    back = report_from_json(report_to_json(_report(chart))).charts[0]
    assert back.series_label_overrides == (("Mies", "Miehet"),)


def test_a_report_saved_before_the_setting_loads_with_none():
    import json
    raw = json.loads(report_to_json(_report(_spec())))
    del raw["charts"][0]["series_label_overrides"]
    assert report_from_json(raw).charts[0].series_label_overrides == ()


def test_a_group_renamed_to_an_answers_name_is_not_renamed_again_as_that_answer():
    after = _compute(category_label_overrides=(("Ei", "E"),),
                     series_label_overrides=(("Mies", "Ei"),))
    assert "E" in after.categories
    assert "Ei" in after.segments, after.segments


def test_row_summary_keys_that_are_answers_are_not_renamed_as_groups():
    from reportbuilder.stats.series import Cell, SeriesResult
    r = SeriesResult(categories=("Kyllä", "Ei"), segments=("Kyllä", "Muu"),
                     cells={(c, s): Cell(pct=1.0) for c in ("Kyllä", "Ei") for s in ("Kyllä", "Muu")},
                     base_n={"Kyllä": 1, "Muu": 1, "Total": 2}, statistic="pct",
                     row_summaries=(1.0, 2.0), row_summary_keys=("Kyllä", "Ei"))
    out = engine._series_relabelled(r, {"Muu": "Muut"})
    assert out.row_summary_keys == ("Kyllä", "Ei")


# ---- the scatter names its axes by group ------------------------------------

@pytest.mark.parametrize("mode", ["image", "native"])
def test_a_scatter_still_draws_when_its_axis_group_is_renamed(mode):
    """`scatter_xy` is saved with the groups' own names; the series carries the
    renamed ones. The scatter must find its axes either way — and title them
    with the name the author gave."""
    import matplotlib
    matplotlib.use("Agg")
    from pptx import Presentation
    from pptx.util import Inches
    from reportbuilder.render.base import RenderContext, Slot, StyleSpec

    model, q, df = _study()
    spec = _spec(chart_type="scatter", scatter_xy=("Mies", "Nainen"),
                 series_label_overrides=(("Mies", "Miehet"),))
    series = engine.compute(q, spec, df, model)
    assert "Miehet" in series.segments
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    ctx = RenderContext(
        slide=slide, slot=Slot(slide_index=0, left=Inches(0.5), top=Inches(1.4),
                               width=Inches(8), height=Inches(4.4), name="s1"),
        style=StyleSpec(), spec=spec, series=series, fmt=spec.number_format)
    if mode == "native":
        from reportbuilder.render.native.scatter import build_scatter
        build_scatter(ctx)
        return
    from matplotlib.figure import Figure
    from reportbuilder.render.image.scatter import build_image_scatter
    seen = {}
    original = Figure.savefig

    def spy(self, *a, **k):
        seen["xlabel"] = self.axes[0].get_xlabel()
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        build_image_scatter(ctx)
    finally:
        Figure.savefig = original
    assert seen["xlabel"] == "Miehet"


def test_the_preview_request_carries_the_legend_names():
    from reportbuilder.api.routes_questions import ChartSpecBody, _chart_spec_from_body
    body = ChartSpecBody(question_ref="q", chart_type="vertical_bar",
                         series_label_overrides=[("Mies", "Miehet")])
    assert _chart_spec_from_body(body).series_label_overrides == (("Mies", "Miehet"),)


# ---- every chart type, through the deck build --------------------------------
#
# Built through `build_presentation` — the function the delivered deck is made
# with — rather than one builder at a time, so nothing between the settings and
# the picture can drop the rename unnoticed.

SPLIT_TYPES = ["vertical_bar", "horizontal_bar", "stacked_vertical_bar",
               "stacked_horizontal_bar", "line", "pie", "doughnut", "radar",
               "funnel", "combo", "combo-no-secondary", "vertical_bar-cross-tab",
               "stacked_horizontal_bar-cross-tab"]


def _deck_texts(chart_type: str, render_mode: str, renames) -> list[str]:
    import re
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from matplotlib.text import Text
    from reportbuilder.export.pptx_build import build_presentation
    from reportbuilder.model.report import Report

    model, q, df = _study()
    model = QuestionModel(variables=model.variables, questions=[q])
    opts = {"combo_secondary": "idx"} if chart_type == "combo" else {}
    kind, _, variant = chart_type.partition("-")
    extra = {"classifying_var_2": "age"} if variant == "cross-tab" else {}
    chart = _spec(chart_type=kind, options=opts, series_label_overrides=renames, **extra)
    report = Report(name="r", render_mode=render_mode, template_ref="", charts=(chart,))
    texts: list[str] = []
    original = Figure.savefig

    def spy(self, *a, **k):
        texts.extend(t.get_text() for t in self.findobj(Text))
        return original(self, *a, **k)

    Figure.savefig = spy
    try:
        prs = build_presentation(report, model, df)
    finally:
        Figure.savefig = original
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_chart", False) and shape.has_chart:
                xml = shape.chart.part.blob.decode("utf-8", "replace")
                texts.extend(re.findall(r"<c:v>([^<]*)</c:v>", xml))
    return texts


def _mentions(texts, word):
    import re
    return any(re.search(rf"(?<!\w){re.escape(word)}(?!\w)", t) for t in texts)


@pytest.mark.parametrize("chart_type", SPLIT_TYPES)
def test_a_renamed_group_is_drawn_under_its_new_name_image(chart_type):
    before = _deck_texts(chart_type, "image", ())
    after = _deck_texts(chart_type, "image", (("Mies", "Miehet"),))
    assert _mentions(before, "Mies"), f"{chart_type} draws no group names: {before}"
    assert _mentions(after, "Miehet"), after
    assert not _mentions(after, "Mies"), after


NATIVE_TYPES = ["vertical_bar", "horizontal_bar", "stacked_vertical_bar",
                "stacked_horizontal_bar", "line", "pie", "doughnut", "radar", "funnel"]


@pytest.mark.parametrize("chart_type", NATIVE_TYPES)
def test_a_renamed_group_is_drawn_under_its_new_name_native(chart_type):
    before = _deck_texts(chart_type, "native", ())
    after = _deck_texts(chart_type, "native", (("Mies", "Miehet"),))
    if not _mentions(before, "Mies"):
        pytest.skip(f"{chart_type} draws no group names natively on this data")
    assert _mentions(after, "Miehet"), after
    assert not _mentions(after, "Mies"), after


@pytest.mark.parametrize("chart_type", SPLIT_TYPES)
def test_a_report_saved_before_the_setting_builds_the_same_deck(chart_type):
    """Every stored report lacks the key. Loaded from JSON without it, the deck
    must say exactly what a slide with no renames says — same texts, same order."""
    import json
    import matplotlib
    matplotlib.use("Agg")
    from matplotlib.figure import Figure
    from matplotlib.text import Text
    from reportbuilder.export.pptx_build import build_presentation
    from reportbuilder.model.report import Report

    model, q, df = _study()
    model = QuestionModel(variables=model.variables, questions=[q])
    opts = {"combo_secondary": "idx"} if chart_type == "combo" else {}
    kind, _, variant = chart_type.partition("-")
    extra = {"classifying_var_2": "age"} if variant == "cross-tab" else {}
    report = Report(name="r", render_mode="image", template_ref="",
                    charts=(_spec(chart_type=kind, options=opts, **extra),))
    raw = json.loads(report_to_json(report))
    del raw["charts"][0]["series_label_overrides"]

    def texts_of(rep):
        seen: list[str] = []
        original = Figure.savefig

        def spy(self, *a, **k):
            seen.extend(t.get_text() for t in self.findobj(Text))
            return original(self, *a, **k)

        Figure.savefig = spy
        try:
            prs = build_presentation(rep, model, df)
        finally:
            Figure.savefig = original
        seen.extend(sh.text_frame.text for sl in prs.slides for sh in sl.shapes
                    if sh.has_text_frame)
        return seen

    assert texts_of(report_from_json(raw)) == texts_of(report)


def test_a_renamed_combo_secondary_series_is_drawn_under_its_new_name():
    model, q, df = _study()
    sec = engine.compute(q, _spec(chart_type="combo", options={"combo_secondary": "idx"}),
                         df, model).secondary_segments[0]
    after = _deck_texts("combo", "image", ((sec, "Työelämäindeksi"),))
    assert _mentions(after, "Työelämäindeksi"), after


def test_groups_on_this_slide_still_select_a_renamed_group():
    """"Groups on this slide" stores the data's own names; the rename is applied
    after the rows are narrowed, so ticking "Mies" still means those people."""
    only = _compute(classifying_values=("Mies",))
    renamed = _compute(classifying_values=("Mies",), series_label_overrides=(("Mies", "Miehet"),))
    assert "Miehet" in renamed.segments
    assert renamed.base_n["Total"] == only.base_n["Total"]
    assert renamed.applied_filter == ("Mies",)


def test_combinations_without_a_primary_map_are_renamed_by_part_too():
    """A battery crossed with two classifiers comes back as combinations
    ("Mies · Samaa mieltä") WITHOUT `segment_primary`. Found on real data: a
    rename of "Mies" changed nothing there. With two classifiers the series are
    combinations whatever else the result carries."""
    from reportbuilder.stats.series import Cell, SeriesResult
    segs = ("Mies · Samaa mieltä", "Nainen · Samaa mieltä")
    r = SeriesResult(categories=("13 Se suojaa",), segments=segs,
                     cells={("13 Se suojaa", s): Cell(pct=10.0) for s in segs},
                     base_n={**{s: 50 for s in segs}, "Total": 100}, statistic="pct")
    out = engine._series_relabelled(r, {"Mies": "Miehet"}, combinations=True)
    assert out.segments == ("Miehet · Samaa mieltä", "Nainen · Samaa mieltä")
    assert out.base_n["Miehet · Samaa mieltä"] == 50
